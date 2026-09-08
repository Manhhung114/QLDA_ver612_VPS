from __future__ import annotations

import functools
import hashlib
import io
import math
import multiprocessing as mp
import os
import tempfile
import threading
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path
from typing import Any


PATCH_MARKER = "V6.22 MULTICORE EXCEL V1"
_DEFAULT_PARALLEL_MIN_MB = 2.0
_DEFAULT_PARALLEL_MIN_SHEETS = 2
_PARALLEL_JOB_SEMAPHORE = threading.BoundedSemaphore(1)


def _env_int(name: str, default: int, minimum: int, maximum: int) -> int:
    try:
        value = int(str(os.environ.get(name, default)).strip())
    except Exception:
        value = int(default)
    return max(int(minimum), min(int(maximum), value))


def _env_float(name: str, default: float, minimum: float, maximum: float) -> float:
    try:
        value = float(str(os.environ.get(name, default)).strip())
    except Exception:
        value = float(default)
    return max(float(minimum), min(float(maximum), value))


def cpu_count() -> int:
    return max(1, int(os.cpu_count() or 1))


def configured_child_workers() -> int:
    """Child processes used by one heavy Excel parse.

    On a 4-vCPU VPS the default is 3 child processes. The Streamlit script
    process remains the fourth worker for BOQ/VO, or parses IPC metadata/GTHT
    while the children build workbook previews. BLAS stays at one thread per
    process to avoid nested oversubscription.
    """
    cpus = cpu_count()
    if cpus <= 1:
        return 1
    default = min(3, cpus - 1)
    requested = _env_int("QLDA_CPU_WORKERS", default, 1, 8)
    return max(1, min(requested, cpus - 1))


def parallel_min_bytes() -> int:
    mb = _env_float("QLDA_PARALLEL_EXCEL_MIN_MB", _DEFAULT_PARALLEL_MIN_MB, 0.0, 512.0)
    return int(mb * 1024 * 1024)


def parallel_min_sheets() -> int:
    return _env_int("QLDA_PARALLEL_MIN_SHEETS", _DEFAULT_PARALLEL_MIN_SHEETS, 1, 100)


def runtime_config() -> dict[str, Any]:
    return {
        "cpu_count": cpu_count(),
        "child_workers": configured_child_workers(),
        "parallel_min_bytes": parallel_min_bytes(),
        "parallel_min_sheets": parallel_min_sheets(),
        "start_method": _start_method(),
    }


def _start_method() -> str:
    available = set(mp.get_all_start_methods())
    requested = str(os.environ.get("QLDA_MP_START_METHOD", "forkserver")).strip().lower()
    if requested in available:
        return requested
    if "forkserver" in available:
        return "forkserver"
    if "spawn" in available:
        return "spawn"
    return next(iter(available))


def _mp_context():
    return mp.get_context(_start_method())


def _should_parallel(raw: bytes, sheet_count: int) -> bool:
    return bool(
        configured_child_workers() > 1
        and len(raw) >= parallel_min_bytes()
        and int(sheet_count) >= parallel_min_sheets()
    )


def _set_child_thread_limits() -> None:
    for name in (
        "OPENBLAS_NUM_THREADS",
        "OMP_NUM_THREADS",
        "MKL_NUM_THREADS",
        "NUMEXPR_NUM_THREADS",
        "VECLIB_MAXIMUM_THREADS",
    ):
        os.environ[name] = "1"
    os.environ.setdefault("MALLOC_ARENA_MAX", "2")


def _write_temp_workbook(raw: bytes, filename: str) -> str:
    suffix = Path(str(filename or "workbook.xlsx")).suffix.lower()
    if suffix not in {".xlsx", ".xlsm"}:
        suffix = ".xlsx"
    handle = tempfile.NamedTemporaryFile(prefix="qlda_excel_", suffix=suffix, delete=False)
    try:
        handle.write(raw)
        handle.flush()
        return handle.name
    finally:
        handle.close()


def _safe_unlink(path: str) -> None:
    try:
        os.unlink(path)
    except Exception:
        pass


def _weights_from_workbook(workbook, names: list[str]) -> dict[str, int]:
    out: dict[str, int] = {}
    for name in names:
        try:
            ws = workbook[name]
            rows = max(1, int(getattr(ws, "max_row", 1) or 1))
            cols = max(1, int(getattr(ws, "max_column", 1) or 1))
            out[name] = rows * min(cols, 40)
        except Exception:
            out[name] = 1
    return out


def _balanced_chunks(names: list[str], bucket_count: int, weights: dict[str, int]) -> list[list[str]]:
    if not names:
        return []
    count = max(1, min(int(bucket_count), len(names)))
    buckets: list[list[str]] = [[] for _ in range(count)]
    loads = [0] * count
    indexed = {name: idx for idx, name in enumerate(names)}
    ordered = sorted(names, key=lambda name: (-int(weights.get(name, 1)), indexed[name]))
    for name in ordered:
        target = min(range(count), key=lambda idx: loads[idx])
        buckets[target].append(name)
        loads[target] += int(weights.get(name, 1))
    return [bucket for bucket in buckets if bucket]


def _cache_parser(function, max_entries: int = 2):
    try:
        import streamlit as st

        @st.cache_data(show_spinner=False, max_entries=max_entries, ttl=3600)
        def cached(data: bytes, filename: str):
            return function(data, filename)

        return cached
    except Exception:
        return functools.lru_cache(maxsize=max_entries)(function)


def _multicore_meta(*, enabled: bool, worker_pids: list[int], strategy: str, fallback: str = "") -> dict[str, Any]:
    child_pids = sorted({int(pid) for pid in worker_pids if int(pid) != os.getpid()})
    return {
        "enabled": bool(enabled),
        "cpu_count": cpu_count(),
        "configured_child_workers": configured_child_workers(),
        "child_processes_used": len(child_pids),
        "worker_pids": child_pids,
        "parent_pid": os.getpid(),
        "start_method": _start_method(),
        "strategy": strategy,
        "fallback": str(fallback or ""),
    }


def _append_fallback_warning(result: dict[str, Any], exc: Exception, label: str) -> dict[str, Any]:
    parsed = dict(result or {})
    warnings = list(parsed.get("warnings") or [])
    warning = f"{label}: đa nhân không khả dụng trong lần xử lý này; hệ thống đã tự chuyển về 1 nhân ({type(exc).__name__})."
    if warning not in warnings:
        warnings.append(warning)
    parsed["warnings"] = warnings
    parsed["_multicore"] = _multicore_meta(
        enabled=False,
        worker_pids=[],
        strategy="automatic single-core fallback",
        fallback=type(exc).__name__,
    )
    return parsed


# ---------------------------------------------------------------------------
# BOQ: split visible sheets across the parent + child processes.
# ---------------------------------------------------------------------------


def _boq_process_loaded(boq, workbook, names: list[str]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for name in names:
        ws = workbook[name]
        records.append(
            {
                "name": name,
                "snapshot": boq._sheet_snapshot(ws),
                "summary_sheet": boq._extract_summary_sheet(ws) if boq._is_summary_sheet_name(name) else None,
                "parsed": boq._parse_sheet(ws),
            }
        )
    return records


def _boq_chunk_worker(path: str, names: list[str]) -> dict[str, Any]:
    _set_child_thread_limits()
    from openpyxl import load_workbook
    import boq_multisheet_v622 as boq

    wb = load_workbook(path, data_only=True, read_only=True, keep_links=False)
    try:
        records = _boq_process_loaded(boq, wb, names)
        return {"pid": os.getpid(), "records": records}
    finally:
        wb.close()


def _assemble_boq_result(boq, raw: bytes, filename: str, visible_names: list[str], records: list[dict[str, Any]], pids: list[int]) -> dict[str, Any]:
    order = {name: idx for idx, name in enumerate(visible_names)}
    records = sorted(records, key=lambda item: order.get(str(item.get("name") or ""), 10**9))
    snapshots = {item["name"]: item["snapshot"] for item in records}

    summary_sheet = None
    for item in records:
        if item.get("summary_sheet"):
            summary_sheet = item["summary_sheet"]
            break

    parsed = [item["parsed"] for item in records if item.get("parsed")]
    if not parsed:
        raise boq.BOQWorkbookError(
            "Không nhận diện được sheet BOQ chi tiết. Cần có cột Nội dung công việc và Thành tiền, "
            "hoặc bộ Khối lượng + Đơn giá."
        )

    summary = [
        {
            "sheet": item["sheet"],
            "line_count": int(item["line_count"]),
            "budget_total": float(item["budget_total"]),
            "header_row": int(item["header_row"]),
        }
        for item in parsed
    ]
    detail_items: list[dict[str, Any]] = []
    for sheet_data in parsed:
        sheet_name = sheet_data["sheet"]
        for item in sheet_data["items"]:
            detail_items.append({"sheet": sheet_name, **item})

    detail_total = float(sum(item["budget_total"] for item in detail_items))
    warnings: list[str] = []
    if summary_sheet:
        warnings.append(
            f"Sheet tổng hợp '{summary_sheet['sheet']}' được hiển thị riêng và không cộng lặp vào BOQ chi tiết."
        )
        before_tax = summary_sheet.get("before_tax_total")
        if before_tax is not None:
            tolerance = max(1.0, abs(float(before_tax)) * 1e-8)
            if abs(detail_total - float(before_tax)) > tolerance:
                warnings.append(
                    "Tổng các dòng BOQ chi tiết lệch so với 'Cộng giá trị trước thuế' trong sheet tổng hợp: "
                    f"{detail_total:,.0f} so với {float(before_tax):,.0f} VND."
                )

    before_tax_total = (
        float(summary_sheet["before_tax_total"])
        if summary_sheet and summary_sheet.get("before_tax_total") is not None
        else detail_total
    )
    vat_total = (
        float(summary_sheet["vat_total"])
        if summary_sheet and summary_sheet.get("vat_total") is not None
        else 0.0
    )
    after_tax_total = (
        float(summary_sheet["after_tax_total"])
        if summary_sheet and summary_sheet.get("after_tax_total") is not None
        else before_tax_total + vat_total
    )

    return {
        "filename": Path(str(filename or "BOQ.xlsx")).name,
        "batch_id": hashlib.sha256(raw).hexdigest()[:16],
        "workbook_sheet_names": list(visible_names),
        "workbook_sheets": snapshots,
        "summary_sheet": summary_sheet,
        "detected_sheets": [item["sheet"] for item in parsed],
        "summary": summary,
        "detail_items": detail_items,
        "detail_line_count": len(detail_items),
        "detail_grand_total": detail_total,
        "before_tax_total": before_tax_total,
        "vat_total": vat_total,
        "after_tax_total": after_tax_total,
        "grand_total": after_tax_total,
        "warnings": warnings,
        "_multicore": _multicore_meta(
            enabled=True,
            worker_pids=pids,
            strategy="BOQ sheets: parent + process pool",
        ),
    }


def _make_boq_parser(base_parse):
    def parse_boq_multicore(data: bytes, filename: str = "BOQ.xlsx"):
        raw = bytes(data or b"")
        if configured_child_workers() <= 1 or len(raw) < parallel_min_bytes():
            return base_parse(data, filename)

        import boq_multisheet_v622 as boq
        from openpyxl import load_workbook

        wb = None
        temp_path = ""
        try:
            wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True, keep_links=False)
            visible = [ws for ws in wb.worksheets if getattr(ws, "sheet_state", "visible") == "visible"]
            visible_names = [ws.title for ws in visible]
            if not _should_parallel(raw, len(visible_names)):
                wb.close()
                wb = None
                return base_parse(data, filename)

            weights = _weights_from_workbook(wb, visible_names)
            bucket_count = min(len(visible_names), configured_child_workers() + 1)
            chunks = _balanced_chunks(visible_names, bucket_count, weights)
            parent_chunk = chunks[0]
            child_chunks = chunks[1:]
            temp_path = _write_temp_workbook(raw, filename)
            pids: list[int] = []
            records: list[dict[str, Any]] = []

            with _PARALLEL_JOB_SEMAPHORE:
                with ProcessPoolExecutor(
                    max_workers=max(1, min(configured_child_workers(), len(child_chunks))),
                    mp_context=_mp_context(),
                ) as pool:
                    futures = [pool.submit(_boq_chunk_worker, temp_path, chunk) for chunk in child_chunks]
                    records.extend(_boq_process_loaded(boq, wb, parent_chunk))
                    pids.append(os.getpid())
                    for future in futures:
                        payload = future.result()
                        pids.append(int(payload.get("pid") or 0))
                        records.extend(payload.get("records") or [])

            wb.close()
            wb = None
            return _assemble_boq_result(boq, raw, filename, visible_names, records, pids)
        except Exception as exc:
            if wb is not None:
                try:
                    wb.close()
                except Exception:
                    pass
            try:
                result = base_parse(data, filename)
            except Exception:
                raise
            return _append_fallback_warning(result, exc, "BOQ")
        finally:
            if temp_path:
                _safe_unlink(temp_path)

    return _cache_parser(parse_boq_multicore, max_entries=2)


# ---------------------------------------------------------------------------
# IPC: child processes build workbook previews while the Streamlit parent
# performs adaptive role detection, payment summary and GTHT parsing.
# ---------------------------------------------------------------------------


def _snapshot_chunk_worker(path: str, names: list[str], kind: str) -> dict[str, Any]:
    _set_child_thread_limits()
    from openpyxl import load_workbook

    if kind == "ipc":
        import ipc_claim_v622 as module
    elif kind == "vo":
        import vo_claim_v622 as module
    else:
        import boq_multisheet_v622 as module

    wb = load_workbook(path, data_only=True, read_only=True, keep_links=False)
    try:
        records = []
        for name in names:
            records.append({"name": name, "snapshot": module._sheet_snapshot(wb[name])})
        return {"pid": os.getpid(), "records": records}
    finally:
        wb.close()


def _make_ipc_parser(base_parse):
    def parse_ipc_multicore(data: bytes, filename: str = "IPC.xlsx"):
        raw = bytes(data or b"")
        if configured_child_workers() <= 1 or len(raw) < parallel_min_bytes():
            return base_parse(data, filename)

        import ipc_claim_v622 as ipc
        from openpyxl import load_workbook

        wb = None
        temp_path = ""
        try:
            wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True, keep_links=False)
            visible = [ws for ws in wb.worksheets if getattr(ws, "sheet_state", "visible") == "visible"]
            visible_names = [ws.title for ws in visible]
            if not _should_parallel(raw, len(visible_names)):
                wb.close()
                wb = None
                return base_parse(data, filename)

            weights = _weights_from_workbook(wb, visible_names)
            child_count = min(configured_child_workers(), len(visible_names))
            chunks = _balanced_chunks(visible_names, child_count, weights)
            temp_path = _write_temp_workbook(raw, filename)
            pids: list[int] = []
            snapshots: dict[str, Any] = {}

            with _PARALLEL_JOB_SEMAPHORE:
                with ProcessPoolExecutor(
                    max_workers=max(1, len(chunks)),
                    mp_context=_mp_context(),
                ) as pool:
                    futures = [pool.submit(_snapshot_chunk_worker, temp_path, chunk, "ipc") for chunk in chunks]

                    # Main Streamlit process uses the fourth CPU on a 4-vCPU VPS.
                    declaration = ipc._first_sheet(wb, ("KHAI BÁO", "KHAI BAO"))
                    payment = ipc._first_sheet(wb, ("Thanh toán", "Thanh toan", "Payment"))
                    gtht = ipc._first_sheet(wb, ("GTHT", "Giá trị hoàn thành", "Gia tri hoan thanh"))
                    metadata = ipc._metadata_from_declaration(declaration)
                    summary = ipc._payment_summary(payment)
                    details = ipc._parse_gtht(gtht)
                    claim_no = str(metadata.get("claim_no") or "").strip()
                    if not claim_no and payment is not None:
                        claim_no = str(ipc._safe_cell(payment, "L8") or "").strip()

                    for future in futures:
                        payload = future.result()
                        pids.append(int(payload.get("pid") or 0))
                        for record in payload.get("records") or []:
                            snapshots[str(record.get("name") or "")] = record.get("snapshot")

            wb.close()
            wb = None
            snapshots = {name: snapshots[name] for name in visible_names if name in snapshots}

            if not claim_no:
                raise ipc.IPCWorkbookError("Không xác định được số Claim/Thanh toán lần trong file Excel.")
            claim_code = f"IPC-{claim_no.zfill(2)}" if claim_no.isdigit() else f"IPC-{claim_no}"
            warnings: list[str] = []
            if not details:
                warnings.append("Chưa nhận diện được dòng chi tiết GTHT; workbook vẫn được lưu và hiển thị đầy đủ theo sheet.")
            if not summary.get("requested_amount"):
                warnings.append("Chưa đọc được 'Giá trị thanh toán kỳ này' từ sheet Thanh toán.")

            return {
                "filename": Path(str(filename or "IPC.xlsx")).name,
                "batch_id": hashlib.sha256(raw).hexdigest()[:20],
                "claim_no": claim_no,
                "claim_code": claim_code,
                "metadata": metadata,
                "summary": summary,
                "workbook_sheet_names": list(visible_names),
                "workbook_sheets": snapshots,
                "detail_items": details,
                "detail_line_count": len(details),
                "warnings": warnings,
                "_multicore": _multicore_meta(
                    enabled=True,
                    worker_pids=pids,
                    strategy="IPC previews in process pool + adaptive/payment/GTHT in parent",
                ),
            }
        except Exception as exc:
            if wb is not None:
                try:
                    wb.close()
                except Exception:
                    pass
            try:
                result = base_parse(data, filename)
            except Exception:
                raise
            return _append_fallback_warning(result, exc, "IPC")
        finally:
            if temp_path:
                _safe_unlink(temp_path)

    return _cache_parser(parse_ipc_multicore, max_entries=2)


# ---------------------------------------------------------------------------
# VO: sequential detail parser per sheet + parallel sheet distribution.
# Negative values are preserved exactly; no abs()/max(0) is used for VO net.
# ---------------------------------------------------------------------------


def _tuple_value(values: tuple[Any, ...], col1: int | None, default: Any = "") -> Any:
    if not col1:
        return default
    idx = int(col1) - 1
    if idx < 0 or idx >= len(values):
        return default
    value = values[idx]
    return default if value is None else value


def _fast_vo_detail_sheet(vo, value_ws, formula_ws) -> list[dict[str, Any]]:
    headers = vo._header_map(value_ws)
    desc_col = vo._pick_col(headers, ("noi dung cong viec", "mo ta", "noi dung"))
    if not desc_col:
        return []
    seq_col = vo._pick_col(headers, ("stt", " tt "))
    unit_col = vo._pick_col(headers, ("don vi",))
    variation_col = vo._pick_col(headers, ("phat sinh tang giam",))
    increase_col = vo._pick_col(headers, ("phat sinh tang",), ("phat sinh tang giam",))
    decrease_col = vo._pick_col(headers, ("phat sinh giam",))
    contract_col = vo._pick_col(headers, ("theo hop dong",))
    actual_col = vo._pick_col(headers, ("thuc te thi cong",))
    mat_col = vo._pick_col(headers, ("don gia vat tu", "don gia vat lieu"))
    labor_col = vo._pick_col(headers, ("nhan cong",))
    amount_col = vo._pick_col(headers, ("thanh tien",))
    spec_col = vo._pick_col(headers, ("quy cach",), ("dieu chinh",))
    code_col = vo._pick_col(headers, ("ma hieu",), ("dieu chinh",))
    brand_col = vo._pick_col(headers, ("thuong hieu",))
    origin_col = vo._pick_col(headers, ("xuat xu",))
    note_col = vo._pick_col(headers, ("ghi chu",))

    cols = [
        col for col in (
            desc_col, seq_col, unit_col, variation_col, increase_col, decrease_col,
            contract_col, actual_col, mat_col, labor_col, amount_col, spec_col,
            code_col, brand_col, origin_col, note_col,
        ) if col
    ]
    max_col = max(cols) if cols else min(int(value_ws.max_column or 1), 40)
    max_row = int(value_ws.max_row or 0)
    values_iter = value_ws.iter_rows(min_row=4, max_row=max_row, min_col=1, max_col=max_col, values_only=True)
    formulas_iter = (
        formula_ws.iter_rows(min_row=4, max_row=max_row, min_col=1, max_col=max_col, values_only=True)
        if formula_ws is not None
        else None
    )

    items: list[dict[str, Any]] = []
    for offset, value_row in enumerate(values_iter, start=4):
        values = tuple(value_row)
        formula_values = tuple(next(formulas_iter)) if formulas_iter is not None else ()
        description = str(_tuple_value(values, desc_col) or "").strip()
        if not description:
            continue
        nd = vo._norm(description)
        if any(x in nd for x in (
            "cong gia tri truoc thue", "tong cong", "thue vat", "ban qlda", "tu van giam sat",
            "tong thau", "nha thau thi cong truc tiep",
        )):
            continue
        if amount_col and formula_values and vo._is_rollup_formula(_tuple_value(formula_values, amount_col)):
            continue

        contract_qty = vo._to_number(_tuple_value(values, contract_col)) if contract_col else None
        actual_qty = vo._to_number(_tuple_value(values, actual_col)) if actual_col else None
        variation_qty = vo._to_number(_tuple_value(values, variation_col)) if variation_col else None
        increase_qty = vo._to_number(_tuple_value(values, increase_col)) if increase_col else None
        decrease_qty = vo._to_number(_tuple_value(values, decrease_col)) if decrease_col else None
        if decrease_qty is not None and decrease_qty > 0:
            decrease_qty = -abs(decrease_qty)
        if variation_qty is None and contract_qty is not None and actual_qty is not None:
            variation_qty = actual_qty - contract_qty
        if variation_qty is None and (increase_qty is not None or decrease_qty is not None):
            variation_qty = float(increase_qty or 0) + float(decrease_qty or 0)

        material_price = vo._to_number(_tuple_value(values, mat_col)) if mat_col else None
        labor_price = vo._to_number(_tuple_value(values, labor_col)) if labor_col else None
        direct_amount = vo._to_number(_tuple_value(values, amount_col)) if amount_col else None
        amount = float(direct_amount or 0)
        qty = float(variation_qty or 0)
        if abs(qty) < 1e-12 and abs(amount) < 0.5:
            continue
        kind = (
            "Tăng" if amount > 0.5 or (abs(amount) <= 0.5 and qty > 0)
            else "Giảm" if amount < -0.5 or qty < 0
            else "Chưa định giá"
        )
        items.append({
            "sheet_name": value_ws.title,
            "row_no": offset,
            "seq": str(_tuple_value(values, seq_col) or "").strip() if seq_col else "",
            "description": description,
            "unit": str(_tuple_value(values, unit_col) or "").strip() if unit_col else "",
            "contract_qty": float(contract_qty or 0),
            "actual_qty": float(actual_qty or 0),
            "increase_qty": float(increase_qty or 0),
            "decrease_qty": float(decrease_qty or 0),
            "variation_qty": qty,
            "spec": str(_tuple_value(values, spec_col) or "").strip() if spec_col else "",
            "item_code": str(_tuple_value(values, code_col) or "").strip() if code_col else "",
            "brand": str(_tuple_value(values, brand_col) or "").strip() if brand_col else "",
            "origin": str(_tuple_value(values, origin_col) or "").strip() if origin_col else "",
            "material_unit_price": float(material_price or 0),
            "labor_unit_price": float(labor_price or 0),
            "unit_price_total": float((material_price or 0) + (labor_price or 0)),
            "variation_amount": amount,
            "variation_kind": kind,
            "note": str(_tuple_value(values, note_col) or "").strip() if note_col else "",
        })
    return items


def _vo_process_loaded(vo, wb, wbf, names: list[str], source_sheets: set[str]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for name in names:
        value_ws = wb[name]
        formula_ws = wbf[name] if wbf is not None and name in wbf.sheetnames else None
        records.append(
            {
                "name": name,
                "snapshot": vo._sheet_snapshot(value_ws),
                "detail_items": _fast_vo_detail_sheet(vo, value_ws, formula_ws) if name in source_sheets else [],
            }
        )
    return records


def _vo_chunk_worker(path: str, names: list[str], source_names: list[str]) -> dict[str, Any]:
    _set_child_thread_limits()
    from openpyxl import load_workbook
    import vo_claim_v622 as vo

    source_sheets = set(source_names)
    wb = load_workbook(path, data_only=True, read_only=True, keep_links=False)
    needs_formula = any(name in source_sheets for name in names)
    wbf = load_workbook(path, data_only=False, read_only=True, keep_links=False) if needs_formula else None
    try:
        return {
            "pid": os.getpid(),
            "records": _vo_process_loaded(vo, wb, wbf, names, source_sheets),
        }
    finally:
        wb.close()
        if wbf is not None:
            wbf.close()


def _make_vo_parser(base_parse):
    def parse_vo_multicore(data: bytes, filename: str = "VO.xlsx"):
        raw = bytes(data or b"")
        if configured_child_workers() <= 1 or len(raw) < parallel_min_bytes():
            return base_parse(data, filename)

        import vo_claim_v622 as vo
        from openpyxl import load_workbook

        wb = wbf = None
        temp_path = ""
        try:
            wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True, keep_links=False)
            wbf = load_workbook(io.BytesIO(raw), data_only=False, read_only=True, keep_links=False)
            visible = [ws for ws in wb.worksheets if getattr(ws, "sheet_state", "visible") == "visible"]
            visible_names = [ws.title for ws in visible]
            if not _should_parallel(raw, len(visible_names)):
                wb.close(); wbf.close()
                wb = wbf = None
                return base_parse(data, filename)

            summary_ws = vo._find_summary_sheet(wb)
            if summary_ws is None:
                raise vo.VOWorkbookError("Không tìm thấy sheet Tổng hợp VO.")
            formula_summary = wbf[summary_ws.title] if summary_ws.title in wbf.sheetnames else None
            metadata = vo._summary_metadata(summary_ws, filename)
            summary, summary_lines, subtotal_row = vo._summary_values(summary_ws)
            valid_names = set(wb.sheetnames)
            source_sheets = vo._formula_source_sheets(formula_summary, subtotal_row, valid_names)
            if not source_sheets:
                source_sheets = vo._fallback_source_sheets(wb, summary_ws.title)
            source_set = set(source_sheets)

            weights = _weights_from_workbook(wb, visible_names)
            bucket_count = min(len(visible_names), configured_child_workers() + 1)
            chunks = _balanced_chunks(visible_names, bucket_count, weights)
            parent_chunk = chunks[0]
            child_chunks = chunks[1:]
            temp_path = _write_temp_workbook(raw, filename)
            pids: list[int] = []
            records: list[dict[str, Any]] = []

            with _PARALLEL_JOB_SEMAPHORE:
                with ProcessPoolExecutor(
                    max_workers=max(1, min(configured_child_workers(), len(child_chunks))),
                    mp_context=_mp_context(),
                ) as pool:
                    futures = [pool.submit(_vo_chunk_worker, temp_path, chunk, list(source_sheets)) for chunk in child_chunks]
                    records.extend(_vo_process_loaded(vo, wb, wbf, parent_chunk, source_set))
                    pids.append(os.getpid())
                    for future in futures:
                        payload = future.result()
                        pids.append(int(payload.get("pid") or 0))
                        records.extend(payload.get("records") or [])

            wb.close(); wbf.close()
            wb = wbf = None
            order = {name: idx for idx, name in enumerate(visible_names)}
            records.sort(key=lambda item: order.get(str(item.get("name") or ""), 10**9))
            snapshots = {item["name"]: item["snapshot"] for item in records}
            detail_items: list[dict[str, Any]] = []
            for item in records:
                detail_items.extend(item.get("detail_items") or [])
            source_order = {name: idx for idx, name in enumerate(source_sheets)}
            detail_items.sort(key=lambda item: (source_order.get(str(item.get("sheet_name") or ""), 10**9), int(item.get("row_no") or 0)))

            priced_items = [x for x in detail_items if abs(float(x.get("variation_amount") or 0)) >= 0.5]
            increase_amount = sum(max(0.0, float(x.get("variation_amount") or 0)) for x in priced_items)
            decrease_amount = sum(min(0.0, float(x.get("variation_amount") or 0)) for x in priced_items)
            detail_net = increase_amount + decrease_amount
            subtotal = float(summary.get("subtotal_before_vat") or 0)
            discrepancy = detail_net - subtotal

            warnings = list(metadata.pop("warnings", []))
            if not source_sheets:
                warnings.append("Không nhận được các sheet chi tiết từ công thức Tổng hợp; parser dùng dò nội dung.")
            if abs(discrepancy) > max(1.0, abs(subtotal) * 0.0001):
                warnings.append(
                    f"Tổng dòng chi tiết có giá {vo._money(detail_net)} VND lệch {vo._money(discrepancy)} VND "
                    f"so với Tổng hợp {vo._money(subtotal)} VND. Không tự sửa số liệu; cần kiểm tra workbook."
                )

            confidence = 55
            confidence += 15 if metadata.get("vo_code") else 0
            confidence += 10 if subtotal_row else 0
            confidence += 10 if source_sheets else 0
            confidence += 10 if abs(discrepancy) <= max(1.0, abs(subtotal) * 0.0001) else 0
            confidence = min(100, confidence)

            return {
                "schema": "qlda_vo_excel_v2",
                "filename": str(filename or "VO.xlsx"),
                "batch_id": hashlib.sha256(raw).hexdigest(),
                "metadata": metadata,
                "vo_no": int(metadata["vo_no"]),
                "vo_code": str(metadata["vo_code"]),
                "summary": {
                    **summary,
                    "increase_amount": float(increase_amount),
                    "decrease_amount": float(decrease_amount),
                    "detail_net_amount": float(detail_net),
                    "detail_discrepancy": float(discrepancy),
                },
                "summary_lines": summary_lines,
                "source_sheets": source_sheets,
                "detail_items": detail_items,
                "workbook": snapshots,
                "workbook_sheet_names": list(visible_names),
                "confidence_pct": int(confidence),
                "warnings": warnings,
                "_multicore": _multicore_meta(
                    enabled=True,
                    worker_pids=pids,
                    strategy="VO sheets: parent + process pool; sequential row streaming per sheet",
                ),
            }
        except Exception as exc:
            for book in (wb, wbf):
                if book is not None:
                    try:
                        book.close()
                    except Exception:
                        pass
            try:
                result = base_parse(data, filename)
            except Exception:
                raise
            return _append_fallback_warning(result, exc, "VO")
        finally:
            if temp_path:
                _safe_unlink(temp_path)

    return _cache_parser(parse_vo_multicore, max_entries=2)


def install_multicore_excel() -> None:
    """Install multicore parsers for BOQ, IPC and VO.

    Call after IPC fast/summary/adaptive patches and before the filename Claim
    identity guard. Database writes, approvals, Google Drive and AI calls remain
    single-transaction/network operations; only CPU-heavy Excel work is split
    across processes.
    """
    import boq_multisheet_v622 as boq
    import ipc_claim_v622 as ipc
    import vo_claim_v622 as vo

    if not getattr(boq, "_qlda_multicore_excel_installed", False):
        boq._qlda_singlecore_parse_boq_workbook = boq.parse_boq_workbook
        boq.parse_boq_workbook = _make_boq_parser(boq.parse_boq_workbook)
        boq._qlda_multicore_excel_installed = True
        boq._qlda_multicore_excel_marker = PATCH_MARKER

    if not getattr(ipc, "_qlda_multicore_excel_installed", False):
        ipc._qlda_singlecore_parse_ipc_workbook = ipc.parse_ipc_workbook
        ipc.parse_ipc_workbook = _make_ipc_parser(ipc.parse_ipc_workbook)
        ipc._qlda_multicore_excel_installed = True
        ipc._qlda_multicore_excel_marker = PATCH_MARKER

    if not getattr(vo, "_qlda_multicore_excel_installed", False):
        vo._qlda_singlecore_parse_vo_workbook = vo.parse_vo_workbook
        vo.parse_vo_workbook = _make_vo_parser(vo.parse_vo_workbook)
        vo._qlda_multicore_excel_installed = True
        vo._qlda_multicore_excel_marker = PATCH_MARKER
