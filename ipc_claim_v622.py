from __future__ import annotations

import base64
import gzip
import hashlib
import io
import json
import math
import re
import unicodedata
import uuid
from datetime import date, datetime
from numbers import Real
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


PATCH_VERSION = "V6.22 IPC CLAIM V1"
PAYLOAD_PREFIX = "gz1:"
MAX_WORKBOOK_BYTES = 120 * 1024 * 1024
MAX_PREVIEW_ROWS = 5000
MAX_PREVIEW_COLS = 32

CLAIMS_TABLE = "payment_claims"
ITEMS_TABLE = "payment_claim_items"
WORKBOOK_TABLE = "payment_claim_workbooks"
REVISIONS_TABLE = "payment_claim_revisions"


class IPCWorkbookError(ValueError):
    pass


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _norm(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.lower().replace("đ", "d")
    text = re.sub(r"[^a-z0-9%]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _to_number(value: Any) -> float | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    if isinstance(value, Real):
        number = float(value)
        return number if math.isfinite(number) else None
    text = str(value).strip()
    if not text or text.startswith("="):
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    text = re.sub(r"[^0-9,\.\-]", "", text)
    if not text or text in {"-", ".", ","}:
        return None
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        parts = text.split(",")
        if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3):
            text = "".join(parts)
        else:
            text = text.replace(",", ".")
    try:
        number = float(text)
        if negative:
            number = -abs(number)
        return number if math.isfinite(number) else None
    except Exception:
        return None


def format_table_number(value: Any) -> Any:
    if isinstance(value, bool) or not isinstance(value, Real):
        return value
    number = float(value)
    if not math.isfinite(number):
        return value
    if abs(number - round(number)) < 1e-9:
        return f"{number:,.0f}"
    return f"{number:,.4f}".rstrip("0").rstrip(".")


def _date_text(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except Exception:
            pass
    return text


def _display_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return value


def _sheet_snapshot(ws) -> dict[str, Any]:
    max_row = min(int(ws.max_row or 1), MAX_PREVIEW_ROWS)
    max_col = min(int(ws.max_column or 1), MAX_PREVIEW_COLS)
    rows: list[list[Any]] = []
    for row_no, values in enumerate(
        ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col, values_only=True),
        start=1,
    ):
        rows.append([row_no] + [_display_value(v) for v in values])
    return {
        "sheet": ws.title,
        "columns": ["Dòng"] + [get_column_letter(i) for i in range(1, max_col + 1)],
        "rows": rows,
        "row_count": int(ws.max_row or 0),
        "col_count": int(ws.max_column or 0),
        "truncated": bool((ws.max_row or 0) > max_row or (ws.max_column or 0) > max_col),
    }


def _safe_cell(ws, ref: str, default: Any = "") -> Any:
    try:
        value = ws[ref].value
        return default if value is None else value
    except Exception:
        return default


def _first_sheet(workbook, aliases: tuple[str, ...]):
    wanted = [_norm(x) for x in aliases]
    for ws in workbook.worksheets:
        name = _norm(ws.title)
        if name in wanted or any(x and x in name for x in wanted):
            return ws
    return None


def _metadata_from_declaration(ws) -> dict[str, Any]:
    if ws is None:
        return {}
    return {
        "project": str(_safe_cell(ws, "B4") or "").strip(),
        "location": str(_safe_cell(ws, "B5") or "").strip(),
        "package": str(_safe_cell(ws, "B6") or "").strip(),
        "contractor": str(_safe_cell(ws, "B7") or "").strip(),
        "account_name": str(_safe_cell(ws, "B8") or "").strip(),
        "account_no": str(_safe_cell(ws, "B9") or "").strip(),
        "bank": str(_safe_cell(ws, "B10") or "").strip(),
        "bank_branch": str(_safe_cell(ws, "B11") or "").strip(),
        "claim_no": str(_safe_cell(ws, "B12") or "").strip(),
        "from_date": _date_text(_safe_cell(ws, "B13")),
        "to_date": _date_text(_safe_cell(ws, "E13")),
        "currency": str(_safe_cell(ws, "B14") or "VNĐ").strip() or "VNĐ",
        "contract_no": str(_safe_cell(ws, "B15") or "").strip(),
    }


def _payment_summary(ws) -> dict[str, Any]:
    if ws is None:
        return {}

    def num(ref: str) -> float:
        return float(_to_number(_safe_cell(ws, ref)) or 0)

    return {
        "contract_value": num("D11"),
        "contract_advance": num("D12"),
        "material_advance": num("D13"),
        "cumulative_acceptance": num("D14"),
        "cumulative_material_acceptance": num("D15"),
        "cumulative_installation_acceptance": num("F16") or num("D16"),
        "cumulative_material_deduction": num("F17") or num("D17"),
        "cumulative_completed": num("D18"),
        "previous_approved": num("D19"),
        "cumulative_retention": num("D25"),
        "cumulative_advance_recovery": num("D28"),
        "cumulative_material_advance_recovery": num("D29"),
        "prior_total_deductions": num("D30"),
        "current_gross": num("D31"),
        "current_deductions": num("K24"),
        "requested_amount": num("K30"),
        "amount_in_words": str(_safe_cell(ws, "C32") or "").strip(),
    }


def _is_total_label(value: Any) -> bool:
    text = _norm(value)
    return bool(
        text in {"tong", "tong cong", "vat 8", "vat 10", "tong cong bao gom thue vat"}
        or text.startswith("tong cong")
        or text.startswith("vat ")
    )


def _parse_gtht(ws) -> list[dict[str, Any]]:
    if ws is None:
        return []
    header_row = None
    for row_no in range(1, min(int(ws.max_row or 1), 45) + 1):
        if "ten cong tac" in _norm(_safe_cell(ws, f"B{row_no}")):
            header_row = row_no
            break
    if header_row is None:
        return []

    start_row = header_row + 4
    out: list[dict[str, Any]] = []
    for row_no in range(start_row, int(ws.max_row or start_row) + 1):
        description = str(_safe_cell(ws, f"B{row_no}") or "").strip()
        if not description:
            continue
        if _is_total_label(description):
            if "tong cong" in _norm(description):
                break
            continue

        seq = _to_number(_safe_cell(ws, f"A{row_no}"))
        contract_qty = _to_number(_safe_cell(ws, f"C{row_no}"))
        material_current_qty = _to_number(_safe_cell(ws, f"M{row_no}")) or 0
        material_cumulative_qty = _to_number(_safe_cell(ws, f"N{row_no}")) or 0
        installation_current_pct = _to_number(_safe_cell(ws, f"P{row_no}")) or 0
        installation_cumulative_pct = _to_number(_safe_cell(ws, f"Q{row_no}")) or 0
        material_current_value = _to_number(_safe_cell(ws, f"S{row_no}")) or 0
        material_cumulative_value = _to_number(_safe_cell(ws, f"T{row_no}")) or 0
        installation_current_value = _to_number(_safe_cell(ws, f"V{row_no}")) or 0
        installation_cumulative_value = _to_number(_safe_cell(ws, f"W{row_no}")) or 0
        deduction_current = _to_number(_safe_cell(ws, f"Y{row_no}")) or 0
        deduction_cumulative = _to_number(_safe_cell(ws, f"Z{row_no}")) or 0

        # Summary/group rows in GTHT repeat child totals. Keep only real BOQ rows.
        is_detail = bool((seq is not None and seq > 0) or (contract_qty is not None and contract_qty > 0))
        if not is_detail:
            continue

        out.append(
            {
                "sheet_name": ws.title,
                "row_no": row_no,
                "seq": seq,
                "boq_item": description,
                "contract_qty": float(contract_qty or 0),
                "unit": str(_safe_cell(ws, f"D{row_no}") or "").strip(),
                "spec": str(_safe_cell(ws, f"E{row_no}") or "").strip(),
                "item_code": str(_safe_cell(ws, f"F{row_no}") or "").strip(),
                "brand": str(_safe_cell(ws, f"G{row_no}") or "").strip(),
                "origin": str(_safe_cell(ws, f"H{row_no}") or "").strip(),
                "material_unit_price": float(_to_number(_safe_cell(ws, f"I{row_no}")) or 0),
                "labor_unit_price": float(_to_number(_safe_cell(ws, f"J{row_no}")) or 0),
                "contract_amount": float(_to_number(_safe_cell(ws, f"K{row_no}")) or 0),
                "material_previous_qty": float(_to_number(_safe_cell(ws, f"L{row_no}")) or 0),
                "material_current_qty": float(material_current_qty),
                "material_cumulative_qty": float(material_cumulative_qty),
                "installation_previous_pct": float(_to_number(_safe_cell(ws, f"O{row_no}")) or 0),
                "installation_current_pct": float(installation_current_pct),
                "installation_cumulative_pct": float(installation_cumulative_pct),
                "material_previous_value": float(_to_number(_safe_cell(ws, f"R{row_no}")) or 0),
                "material_current_value": float(material_current_value),
                "material_cumulative_value": float(material_cumulative_value),
                "installation_previous_value": float(_to_number(_safe_cell(ws, f"U{row_no}")) or 0),
                "installation_current_value": float(installation_current_value),
                "installation_cumulative_value": float(installation_cumulative_value),
                "deduction_previous": float(_to_number(_safe_cell(ws, f"X{row_no}")) or 0),
                "deduction_current": float(deduction_current),
                "deduction_cumulative": float(deduction_cumulative),
                "current_value": float(material_current_value + installation_current_value - deduction_current),
                "cumulative_value": float(material_cumulative_value + installation_cumulative_value - deduction_cumulative),
                "completion_ratio": float(_to_number(_safe_cell(ws, f"AA{row_no}")) or 0),
                "note": str(_safe_cell(ws, f"AB{row_no}") or "").strip(),
                "cost_code": str(_safe_cell(ws, f"AC{row_no}") or "").strip(),
                "system": str(_safe_cell(ws, f"AD{row_no}") or "").strip(),
            }
        )
    return out


def parse_ipc_workbook(data: bytes, filename: str = "IPC.xlsx") -> dict[str, Any]:
    raw = bytes(data or b"")
    if not raw:
        raise IPCWorkbookError("File IPC/Claim đang trống.")
    if len(raw) > MAX_WORKBOOK_BYTES:
        raise IPCWorkbookError("File IPC lớn hơn 120 MB; hãy tối ưu hoặc tách file trước khi nhập.")
    suffix = Path(str(filename or "")).suffix.lower()
    if suffix and suffix not in {".xlsx", ".xlsm"}:
        raise IPCWorkbookError("Chỉ hỗ trợ file IPC Excel .xlsx hoặc .xlsm.")

    try:
        wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    except Exception as exc:
        raise IPCWorkbookError(f"Không đọc được workbook IPC: {exc}") from exc

    try:
        visible = [ws for ws in wb.worksheets if getattr(ws, "sheet_state", "visible") == "visible"]
        snapshots = {ws.title: _sheet_snapshot(ws) for ws in visible}
        declaration = _first_sheet(wb, ("KHAI BÁO", "KHAI BAO"))
        payment = _first_sheet(wb, ("Thanh toán", "Thanh toan", "Payment"))
        gtht = _first_sheet(wb, ("GTHT", "Giá trị hoàn thành", "Gia tri hoan thanh"))
        metadata = _metadata_from_declaration(declaration)
        summary = _payment_summary(payment)
        details = _parse_gtht(gtht)
    finally:
        wb.close()

    claim_no = str(metadata.get("claim_no") or "").strip()
    if not claim_no and payment is not None:
        claim_no = str(_safe_cell(payment, "L8") or "").strip()
    if not claim_no:
        raise IPCWorkbookError("Không xác định được số Claim/Thanh toán lần trong file Excel.")

    claim_code = f"IPC-{claim_no.zfill(2)}" if claim_no.isdigit() else f"IPC-{claim_no}"
    warnings: list[str] = []
    if not details:
        warnings.append("Chưa nhận diện được dòng chi tiết GTHT; workbook vẫn được lưu và hiển thị đầy đủ theo sheet.")
    if not summary.get("requested_amount"):
        warnings.append("Chưa đọc được 'Giá trị thanh toán kỳ này' từ sheet Thanh toán.")

    return {
        "filename": Path(str(filename or "IPC.xlsx")).name,
        "batch_id": hashlib.sha256(raw).hexdigest()[:20],
        "claim_no": claim_no,
        "claim_code": claim_code,
        "metadata": metadata,
        "summary": summary,
        "workbook_sheet_names": [ws.title for ws in visible],
        "workbook_sheets": snapshots,
        "detail_items": details,
        "detail_line_count": len(details),
        "warnings": warnings,
    }


def _rowdict(row: Any) -> dict[str, Any]:
    if row is None:
        return {}
    if isinstance(row, dict):
        return dict(row)
    try:
        return {str(k): row[k] for k in row.keys()}
    except Exception:
        try:
            return dict(row)
        except Exception:
            return {}


def _ensure_tables(connection) -> None:
    connection.execute(
        f"""
        CREATE TABLE IF NOT EXISTS {CLAIMS_TABLE}(
            claim_id TEXT PRIMARY KEY,
            project_id INTEGER NOT NULL,
            claim_no TEXT NOT NULL,
            claim_code TEXT NOT NULL,
            filename TEXT DEFAULT '',
            batch_id TEXT DEFAULT '',
            contractor TEXT DEFAULT '',
            contract_no TEXT DEFAULT '',
            package_name TEXT DEFAULT '',
            from_date TEXT DEFAULT '',
            to_date TEXT DEFAULT '',
            contract_value REAL DEFAULT 0,
            requested_amount REAL DEFAULT 0,
            approved_amount REAL DEFAULT 0,
            disbursed_amount REAL DEFAULT 0,
            certified_cumulative REAL DEFAULT 0,
            previous_approved REAL DEFAULT 0,
            retention_cumulative REAL DEFAULT 0,
            advance_amount REAL DEFAULT 0,
            advance_recovery REAL DEFAULT 0,
            current_deductions REAL DEFAULT 0,
            payment_status TEXT DEFAULT 'Nháp',
            disbursement_date TEXT DEFAULT '',
            latest_revision INTEGER DEFAULT 0,
            note TEXT DEFAULT '',
            created_at TEXT DEFAULT '',
            updated_at TEXT DEFAULT '',
            UNIQUE(project_id, claim_no)
        )
        """
    )
    connection.execute(
        f"""
        CREATE TABLE IF NOT EXISTS {ITEMS_TABLE}(
            claim_id TEXT NOT NULL,
            project_id INTEGER NOT NULL,
            sheet_name TEXT DEFAULT '',
            row_no INTEGER NOT NULL,
            seq REAL DEFAULT 0,
            boq_item TEXT NOT NULL,
            contract_qty REAL DEFAULT 0,
            unit TEXT DEFAULT '',
            spec TEXT DEFAULT '',
            item_code TEXT DEFAULT '',
            brand TEXT DEFAULT '',
            origin TEXT DEFAULT '',
            material_unit_price REAL DEFAULT 0,
            labor_unit_price REAL DEFAULT 0,
            contract_amount REAL DEFAULT 0,
            material_previous_qty REAL DEFAULT 0,
            material_current_qty REAL DEFAULT 0,
            material_cumulative_qty REAL DEFAULT 0,
            installation_previous_pct REAL DEFAULT 0,
            installation_current_pct REAL DEFAULT 0,
            installation_cumulative_pct REAL DEFAULT 0,
            material_previous_value REAL DEFAULT 0,
            material_current_value REAL DEFAULT 0,
            material_cumulative_value REAL DEFAULT 0,
            installation_previous_value REAL DEFAULT 0,
            installation_current_value REAL DEFAULT 0,
            installation_cumulative_value REAL DEFAULT 0,
            deduction_previous REAL DEFAULT 0,
            deduction_current REAL DEFAULT 0,
            deduction_cumulative REAL DEFAULT 0,
            current_value REAL DEFAULT 0,
            cumulative_value REAL DEFAULT 0,
            completion_ratio REAL DEFAULT 0,
            note TEXT DEFAULT '',
            cost_code TEXT DEFAULT '',
            system TEXT DEFAULT '',
            PRIMARY KEY(claim_id, sheet_name, row_no)
        )
        """
    )
    connection.execute(
        f"""
        CREATE TABLE IF NOT EXISTS {WORKBOOK_TABLE}(
            claim_id TEXT PRIMARY KEY,
            project_id INTEGER NOT NULL,
            filename TEXT NOT NULL,
            batch_id TEXT NOT NULL,
            payload TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
        """
    )
    connection.execute(
        f"""
        CREATE TABLE IF NOT EXISTS {REVISIONS_TABLE}(
            revision_id TEXT PRIMARY KEY,
            claim_id TEXT NOT NULL,
            project_id INTEGER NOT NULL,
            revision_no INTEGER NOT NULL,
            filename TEXT NOT NULL,
            batch_id TEXT NOT NULL,
            payload TEXT NOT NULL,
            created_at TEXT NOT NULL,
            UNIQUE(claim_id, revision_no)
        )
        """
    )
    connection.execute(f"CREATE INDEX IF NOT EXISTS idx_payment_claim_project ON {CLAIMS_TABLE}(project_id, claim_no)")
    connection.execute(f"CREATE INDEX IF NOT EXISTS idx_payment_claim_items_project ON {ITEMS_TABLE}(project_id, claim_id)")
    connection.execute(f"CREATE INDEX IF NOT EXISTS idx_payment_claim_revision ON {REVISIONS_TABLE}(claim_id, revision_no)")


def _encode_result(result: dict[str, Any]) -> str:
    raw = json.dumps(result, ensure_ascii=False, separators=(",", ":"), allow_nan=False).encode("utf-8")
    return PAYLOAD_PREFIX + base64.b64encode(gzip.compress(raw, compresslevel=6)).decode("ascii")


def _decode_result(payload: str) -> dict[str, Any]:
    text = str(payload or "")
    if not text.startswith(PAYLOAD_PREFIX):
        raise ValueError("Dữ liệu workbook IPC đã lưu không hợp lệ.")
    packed = base64.b64decode(text[len(PAYLOAD_PREFIX):].encode("ascii"))
    data = json.loads(gzip.decompress(packed).decode("utf-8"))
    if not isinstance(data, dict):
        raise ValueError("Dữ liệu IPC đã lưu không phải workbook hợp lệ.")
    return data


def _sync_payment_tracking(connection, claim: dict[str, Any]) -> None:
    project_id = int(claim["project_id"])
    payment_code = str(claim["claim_code"])
    connection.execute(
        "DELETE FROM payment_tracking WHERE project_id=? AND payment_code=?",
        (project_id, payment_code),
    )
    connection.execute(
        """INSERT INTO payment_tracking(
               project_id,payment_code,task_ref,installment,certified_cumulative,paid_amount,
               advance_amount,advance_recovery,planned_disbursement_pct,payment_status,
               payment_date,note,created_at,updated_at
           ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (
            project_id,
            payment_code,
            "",
            str(claim.get("claim_no") or ""),
            float(claim.get("certified_cumulative") or 0),
            float(claim.get("disbursed_amount") or 0),
            float(claim.get("advance_amount") or 0),
            float(claim.get("advance_recovery") or 0),
            0,
            str(claim.get("payment_status") or ""),
            str(claim.get("disbursement_date") or claim.get("to_date") or ""),
            f"[IPC_CLAIM] requested={float(claim.get('requested_amount') or 0):.4f}; approved={float(claim.get('approved_amount') or 0):.4f}; {str(claim.get('note') or '')}",
            str(claim.get("created_at") or _now()),
            str(claim.get("updated_at") or _now()),
        ),
    )


def save_ipc_claim(db, project_id: int, result: dict[str, Any]) -> dict[str, Any]:
    pid = int(project_id)
    claim_no = str(result.get("claim_no") or "").strip()
    if not claim_no:
        raise ValueError("Claim chưa có số/kỳ thanh toán.")
    claim_code = str(result.get("claim_code") or f"IPC-{claim_no}").strip()
    batch_id = str(result.get("batch_id") or "")
    filename = str(result.get("filename") or "IPC.xlsx")
    metadata = dict(result.get("metadata") or {})
    summary = dict(result.get("summary") or {})
    payload = _encode_result(result)
    now = _now()

    with db.connect() as connection:
        _ensure_tables(connection)
        old_row = connection.execute(
            f"SELECT * FROM {CLAIMS_TABLE} WHERE project_id=? AND claim_no=?",
            (pid, claim_no),
        ).fetchone()
        old = _rowdict(old_row)
        claim_id = str(old.get("claim_id") or uuid.uuid4().hex)
        old_batch = str(old.get("batch_id") or "")
        latest_revision = int(old.get("latest_revision") or 0)
        new_revision = latest_revision + 1 if old_batch != batch_id else latest_revision
        if old_row is None:
            new_revision = 0

        claim = {
            "claim_id": claim_id,
            "project_id": pid,
            "claim_no": claim_no,
            "claim_code": claim_code,
            "filename": filename,
            "batch_id": batch_id,
            "contractor": str(metadata.get("contractor") or ""),
            "contract_no": str(metadata.get("contract_no") or ""),
            "package_name": str(metadata.get("package") or ""),
            "from_date": str(metadata.get("from_date") or ""),
            "to_date": str(metadata.get("to_date") or ""),
            "contract_value": float(summary.get("contract_value") or 0),
            "requested_amount": float(summary.get("requested_amount") or 0),
            "approved_amount": float(old.get("approved_amount") or 0),
            "disbursed_amount": float(old.get("disbursed_amount") or 0),
            "certified_cumulative": float(summary.get("cumulative_completed") or summary.get("cumulative_acceptance") or 0),
            "previous_approved": float(summary.get("previous_approved") or 0),
            "retention_cumulative": float(summary.get("cumulative_retention") or 0),
            "advance_amount": float(summary.get("contract_advance") or 0),
            "advance_recovery": float(summary.get("cumulative_advance_recovery") or 0),
            "current_deductions": float(summary.get("current_deductions") or 0),
            "payment_status": str(old.get("payment_status") or "Nháp"),
            "disbursement_date": str(old.get("disbursement_date") or ""),
            "latest_revision": int(new_revision),
            "note": str(old.get("note") or ""),
            "created_at": str(old.get("created_at") or now),
            "updated_at": now,
        }

        connection.execute(f"DELETE FROM {CLAIMS_TABLE} WHERE claim_id=?", (claim_id,))
        fields = list(claim.keys())
        connection.execute(
            f"INSERT INTO {CLAIMS_TABLE}({','.join(fields)}) VALUES({','.join('?' for _ in fields)})",
            tuple(claim[f] for f in fields),
        )

        connection.execute(f"DELETE FROM {WORKBOOK_TABLE} WHERE claim_id=?", (claim_id,))
        connection.execute(
            f"INSERT INTO {WORKBOOK_TABLE}(claim_id,project_id,filename,batch_id,payload,updated_at) VALUES(?,?,?,?,?,?)",
            (claim_id, pid, filename, batch_id, payload, now),
        )

        if old_row is None or old_batch != batch_id:
            revision_id = uuid.uuid4().hex
            connection.execute(
                f"INSERT INTO {REVISIONS_TABLE}(revision_id,claim_id,project_id,revision_no,filename,batch_id,payload,created_at) VALUES(?,?,?,?,?,?,?,?)",
                (revision_id, claim_id, pid, int(new_revision), filename, batch_id, payload, now),
            )

        connection.execute(f"DELETE FROM {ITEMS_TABLE} WHERE claim_id=?", (claim_id,))
        item_fields = [
            "claim_id","project_id","sheet_name","row_no","seq","boq_item","contract_qty","unit","spec","item_code","brand","origin",
            "material_unit_price","labor_unit_price","contract_amount","material_previous_qty","material_current_qty","material_cumulative_qty",
            "installation_previous_pct","installation_current_pct","installation_cumulative_pct","material_previous_value","material_current_value",
            "material_cumulative_value","installation_previous_value","installation_current_value","installation_cumulative_value",
            "deduction_previous","deduction_current","deduction_cumulative","current_value","cumulative_value","completion_ratio","note","cost_code","system",
        ]
        params = []
        for item in result.get("detail_items") or []:
            row = {**item, "claim_id": claim_id, "project_id": pid}
            params.append(tuple(row.get(field, "") for field in item_fields))
        if params:
            sql = f"INSERT INTO {ITEMS_TABLE}({','.join(item_fields)}) VALUES({','.join('?' for _ in item_fields)})"
            connection.executemany(sql, params)

        _sync_payment_tracking(connection, claim)

    return {
        "claim_id": claim_id,
        "claim_no": claim_no,
        "claim_code": claim_code,
        "revision_no": int(new_revision),
        "inserted_items": len(result.get("detail_items") or []),
        "requested_amount": float(summary.get("requested_amount") or 0),
        "updated_at": now,
    }


def list_ipc_claims(db, project_id: int) -> list[dict[str, Any]]:
    with db.connect() as connection:
        _ensure_tables(connection)
        rows = connection.execute(
            f"SELECT * FROM {CLAIMS_TABLE} WHERE project_id=? ORDER BY claim_no,created_at",
            (int(project_id),),
        ).fetchall()
    return [_rowdict(row) for row in rows]


def get_ipc_claim(db, claim_id: str) -> dict[str, Any] | None:
    with db.connect() as connection:
        _ensure_tables(connection)
        row = connection.execute(f"SELECT * FROM {CLAIMS_TABLE} WHERE claim_id=?", (str(claim_id),)).fetchone()
    return _rowdict(row) if row is not None else None


def load_ipc_workbook(db, claim_id: str) -> dict[str, Any] | None:
    with db.connect() as connection:
        _ensure_tables(connection)
        row = connection.execute(
            f"SELECT filename,batch_id,payload,updated_at FROM {WORKBOOK_TABLE} WHERE claim_id=?",
            (str(claim_id),),
        ).fetchone()
    if row is None:
        return None
    data = _rowdict(row)
    if not data:
        data = {"filename": row[0], "batch_id": row[1], "payload": row[2], "updated_at": row[3]}
    result = _decode_result(str(data.get("payload") or ""))
    result["_persisted"] = True
    result["_persisted_at"] = str(data.get("updated_at") or "")
    return result


def ipc_claim_items(db, claim_id: str) -> list[dict[str, Any]]:
    with db.connect() as connection:
        _ensure_tables(connection)
        rows = connection.execute(
            f"SELECT * FROM {ITEMS_TABLE} WHERE claim_id=? ORDER BY row_no",
            (str(claim_id),),
        ).fetchall()
    return [_rowdict(row) for row in rows]


def ipc_claim_revisions(db, claim_id: str) -> list[dict[str, Any]]:
    with db.connect() as connection:
        _ensure_tables(connection)
        rows = connection.execute(
            f"SELECT revision_id,claim_id,project_id,revision_no,filename,batch_id,created_at FROM {REVISIONS_TABLE} WHERE claim_id=? ORDER BY revision_no DESC",
            (str(claim_id),),
        ).fetchall()
    return [_rowdict(row) for row in rows]


def update_ipc_claim_finance(
    db,
    claim_id: str,
    *,
    approved_amount: float,
    disbursed_amount: float,
    payment_status: str,
    disbursement_date: str = "",
    note: str = "",
) -> None:
    now = _now()
    with db.connect() as connection:
        _ensure_tables(connection)
        connection.execute(
            f"""UPDATE {CLAIMS_TABLE}
                   SET approved_amount=?,disbursed_amount=?,payment_status=?,disbursement_date=?,note=?,updated_at=?
                   WHERE claim_id=?""",
            (
                float(approved_amount or 0), float(disbursed_amount or 0), str(payment_status or ""),
                str(disbursement_date or ""), str(note or ""), now, str(claim_id),
            ),
        )
        row = connection.execute(f"SELECT * FROM {CLAIMS_TABLE} WHERE claim_id=?", (str(claim_id),)).fetchone()
        if row is not None:
            _sync_payment_tracking(connection, _rowdict(row))


def compare_claim_to_boq(db, project_id: int, claim_id: str) -> list[dict[str, Any]]:
    items = ipc_claim_items(db, claim_id)
    with db.connect() as connection:
        boq_rows = connection.execute(
            "SELECT id,task_ref,boq_item,quantity,unit,unit_price,budget_total FROM cost_budgets WHERE project_id=? ORDER BY id",
            (int(project_id),),
        ).fetchall()
    boq = [_rowdict(row) for row in boq_rows]
    lookup: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for row in boq:
        key = (_norm(row.get("boq_item")), _norm(row.get("unit")))
        lookup.setdefault(key, []).append(row)

    out: list[dict[str, Any]] = []
    for item in items:
        key = (_norm(item.get("boq_item")), _norm(item.get("unit")))
        matches = lookup.get(key) or lookup.get((key[0], "")) or []
        boq_row = matches[0] if matches else {}
        boq_qty = float(boq_row.get("quantity") or 0)
        cumulative_qty = float(item.get("material_cumulative_qty") or 0)
        qty_ratio = cumulative_qty / boq_qty if boq_qty else 0
        out.append(
            {
                "Khớp BOQ": "Có" if boq_row else "Chưa",
                "Hạng mục": item.get("boq_item") or "",
                "ĐVT": item.get("unit") or "",
                "KL BOQ": boq_qty,
                "KL HĐ trong Claim": float(item.get("contract_qty") or 0),
                "KL vật tư lũy kế": cumulative_qty,
                "% KL vật tư/BOQ": qty_ratio * 100 if boq_qty else 0,
                "Giá trị BOQ": float(boq_row.get("budget_total") or 0),
                "Giá trị kỳ này": float(item.get("current_value") or 0),
                "Giá trị lũy kế": float(item.get("cumulative_value") or 0),
            }
        )
    return out


def _parse_ui_date(value: Any):
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except Exception:
            pass
    return date.today()


def render_ipc_claim_ui(db, project_id: int, *, can_update: bool = True) -> None:
    import pandas as pd
    import streamlit as st

    pid = int(project_id)
    claims = list_ipc_claims(db, pid)

    st.markdown("### 💰 Thanh toán & Giải ngân theo IPC (Claim)")
    st.caption(
        "Mỗi lần thanh toán là một Claim riêng. Mỗi Claim lưu workbook Excel nhiều sheet ở cấp dự án, "
        "không phụ thuộc tài khoản đăng nhập. Claim cũ không bị ghi đè; file mới của cùng Claim được lưu thành revision."
    )

    with st.expander("➕ Tạo Claim mới / cập nhật revision từ Excel", expanded=not bool(claims)):
        upload = st.file_uploader(
            "Chọn file IPC/Claim (.xlsx / .xlsm)",
            type=["xlsx", "xlsm"],
            key=f"ipc_claim_upload_{pid}",
        )
        parsed = None
        if upload is not None:
            try:
                parsed = parse_ipc_workbook(upload.getvalue(), upload.name)
            except Exception as exc:
                st.error(f"Không thể đọc file IPC: {exc}")
            if parsed is not None:
                meta = parsed.get("metadata") or {}
                summary = parsed.get("summary") or {}
                st.success(
                    f"Nhận diện {parsed.get('claim_code')} · {len(parsed.get('workbook_sheet_names') or []):,} sheet hiển thị · "
                    f"{int(parsed.get('detail_line_count') or 0):,} dòng GTHT chi tiết."
                )
                for warning in parsed.get("warnings") or []:
                    st.info(warning)
                a, b, c = st.columns(3)
                a.metric("Giá trị hợp đồng", f"{format_table_number(summary.get('contract_value', 0))} VND")
                b.metric("Giá trị Claim kỳ này", f"{format_table_number(summary.get('requested_amount', 0))} VND")
                c.metric("Lũy kế nghiệm thu", f"{format_table_number(summary.get('cumulative_completed', 0))} VND")
                st.caption(
                    f"Nhà thầu: {meta.get('contractor','')} · Hợp đồng: {meta.get('contract_no','')} · "
                    f"Kỳ {meta.get('from_date','')} → {meta.get('to_date','')}"
                )
                if st.button(
                    f"💾 Lưu {parsed.get('claim_code')} vào dự án",
                    type="primary",
                    disabled=not bool(can_update),
                    key=f"ipc_claim_save_{pid}_{parsed.get('batch_id','')}",
                    width="stretch",
                ):
                    try:
                        stats = save_ipc_claim(db, pid, parsed)
                        st.success(
                            f"Đã lưu {stats['claim_code']} · revision {stats['revision_no']} · "
                            f"{stats['inserted_items']:,} dòng chi tiết · đề nghị {format_table_number(stats['requested_amount'])} VND."
                        )
                        st.rerun()
                    except Exception as exc:
                        st.error(f"Không thể lưu Claim: {exc}")

    claims = list_ipc_claims(db, pid)
    labels = ["📊 Tổng hợp"] + [f"Claim {str(c.get('claim_no') or '')}" for c in claims]
    tabs = st.tabs(labels)

    with tabs[0]:
        contract_value = max([float(c.get("contract_value") or 0) for c in claims] + [0])
        requested = sum(float(c.get("requested_amount") or 0) for c in claims)
        approved = sum(float(c.get("approved_amount") or 0) for c in claims)
        disbursed = sum(float(c.get("disbursed_amount") or 0) for c in claims)
        remaining = max(0.0, contract_value - approved) if contract_value else 0.0
        m1, m2, m3, m4, m5 = st.columns(5)
        m1.metric("Số Claim", f"{len(claims):,}")
        m2.metric("Tổng đề nghị", f"{format_table_number(requested)} VND")
        m3.metric("Được duyệt", f"{format_table_number(approved)} VND")
        m4.metric("Đã giải ngân", f"{format_table_number(disbursed)} VND")
        m5.metric("Còn lại HĐ", f"{format_table_number(remaining)} VND")
        if claims:
            summary_df = pd.DataFrame(
                [
                    {
                        "Claim": c.get("claim_code"),
                        "Nhà thầu": c.get("contractor"),
                        "Từ ngày": c.get("from_date"),
                        "Đến ngày": c.get("to_date"),
                        "Đề nghị (VND)": float(c.get("requested_amount") or 0),
                        "Duyệt (VND)": float(c.get("approved_amount") or 0),
                        "Giải ngân (VND)": float(c.get("disbursed_amount") or 0),
                        "Revision": int(c.get("latest_revision") or 0),
                        "Trạng thái": c.get("payment_status"),
                    }
                    for c in claims
                ]
            )
            for col in ("Đề nghị (VND)", "Duyệt (VND)", "Giải ngân (VND)"):
                summary_df[col] = summary_df[col].map(format_table_number)
            st.dataframe(summary_df, hide_index=True, width="stretch")
        else:
            st.info("Chưa có Claim. Hãy tải file IPC Excel đầu tiên ở phần phía trên.")

    for claim, tab in zip(claims, tabs[1:]):
        claim_id = str(claim.get("claim_id") or "")
        with tab:
            workbook = load_ipc_workbook(db, claim_id)
            sub = st.tabs(["Tổng quan", "Excel", "Chi tiết nghiệm thu", "Đối chiếu BOQ", "Giải ngân", "Lịch sử"])

            with sub[0]:
                c1, c2, c3, c4 = st.columns(4)
                c1.metric("Đề nghị kỳ này", f"{format_table_number(claim.get('requested_amount', 0))} VND")
                c2.metric("Được duyệt", f"{format_table_number(claim.get('approved_amount', 0))} VND")
                c3.metric("Đã giải ngân", f"{format_table_number(claim.get('disbursed_amount', 0))} VND")
                c4.metric("Lũy kế nghiệm thu", f"{format_table_number(claim.get('certified_cumulative', 0))} VND")
                st.markdown(
                    f"**Mã Claim:** {claim.get('claim_code','')}  \n"
                    f"**Nhà thầu:** {claim.get('contractor','')}  \n"
                    f"**Hợp đồng:** {claim.get('contract_no','')}  \n"
                    f"**Kỳ:** {claim.get('from_date','')} → {claim.get('to_date','')}  \n"
                    f"**File:** {claim.get('filename','')} · Revision {int(claim.get('latest_revision') or 0)}  \n"
                    f"**Trạng thái:** {claim.get('payment_status','')}"
                )

            with sub[1]:
                if workbook is None:
                    st.warning("Claim này chưa có snapshot workbook Excel.")
                else:
                    names = list(workbook.get("workbook_sheet_names") or [])
                    data = workbook.get("workbook_sheets") or {}
                    if names:
                        excel_tabs = st.tabs(names)
                        for name, excel_tab in zip(names, excel_tabs):
                            with excel_tab:
                                snap = data.get(name) or {}
                                df = pd.DataFrame(snap.get("rows") or [], columns=snap.get("columns") or None)
                                if not df.empty:
                                    df = df.map(format_table_number)
                                st.caption(f"{int(snap.get('row_count') or 0):,} dòng × {int(snap.get('col_count') or 0):,} cột")
                                if snap.get("truncated"):
                                    st.warning("Sheet quá lớn; phần xem trước đã được giới hạn để bảo vệ bộ nhớ.")
                                st.dataframe(df, hide_index=True, width="stretch", height=560)

            with sub[2]:
                items = ipc_claim_items(db, claim_id)
                if items:
                    df = pd.DataFrame(items)
                    show_cols = [
                        "row_no","boq_item","contract_qty","unit","material_current_qty","material_cumulative_qty",
                        "installation_current_pct","installation_cumulative_pct","current_value","cumulative_value","system","note",
                    ]
                    df = df[[c for c in show_cols if c in df.columns]].rename(columns={
                        "row_no":"Dòng","boq_item":"Tên công tác / Diễn giải","contract_qty":"KL hợp đồng","unit":"ĐVT",
                        "material_current_qty":"KL vật tư kỳ này","material_cumulative_qty":"KL vật tư lũy kế",
                        "installation_current_pct":"Lắp đặt kỳ này (%)","installation_cumulative_pct":"Lắp đặt lũy kế (%)",
                        "current_value":"Giá trị kỳ này (VND)","cumulative_value":"Giá trị lũy kế (VND)","system":"Hệ thống","note":"Ghi chú",
                    })
                    for col in df.columns:
                        if col in {"Dòng","KL hợp đồng","KL vật tư kỳ này","KL vật tư lũy kế","Lắp đặt kỳ này (%)","Lắp đặt lũy kế (%)","Giá trị kỳ này (VND)","Giá trị lũy kế (VND)"}:
                            df[col] = df[col].map(format_table_number)
                    st.caption(f"{len(items):,} dòng GTHT chi tiết đã lưu trong Claim.")
                    st.dataframe(df, hide_index=True, width="stretch", height=560)
                else:
                    st.info("Claim này chưa có dữ liệu GTHT chuẩn hóa.")

            with sub[3]:
                rows = compare_claim_to_boq(db, pid, claim_id)
                if rows:
                    matched = sum(1 for row in rows if row.get("Khớp BOQ") == "Có")
                    st.caption(f"Đối chiếu tên hạng mục + ĐVT: khớp {matched:,}/{len(rows):,} dòng Claim với BOQ dự án.")
                    df = pd.DataFrame(rows)
                    for col in ["KL BOQ","KL HĐ trong Claim","KL vật tư lũy kế","% KL vật tư/BOQ","Giá trị BOQ","Giá trị kỳ này","Giá trị lũy kế"]:
                        if col in df.columns:
                            df[col] = df[col].map(format_table_number)
                    st.dataframe(df, hide_index=True, width="stretch", height=560)
                else:
                    st.info("Chưa có dòng Claim để đối chiếu BOQ.")

            with sub[4]:
                statuses = ["Nháp", "Đã nộp", "Đang kiểm tra", "Yêu cầu chỉnh sửa", "Đã duyệt", "Chờ giải ngân", "Đã giải ngân", "Từ chối"]
                current_status = str(claim.get("payment_status") or "Nháp")
                status_index = statuses.index(current_status) if current_status in statuses else 0
                with st.form(f"ipc_finance_{claim_id}"):
                    approved = st.number_input(
                        "Giá trị được duyệt (VND)", min_value=0.0,
                        value=float(claim.get("approved_amount") or 0), step=1000000.0,
                    )
                    disbursed = st.number_input(
                        "Giá trị đã giải ngân (VND)", min_value=0.0,
                        value=float(claim.get("disbursed_amount") or 0), step=1000000.0,
                    )
                    status = st.selectbox("Trạng thái Claim", statuses, index=status_index)
                    disbursement_date = st.date_input(
                        "Ngày giải ngân", value=_parse_ui_date(claim.get("disbursement_date")),
                    )
                    note = st.text_area("Ghi chú", value=str(claim.get("note") or ""))
                    submitted = st.form_submit_button("💾 Cập nhật duyệt / giải ngân", disabled=not bool(can_update), use_container_width=True)
                if submitted:
                    try:
                        update_ipc_claim_finance(
                            db, claim_id,
                            approved_amount=approved,
                            disbursed_amount=disbursed,
                            payment_status=status,
                            disbursement_date=disbursement_date.strftime("%Y-%m-%d"),
                            note=note,
                        )
                        st.success("Đã cập nhật Claim và đồng bộ bảng thanh toán của dự án.")
                        st.rerun()
                    except Exception as exc:
                        st.error(f"Không thể cập nhật Claim: {exc}")

            with sub[5]:
                revisions = ipc_claim_revisions(db, claim_id)
                if revisions:
                    rev_df = pd.DataFrame(revisions).rename(columns={
                        "revision_no":"Revision","filename":"File Excel","batch_id":"Batch","created_at":"Ngày cập nhật",
                    })
                    cols = [c for c in ["Revision","File Excel","Batch","Ngày cập nhật"] if c in rev_df.columns]
                    st.dataframe(rev_df[cols], hide_index=True, width="stretch")
                else:
                    st.info("Chưa có lịch sử revision.")
