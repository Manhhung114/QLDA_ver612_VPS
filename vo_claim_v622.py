from __future__ import annotations

import base64
import gzip
import hashlib
import io
import json
import math
import re
import unicodedata
import uuid
from datetime import date, datetime
from numbers import Real
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

PATCH_VERSION = "V6.22 VO EXCEL V2 SIGNED"
PAYLOAD_PREFIX = "gz1:"
MAX_WORKBOOK_BYTES = 120 * 1024 * 1024
MAX_PREVIEW_ROWS = 1800
MAX_PREVIEW_COLS = 36

VO_TABLE = "variation_orders"
ITEMS_TABLE = "variation_order_items"
WORKBOOK_TABLE = "variation_order_workbooks"
REVISIONS_TABLE = "variation_order_revisions"


class VOWorkbookError(ValueError):
    pass


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _rowdict(row: Any) -> dict[str, Any]:
    if row is None:
        return {}
    if isinstance(row, dict):
        return dict(row)
    try:
        return {str(k): row[k] for k in row.keys()}
    except Exception:
        try:
            return dict(row)
        except Exception:
            return {}


def _norm(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.lower().replace("đ", "d")
    text = re.sub(r"[^a-z0-9%]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _to_number(value: Any) -> float | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    if isinstance(value, Real):
        n = float(value)
        return n if math.isfinite(n) else None
    text = str(value).strip()
    if not text or text.startswith("=") or text.startswith("#"):
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    text = re.sub(r"[^0-9,\.\-]", "", text)
    if not text or text in {"-", ".", ","}:
        return None
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        parts = text.split(",")
        if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3):
            text = "".join(parts)
        else:
            text = text.replace(",", ".")
    try:
        n = float(text)
        if negative:
            n = -abs(n)
        return n if math.isfinite(n) else None
    except Exception:
        return None


def _money(value: Any) -> str:
    try:
        n = float(value or 0)
        if abs(n - round(n)) < 1e-9:
            return f"{n:,.0f}"
        return f"{n:,.2f}"
    except Exception:
        return "0"


def _table_value(value: Any) -> Any:
    if isinstance(value, bool) or not isinstance(value, Real):
        return value
    n = float(value)
    if not math.isfinite(n):
        return value
    if abs(n - round(n)) < 1e-9:
        return f"{n:,.0f}"
    return f"{n:,.4f}".rstrip("0").rstrip(".")


def _date_text(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y.%m.%d"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except Exception:
            pass
    return ""


def _date_from_filename(filename: str) -> str:
    text = str(filename or "")
    for pat in (
        r"(?<!\d)(20\d{2})[._-](\d{1,2})[._-](\d{1,2})(?!\d)",
        r"(?<!\d)(20\d{2})(\d{2})(\d{2})(?!\d)",
    ):
        m = re.search(pat, text)
        if not m:
            continue
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3))).strftime("%Y-%m-%d")
        except Exception:
            pass
    return ""


def _vo_number_from_text(value: Any) -> int | None:
    text = str(value or "")
    m = re.search(r"\bVO\s*[-#:/]?\s*0*([0-9]{1,4})\b", text, re.I)
    return int(m.group(1)) if m else None


def _revision_from_text(value: Any) -> str:
    m = re.search(r"\bR\s*([0-9]+)\b", str(value or ""), re.I)
    return f"R{int(m.group(1))}" if m else ""


def _sheet_snapshot(ws) -> dict[str, Any]:
    max_row = min(int(ws.max_row or 1), MAX_PREVIEW_ROWS)
    max_col = min(int(ws.max_column or 1), MAX_PREVIEW_COLS)
    rows: list[list[Any]] = []
    for row_no, values in enumerate(
        ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col, values_only=True),
        start=1,
    ):
        shown = []
        for value in values:
            if isinstance(value, datetime):
                value = value.strftime("%d/%m/%Y %H:%M:%S")
            elif isinstance(value, date):
                value = value.strftime("%d/%m/%Y")
            shown.append(value)
        rows.append([row_no] + shown)
    return {
        "sheet": ws.title,
        "columns": ["Dòng"] + [get_column_letter(i) for i in range(1, max_col + 1)],
        "rows": rows,
        "row_count": int(ws.max_row or 0),
        "col_count": int(ws.max_column or 0),
        "truncated": bool((ws.max_row or 0) > max_row or (ws.max_column or 0) > max_col),
    }


def _find_summary_sheet(workbook):
    best = None
    best_score = -1
    for ws in workbook.worksheets:
        score = 0
        if "tong hop" in _norm(ws.title):
            score += 5
        cells = []
        for row in ws.iter_rows(min_row=1, max_row=min(int(ws.max_row or 1), 25), min_col=1, max_col=min(int(ws.max_column or 1), 8), values_only=True):
            cells.extend(_norm(v) for v in row if v not in (None, ""))
        blob = " ".join(cells)
        if "tong hop gia phat sinh" in blob:
            score += 8
        if "vo" in blob:
            score += 3
        if "tong cong" in blob and "vat" in blob:
            score += 3
        if score > best_score:
            best, best_score = ws, score
    return best if best_score >= 5 else None


def _summary_metadata(ws, filename: str) -> dict[str, Any]:
    blob_parts = [str(filename or "")]
    project = ""
    package = ""
    revision = ""
    if ws is not None:
        for row in ws.iter_rows(min_row=1, max_row=min(int(ws.max_row or 1), 30), min_col=1, max_col=min(int(ws.max_column or 1), 10), values_only=True):
            for value in row:
                if value not in (None, ""):
                    blob_parts.append(str(value))
                    n = _norm(value)
                    if n.startswith("du an") and not project:
                        project = re.sub(r"^\s*DỰ\s*ÁN\s*:\s*", "", str(value), flags=re.I).strip()
                    if n.startswith("goi thau") and not package:
                        package = re.sub(r"^\s*GÓI\s*THẦU\s*:\s*", "", str(value), flags=re.I).strip()
                    if "lan sua doi" in n and not revision:
                        revision = _revision_from_text(value)
    blob = " | ".join(blob_parts)
    file_no = _vo_number_from_text(filename)
    content_no = _vo_number_from_text(blob)
    vo_no = file_no or content_no
    if vo_no is None:
        raise VOWorkbookError("Không xác định được số VO trong tên file hoặc sheet tổng hợp.")
    warnings: list[str] = []
    if file_no and content_no and file_no != content_no:
        warnings.append(
            f"Tên file ghi VO-{file_no:02d} nhưng nội dung workbook ghi VO-{content_no:02d}; "
            f"hệ thống ưu tiên tên file để tránh cập nhật nhầm VO."
        )
    return {
        "vo_no": int(vo_no),
        "vo_code": f"VO-{int(vo_no):02d}",
        "revision_label": _revision_from_text(filename) or revision or "R0",
        "vo_date": _date_from_filename(filename),
        "project": project,
        "package": package,
        "warnings": warnings,
    }


def _summary_values(ws) -> tuple[dict[str, float], list[dict[str, Any]], int]:
    if ws is None:
        return {}, [], 0
    subtotal = vat = total = None
    summary_lines: list[dict[str, Any]] = []
    subtotal_row = 0
    for r in range(1, min(int(ws.max_row or 1), 120) + 1):
        vals = [ws.cell(r, c).value for c in range(1, min(int(ws.max_column or 1), 12) + 1)]
        text_cells = [(c + 1, str(v).strip(), _norm(v)) for c, v in enumerate(vals) if isinstance(v, str) and str(v).strip()]
        nums = [(c + 1, _to_number(v)) for c, v in enumerate(vals) if _to_number(v) is not None]
        if text_cells and nums:
            label = max(text_cells, key=lambda x: len(x[1]))[1]
            amount = nums[-1][1]
            if amount is not None:
                summary_lines.append({"row_no": r, "description": label, "amount": float(amount)})
        row_norm = " ".join(x[2] for x in text_cells)
        amount = nums[-1][1] if nums else None
        if amount is None:
            continue
        if "tong cong" in row_norm and ("chua bao gom vat" in row_norm or "truoc thue" in row_norm):
            subtotal = float(amount)
            subtotal_row = r
        elif "tong cong" in row_norm and ("sau thue" in row_norm or "sau vat" in row_norm):
            total = float(amount)
        elif ("thue vat" in row_norm or ("vat" in row_norm and "10" in row_norm)) and "tong cong" not in row_norm:
            vat = float(amount)
    if subtotal is None:
        for row in reversed(summary_lines):
            n = _norm(row["description"])
            if "tong cong" in n and "vat" not in n:
                subtotal = float(row["amount"])
                subtotal_row = int(row["row_no"])
                break
    if total is None and subtotal is not None and vat is not None:
        total = subtotal + vat
    return {
        "subtotal_before_vat": float(subtotal or 0),
        "vat_amount": float(vat or 0),
        "total_after_vat": float(total if total is not None else subtotal or 0),
    }, summary_lines, subtotal_row


def _formula_source_sheets(formula_ws, subtotal_row: int, valid_names: set[str]) -> list[str]:
    if formula_ws is None:
        return []
    refs: list[str] = []
    end = max(1, subtotal_row - 1) if subtotal_row else min(int(formula_ws.max_row or 1), 80)
    for row in formula_ws.iter_rows(min_row=1, max_row=end, min_col=1, max_col=min(int(formula_ws.max_column or 1), 20), values_only=True):
        for value in row:
            if not isinstance(value, str) or not value.startswith("="):
                continue
            for name in re.findall(r"'([^']+)'!", value):
                if name in valid_names and name not in refs:
                    refs.append(name)
            for name in re.findall(r"(?<!')([A-Za-z0-9_. -]+)!", value):
                name = name.strip(" =()+-*/")
                if name in valid_names and name not in refs:
                    refs.append(name)
    return refs


def _fallback_source_sheets(workbook, summary_title: str) -> list[str]:
    out = []
    for ws in workbook.worksheets:
        if ws.title == summary_title or getattr(ws, "sheet_state", "visible") != "visible":
            continue
        nname = _norm(ws.title)
        if any(x in nname for x in ("goc", "pa1", "dgkl", "kangatang")):
            continue
        blob = []
        for row in ws.iter_rows(min_row=1, max_row=min(int(ws.max_row or 1), 5), min_col=1, max_col=min(int(ws.max_column or 1), 20), values_only=True):
            blob.extend(_norm(v) for v in row if v not in (None, ""))
        text = " ".join(blob)
        if "phat sinh" in text and ("tang giam" in text or "gia tri chi tiet" in text):
            out.append(ws.title)
    return out


def _header_map(ws) -> dict[int, str]:
    max_col = min(int(ws.max_column or 1), 40)
    header_row = None
    for r in range(1, min(int(ws.max_row or 1), 12) + 1):
        row_text = " ".join(
            _norm(ws.cell(r, c).value)
            for c in range(1, max_col + 1)
            if ws.cell(r, c).value not in (None, "")
        )
        if "noi dung cong viec" in row_text or ("mo ta" in row_text and "don vi" in row_text):
            header_row = r
            break
    if header_row is None:
        header_row = 2
    end_row = min(int(ws.max_row or header_row), header_row + 1)
    out = {}
    for c in range(1, max_col + 1):
        text = " ".join(
            _norm(ws.cell(r, c).value)
            for r in range(header_row, end_row + 1)
            if ws.cell(r, c).value not in (None, "")
        )
        out[c] = text
    return out


def _pick_col(headers: dict[int, str], patterns: tuple[str, ...], excludes: tuple[str, ...] = ()) -> int | None:
    candidates = []
    for c, text in headers.items():
        if any(p in text for p in patterns) and not any(x in text for x in excludes):
            candidates.append((len(text), c))
    return sorted(candidates)[0][1] if candidates else None


def _is_rollup_formula(value: Any) -> bool:
    if not isinstance(value, str) or not value.startswith("="):
        return False
    upper = value.upper()
    if "SUBTOTAL" in upper:
        return True
    for _, r1, _, r2 in re.findall(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", upper):
        if r1 != r2:
            return True
    return False


def _parse_detail_sheet(value_ws, formula_ws) -> list[dict[str, Any]]:
    headers = _header_map(value_ws)
    desc_col = _pick_col(headers, ("noi dung cong viec", "mo ta", "noi dung"))
    if not desc_col:
        return []
    seq_col = _pick_col(headers, ("stt", " tt "))
    unit_col = _pick_col(headers, ("don vi",))
    variation_col = _pick_col(headers, ("phat sinh tang giam",))
    increase_col = _pick_col(headers, ("phat sinh tang",), ("phat sinh tang giam",))
    decrease_col = _pick_col(headers, ("phat sinh giam",))
    contract_col = _pick_col(headers, ("theo hop dong",))
    actual_col = _pick_col(headers, ("thuc te thi cong",))
    mat_col = _pick_col(headers, ("don gia vat tu", "don gia vat lieu"))
    labor_col = _pick_col(headers, ("nhan cong",))
    amount_col = _pick_col(headers, ("thanh tien",))
    spec_col = _pick_col(headers, ("quy cach",), ("dieu chinh",))
    code_col = _pick_col(headers, ("ma hieu",), ("dieu chinh",))
    brand_col = _pick_col(headers, ("thuong hieu",))
    origin_col = _pick_col(headers, ("xuat xu",))
    note_col = _pick_col(headers, ("ghi chu",))

    items: list[dict[str, Any]] = []
    for r in range(4, int(value_ws.max_row or 4) + 1):
        description = str(value_ws.cell(r, desc_col).value or "").strip()
        if not description:
            continue
        nd = _norm(description)
        if any(x in nd for x in (
            "cong gia tri truoc thue", "tong cong", "thue vat", "ban qlda", "tu van giam sat",
            "tong thau", "nha thau thi cong truc tiep",
        )):
            continue
        if amount_col and formula_ws is not None and _is_rollup_formula(formula_ws.cell(r, amount_col).value):
            continue

        contract_qty = _to_number(value_ws.cell(r, contract_col).value) if contract_col else None
        actual_qty = _to_number(value_ws.cell(r, actual_col).value) if actual_col else None
        variation_qty = _to_number(value_ws.cell(r, variation_col).value) if variation_col else None
        increase_qty = _to_number(value_ws.cell(r, increase_col).value) if increase_col else None
        decrease_qty = _to_number(value_ws.cell(r, decrease_col).value) if decrease_col else None

        if decrease_qty is not None and decrease_qty > 0:
            decrease_qty = -abs(decrease_qty)
        if variation_qty is None and contract_qty is not None and actual_qty is not None:
            variation_qty = actual_qty - contract_qty
        if variation_qty is None and (increase_qty is not None or decrease_qty is not None):
            variation_qty = float(increase_qty or 0) + float(decrease_qty or 0)

        material_price = _to_number(value_ws.cell(r, mat_col).value) if mat_col else None
        labor_price = _to_number(value_ws.cell(r, labor_col).value) if labor_col else None
        direct_amount = _to_number(value_ws.cell(r, amount_col).value) if amount_col else None

        amount = float(direct_amount or 0)
        qty = float(variation_qty or 0)
        if abs(qty) < 1e-12 and abs(amount) < 0.5:
            continue

        kind = "Tăng" if amount > 0.5 or (abs(amount) <= 0.5 and qty > 0) else "Giảm" if amount < -0.5 or qty < 0 else "Chưa định giá"
        items.append({
            "sheet_name": value_ws.title,
            "row_no": r,
            "seq": str(value_ws.cell(r, seq_col).value or "").strip() if seq_col else "",
            "description": description,
            "unit": str(value_ws.cell(r, unit_col).value or "").strip() if unit_col else "",
            "contract_qty": float(contract_qty or 0),
            "actual_qty": float(actual_qty or 0),
            "increase_qty": float(increase_qty or 0),
            "decrease_qty": float(decrease_qty or 0),
            "variation_qty": qty,
            "spec": str(value_ws.cell(r, spec_col).value or "").strip() if spec_col else "",
            "item_code": str(value_ws.cell(r, code_col).value or "").strip() if code_col else "",
            "brand": str(value_ws.cell(r, brand_col).value or "").strip() if brand_col else "",
            "origin": str(value_ws.cell(r, origin_col).value or "").strip() if origin_col else "",
            "material_unit_price": float(material_price or 0),
            "labor_unit_price": float(labor_price or 0),
            "unit_price_total": float((material_price or 0) + (labor_price or 0)),
            "variation_amount": amount,
            "variation_kind": kind,
            "note": str(value_ws.cell(r, note_col).value or "").strip() if note_col else "",
        })
    return items


def parse_vo_workbook(data: bytes, filename: str = "VO.xlsx") -> dict[str, Any]:
    raw = bytes(data or b"")
    if not raw:
        raise VOWorkbookError("File VO đang trống.")
    if len(raw) > MAX_WORKBOOK_BYTES:
        raise VOWorkbookError("File VO lớn hơn 120 MB; hãy tối ưu hoặc tách file trước khi nhập.")
    suffix = Path(str(filename or "")).suffix.lower()
    if suffix and suffix not in {".xlsx", ".xlsm"}:
        raise VOWorkbookError("Chỉ hỗ trợ file VO Excel .xlsx hoặc .xlsm.")

    try:
        wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        wbf = load_workbook(io.BytesIO(raw), data_only=False, read_only=True)
    except Exception as exc:
        raise VOWorkbookError(f"Không đọc được workbook VO: {exc}") from exc

    try:
        summary_ws = _find_summary_sheet(wb)
        if summary_ws is None:
            raise VOWorkbookError("Không tìm thấy sheet Tổng hợp VO.")
        formula_summary = wbf[summary_ws.title] if summary_ws.title in wbf.sheetnames else None
        metadata = _summary_metadata(summary_ws, filename)
        summary, summary_lines, subtotal_row = _summary_values(summary_ws)
        valid_names = set(wb.sheetnames)
        source_sheets = _formula_source_sheets(formula_summary, subtotal_row, valid_names)
        if not source_sheets:
            source_sheets = _fallback_source_sheets(wb, summary_ws.title)

        detail_items: list[dict[str, Any]] = []
        for name in source_sheets:
            if name not in wb.sheetnames:
                continue
            detail_items.extend(_parse_detail_sheet(wb[name], wbf[name] if name in wbf.sheetnames else None))

        visible = [ws for ws in wb.worksheets if getattr(ws, "sheet_state", "visible") == "visible"]
        snapshots = {ws.title: _sheet_snapshot(ws) for ws in visible}
    finally:
        wb.close()
        wbf.close()

    priced_items = [x for x in detail_items if abs(float(x.get("variation_amount") or 0)) >= 0.5]
    increase_amount = sum(max(0.0, float(x.get("variation_amount") or 0)) for x in priced_items)
    decrease_amount = sum(min(0.0, float(x.get("variation_amount") or 0)) for x in priced_items)
    detail_net = increase_amount + decrease_amount
    subtotal = float(summary.get("subtotal_before_vat") or 0)
    discrepancy = detail_net - subtotal

    warnings = list(metadata.pop("warnings", []))
    if not source_sheets:
        warnings.append("Không nhận được các sheet chi tiết từ công thức Tổng hợp; parser dùng dò nội dung.")
    if abs(discrepancy) > max(1.0, abs(subtotal) * 0.0001):
        warnings.append(
            f"Tổng dòng chi tiết có giá {_money(detail_net)} VND lệch {_money(discrepancy)} VND "
            f"so với Tổng hợp {_money(subtotal)} VND. Không tự sửa số liệu; cần kiểm tra workbook."
        )

    confidence = 55
    confidence += 15 if metadata.get("vo_code") else 0
    confidence += 10 if subtotal_row else 0
    confidence += 10 if source_sheets else 0
    confidence += 10 if abs(discrepancy) <= max(1.0, abs(subtotal) * 0.0001) else 0
    confidence = min(100, confidence)

    batch_id = hashlib.sha256(raw).hexdigest()
    return {
        "schema": "qlda_vo_excel_v2",
        "filename": str(filename or "VO.xlsx"),
        "batch_id": batch_id,
        "metadata": metadata,
        "vo_no": int(metadata["vo_no"]),
        "vo_code": str(metadata["vo_code"]),
        "summary": {
            **summary,
            "increase_amount": float(increase_amount),
            "decrease_amount": float(decrease_amount),
            "detail_net_amount": float(detail_net),
            "detail_discrepancy": float(discrepancy),
        },
        "summary_lines": summary_lines,
        "source_sheets": source_sheets,
        "detail_items": detail_items,
        "workbook": snapshots,
        "workbook_sheet_names": [ws.title for ws in visible],
        "confidence_pct": int(confidence),
        "warnings": warnings,
    }


def _ensure_tables(connection) -> None:
    connection.execute(f"""
        CREATE TABLE IF NOT EXISTS {VO_TABLE}(
            vo_id TEXT PRIMARY KEY,
            project_id INTEGER NOT NULL,
            vo_no INTEGER NOT NULL,
            vo_code TEXT NOT NULL,
            filename TEXT NOT NULL,
            batch_id TEXT NOT NULL,
            revision_label TEXT DEFAULT 'R0',
            vo_date TEXT DEFAULT '',
            project_name TEXT DEFAULT '',
            package_name TEXT DEFAULT '',
            subtotal_before_vat REAL DEFAULT 0,
            vat_amount REAL DEFAULT 0,
            total_after_vat REAL DEFAULT 0,
            increase_amount REAL DEFAULT 0,
            decrease_amount REAL DEFAULT 0,
            proposed_amount REAL DEFAULT 0,
            approved_amount REAL DEFAULT 0,
            funding_source TEXT DEFAULT '',
            status TEXT DEFAULT 'Dự thảo',
            latest_revision INTEGER DEFAULT 0,
            note TEXT DEFAULT '',
            created_at TEXT DEFAULT '',
            updated_at TEXT DEFAULT '',
            UNIQUE(project_id, vo_code)
        )
    """)
    connection.execute(f"""
        CREATE TABLE IF NOT EXISTS {ITEMS_TABLE}(
            vo_id TEXT NOT NULL,
            project_id INTEGER NOT NULL,
            sheet_name TEXT NOT NULL,
            row_no INTEGER NOT NULL,
            seq TEXT DEFAULT '',
            description TEXT NOT NULL,
            unit TEXT DEFAULT '',
            contract_qty REAL DEFAULT 0,
            actual_qty REAL DEFAULT 0,
            increase_qty REAL DEFAULT 0,
            decrease_qty REAL DEFAULT 0,
            variation_qty REAL DEFAULT 0,
            spec TEXT DEFAULT '',
            item_code TEXT DEFAULT '',
            brand TEXT DEFAULT '',
            origin TEXT DEFAULT '',
            material_unit_price REAL DEFAULT 0,
            labor_unit_price REAL DEFAULT 0,
            unit_price_total REAL DEFAULT 0,
            variation_amount REAL DEFAULT 0,
            variation_kind TEXT DEFAULT '',
            note TEXT DEFAULT '',
            PRIMARY KEY(vo_id, sheet_name, row_no)
        )
    """)
    connection.execute(f"""
        CREATE TABLE IF NOT EXISTS {WORKBOOK_TABLE}(
            vo_id TEXT PRIMARY KEY,
            project_id INTEGER NOT NULL,
            filename TEXT NOT NULL,
            batch_id TEXT NOT NULL,
            payload TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )
    """)
    connection.execute(f"""
        CREATE TABLE IF NOT EXISTS {REVISIONS_TABLE}(
            revision_id TEXT PRIMARY KEY,
            vo_id TEXT NOT NULL,
            project_id INTEGER NOT NULL,
            revision_no INTEGER NOT NULL,
            revision_label TEXT DEFAULT '',
            filename TEXT NOT NULL,
            batch_id TEXT NOT NULL,
            payload TEXT NOT NULL,
            created_at TEXT NOT NULL,
            UNIQUE(vo_id, revision_no)
        )
    """)
    connection.execute(f"CREATE INDEX IF NOT EXISTS idx_vo_project ON {VO_TABLE}(project_id, vo_no)")
    connection.execute(f"CREATE INDEX IF NOT EXISTS idx_vo_items_project ON {ITEMS_TABLE}(project_id, vo_id)")
    connection.execute(f"CREATE INDEX IF NOT EXISTS idx_vo_revision ON {REVISIONS_TABLE}(vo_id, revision_no)")


def _encode_result(result: dict[str, Any]) -> str:
    raw = json.dumps(result, ensure_ascii=False, separators=(",", ":"), allow_nan=False).encode("utf-8")
    return PAYLOAD_PREFIX + base64.b64encode(gzip.compress(raw, compresslevel=6)).decode("ascii")


def _decode_result(payload: str) -> dict[str, Any]:
    text = str(payload or "")
    if not text.startswith(PAYLOAD_PREFIX):
        raise ValueError("Dữ liệu workbook VO đã lưu không hợp lệ.")
    packed = base64.b64decode(text[len(PAYLOAD_PREFIX):].encode("ascii"))
    data = json.loads(gzip.decompress(packed).decode("utf-8"))
    if not isinstance(data, dict):
        raise ValueError("Dữ liệu VO đã lưu không hợp lệ.")
    return data


def _replace_items(connection, project_id: int, vo_id: str, result: dict[str, Any]) -> int:
    connection.execute(f"DELETE FROM {ITEMS_TABLE} WHERE vo_id=?", (str(vo_id),))
    fields = [
        "vo_id", "project_id", "sheet_name", "row_no", "seq", "description", "unit",
        "contract_qty", "actual_qty", "increase_qty", "decrease_qty", "variation_qty",
        "spec", "item_code", "brand", "origin", "material_unit_price", "labor_unit_price",
        "unit_price_total", "variation_amount", "variation_kind", "note",
    ]
    params = []
    for item in result.get("detail_items") or []:
        row = {**item, "vo_id": str(vo_id), "project_id": int(project_id)}
        params.append(tuple(row.get(f, "") for f in fields))
    if params:
        connection.executemany(
            f"INSERT INTO {ITEMS_TABLE}({','.join(fields)}) VALUES({','.join('?' for _ in fields)})",
            params,
        )
    return len(params)


def _sync_cost_variations(connection, order: dict[str, Any]) -> None:
    pid = int(order["project_id"])
    code = str(order["vo_code"])
    old_row = connection.execute(
        "SELECT * FROM cost_variations WHERE project_id=? AND vo_code=?",
        (pid, code),
    ).fetchone()
    old = _rowdict(old_row)
    connection.execute("DELETE FROM cost_variations WHERE project_id=? AND vo_code=?", (pid, code))
    now = _now()
    connection.execute(
        """INSERT INTO cost_variations(
               project_id,vo_code,task_ref,description,proposed_amount,approved_amount,
               funding_source,status,vo_date,note,created_at,updated_at
           ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
        (
            pid, code, str(old.get("task_ref") or ""),
            f"[VO_EXCEL] {order.get('filename','')} | {order.get('package_name','')}".strip(),
            float(order.get("proposed_amount") or 0),
            float(order.get("approved_amount") or 0),
            str(order.get("funding_source") or ""),
            str(order.get("status") or "Dự thảo"),
            str(order.get("vo_date") or ""),
            str(order.get("note") or ""),
            str(old.get("created_at") or order.get("created_at") or now),
            str(order.get("updated_at") or now),
        ),
    )


def save_vo(db, project_id: int, result: dict[str, Any]) -> dict[str, Any]:
    pid = int(project_id)
    vo_code = str(result.get("vo_code") or "").strip()
    if not vo_code:
        raise ValueError("VO chưa có mã.")
    metadata = dict(result.get("metadata") or {})
    summary = dict(result.get("summary") or {})
    filename = str(result.get("filename") or "VO.xlsx")
    batch_id = str(result.get("batch_id") or "")
    now = _now()
    payload = _encode_result(result)

    with db.connect() as connection:
        _ensure_tables(connection)
        old_row = connection.execute(
            f"SELECT * FROM {VO_TABLE} WHERE project_id=? AND vo_code=?",
            (pid, vo_code),
        ).fetchone()
        old = _rowdict(old_row)
        vo_id = str(old.get("vo_id") or uuid.uuid4().hex)
        old_batch = str(old.get("batch_id") or "")
        latest = int(old.get("latest_revision") or 0)
        revision_no = 0 if old_row is None else latest + 1 if old_batch != batch_id else latest

        order = {
            "vo_id": vo_id,
            "project_id": pid,
            "vo_no": int(result.get("vo_no") or 0),
            "vo_code": vo_code,
            "filename": filename,
            "batch_id": batch_id,
            "revision_label": str(metadata.get("revision_label") or f"R{revision_no}"),
            "vo_date": str(metadata.get("vo_date") or old.get("vo_date") or ""),
            "project_name": str(metadata.get("project") or ""),
            "package_name": str(metadata.get("package") or ""),
            "subtotal_before_vat": float(summary.get("subtotal_before_vat") or 0),
            "vat_amount": float(summary.get("vat_amount") or 0),
            "total_after_vat": float(summary.get("total_after_vat") or 0),
            "increase_amount": float(summary.get("increase_amount") or 0),
            "decrease_amount": float(summary.get("decrease_amount") or 0),
            "proposed_amount": float(summary.get("total_after_vat") or 0),
            "approved_amount": float(old.get("approved_amount") or 0),
            "funding_source": str(old.get("funding_source") or ""),
            "status": str(old.get("status") or "Dự thảo"),
            "latest_revision": int(revision_no),
            "note": str(old.get("note") or ""),
            "created_at": str(old.get("created_at") or now),
            "updated_at": now,
        }
        connection.execute(f"DELETE FROM {VO_TABLE} WHERE vo_id=?", (vo_id,))
        fields = list(order.keys())
        connection.execute(
            f"INSERT INTO {VO_TABLE}({','.join(fields)}) VALUES({','.join('?' for _ in fields)})",
            tuple(order[f] for f in fields),
        )
        connection.execute(f"DELETE FROM {WORKBOOK_TABLE} WHERE vo_id=?", (vo_id,))
        connection.execute(
            f"INSERT INTO {WORKBOOK_TABLE}(vo_id,project_id,filename,batch_id,payload,updated_at) VALUES(?,?,?,?,?,?)",
            (vo_id, pid, filename, batch_id, payload, now),
        )
        if old_row is None or old_batch != batch_id:
            connection.execute(
                f"""INSERT INTO {REVISIONS_TABLE}(
                    revision_id,vo_id,project_id,revision_no,revision_label,filename,batch_id,payload,created_at
                ) VALUES(?,?,?,?,?,?,?,?,?)""",
                (
                    uuid.uuid4().hex, vo_id, pid, revision_no, str(metadata.get("revision_label") or f"R{revision_no}"),
                    filename, batch_id, payload, now,
                ),
            )
        inserted = _replace_items(connection, pid, vo_id, result)
        _sync_cost_variations(connection, order)

    return {
        "vo_id": vo_id,
        "vo_code": vo_code,
        "revision_no": revision_no,
        "inserted_items": inserted,
        "proposed_amount": float(summary.get("total_after_vat") or 0),
        "updated_at": now,
    }


def list_vos(db, project_id: int) -> list[dict[str, Any]]:
    with db.connect() as connection:
        _ensure_tables(connection)
        rows = connection.execute(
            f"SELECT * FROM {VO_TABLE} WHERE project_id=? ORDER BY vo_no,created_at",
            (int(project_id),),
        ).fetchall()
    return [_rowdict(row) for row in rows]


def get_vo(db, vo_id: str) -> dict[str, Any] | None:
    with db.connect() as connection:
        _ensure_tables(connection)
        row = connection.execute(f"SELECT * FROM {VO_TABLE} WHERE vo_id=?", (str(vo_id),)).fetchone()
    return _rowdict(row) if row is not None else None


def load_vo_workbook(db, vo_id: str) -> dict[str, Any] | None:
    with db.connect() as connection:
        _ensure_tables(connection)
        row = connection.execute(
            f"SELECT payload,updated_at FROM {WORKBOOK_TABLE} WHERE vo_id=?",
            (str(vo_id),),
        ).fetchone()
    if row is None:
        return None
    data = _rowdict(row)
    payload = str(data.get("payload") or (row[0] if not data else ""))
    result = _decode_result(payload)
    result["_persisted"] = True
    result["_persisted_at"] = str(data.get("updated_at") or "")
    return result


def vo_items(db, vo_id: str) -> list[dict[str, Any]]:
    with db.connect() as connection:
        _ensure_tables(connection)
        rows = connection.execute(
            f"SELECT * FROM {ITEMS_TABLE} WHERE vo_id=? ORDER BY sheet_name,row_no",
            (str(vo_id),),
        ).fetchall()
    return [_rowdict(row) for row in rows]


def vo_revisions(db, vo_id: str) -> list[dict[str, Any]]:
    with db.connect() as connection:
        _ensure_tables(connection)
        rows = connection.execute(
            f"""SELECT revision_id,vo_id,project_id,revision_no,revision_label,filename,batch_id,created_at
                FROM {REVISIONS_TABLE} WHERE vo_id=? ORDER BY revision_no DESC""",
            (str(vo_id),),
        ).fetchall()
    return [_rowdict(row) for row in rows]


def update_vo_finance(
    db,
    vo_id: str,
    *,
    approved_amount: float,
    status: str,
    funding_source: str = "",
    vo_date: str = "",
    note: str = "",
) -> None:
    now = _now()
    with db.connect() as connection:
        _ensure_tables(connection)
        connection.execute(
            f"""UPDATE {VO_TABLE}
                SET approved_amount=?,status=?,funding_source=?,vo_date=?,note=?,updated_at=?
                WHERE vo_id=?""",
            (
                float(approved_amount or 0), str(status or ""), str(funding_source or ""),
                str(vo_date or ""), str(note or ""), now, str(vo_id),
            ),
        )
        row = connection.execute(f"SELECT * FROM {VO_TABLE} WHERE vo_id=?", (str(vo_id),)).fetchone()
        if row is not None:
            _sync_cost_variations(connection, _rowdict(row))


def delete_vo(db, vo_id: str) -> None:
    with db.connect() as connection:
        _ensure_tables(connection)
        row = connection.execute(f"SELECT project_id,vo_code FROM {VO_TABLE} WHERE vo_id=?", (str(vo_id),)).fetchone()
        data = _rowdict(row)
        if row is None:
            return
        connection.execute(f"DELETE FROM {ITEMS_TABLE} WHERE vo_id=?", (str(vo_id),))
        connection.execute(f"DELETE FROM {WORKBOOK_TABLE} WHERE vo_id=?", (str(vo_id),))
        connection.execute(f"DELETE FROM {REVISIONS_TABLE} WHERE vo_id=?", (str(vo_id),))
        connection.execute(f"DELETE FROM {VO_TABLE} WHERE vo_id=?", (str(vo_id),))
        connection.execute(
            "DELETE FROM cost_variations WHERE project_id=? AND vo_code=?",
            (int(data.get("project_id") or 0), str(data.get("vo_code") or "")),
        )


def compare_vo_to_boq(db, project_id: int, vo_id: str) -> list[dict[str, Any]]:
    items = vo_items(db, vo_id)
    with db.connect() as connection:
        rows = connection.execute(
            "SELECT id,task_ref,boq_item,quantity,unit,unit_price,budget_total FROM cost_budgets WHERE project_id=? ORDER BY id",
            (int(project_id),),
        ).fetchall()
    boq = [_rowdict(r) for r in rows]
    lookup: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for row in boq:
        lookup.setdefault((_norm(row.get("boq_item")), _norm(row.get("unit"))), []).append(row)
    out = []
    for item in items:
        key = (_norm(item.get("description")), _norm(item.get("unit")))
        matches = lookup.get(key) or lookup.get((key[0], "")) or []
        b = matches[0] if matches else {}
        out.append({
            "Khớp BOQ": "Có" if b else "Chưa",
            "Hạng mục": item.get("description") or "",
            "ĐVT": item.get("unit") or "",
            "KL BOQ": float(b.get("quantity") or 0),
            "KL phát sinh": float(item.get("variation_qty") or 0),
            "Loại": item.get("variation_kind") or "",
            "Đơn giá BOQ": float(b.get("unit_price") or 0),
            "Giá trị BOQ": float(b.get("budget_total") or 0),
            "Giá trị VO": float(item.get("variation_amount") or 0),
            "Sheet": item.get("sheet_name") or "",
        })
    return out


def _parse_ui_date(value: Any):
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except Exception:
            pass
    return date.today()


def render_vo_ui(db, project_id: int, *, can_update: bool = True) -> None:
    import pandas as pd
    import streamlit as st

    pid = int(project_id)
    st.markdown("## 🧾 Chi phí phát sinh theo VO")
    st.caption(
        "Mỗi VO là một hồ sơ Excel riêng. VO có thể **tăng (+)** hoặc **giảm (-)**; "
        "hệ thống giữ nguyên dấu âm từ Excel và không ép giá trị về 0."
    )

    with st.expander("➕ Tạo VO mới / cập nhật revision từ Excel", expanded=True):
        uploaded = st.file_uploader(
            "Chọn file VO (.xlsx / .xlsm)",
            type=["xlsx", "xlsm"],
            key=f"vo_upload_{pid}",
            disabled=not can_update,
        )
        parsed = None
        if uploaded is not None:
            try:
                raw = uploaded.getvalue()
                parsed = parse_vo_workbook(raw, uploaded.name)
                s = parsed["summary"]
                st.success(
                    f"Nhận diện {parsed['vo_code']} · {len(parsed.get('workbook_sheet_names') or [])} sheet hiển thị · "
                    f"{len(parsed.get('detail_items') or []):,} dòng phát sinh · độ tin cậy {parsed.get('confidence_pct',0)}%."
                )
                for warning in parsed.get("warnings") or []:
                    st.warning(warning)
                a, b, c, d, e = st.columns(5)
                a.metric("Phát sinh tăng", f"{_money(s.get('increase_amount'))} VND")
                b.metric("Phát sinh giảm", f"{_money(s.get('decrease_amount'))} VND")
                c.metric("Ròng trước VAT", f"{_money(s.get('subtotal_before_vat'))} VND")
                d.metric("VAT", f"{_money(s.get('vat_amount'))} VND")
                e.metric("Ròng sau VAT", f"{_money(s.get('total_after_vat'))} VND")
                meta = parsed.get("metadata") or {}
                st.caption(
                    f"Ngày VO: {meta.get('vo_date') or 'chưa nhận dạng'} · Revision: {meta.get('revision_label','R0')} · "
                    f"Sheet tính VO: {', '.join(parsed.get('source_sheets') or []) or 'chưa xác định'}"
                )
                if st.button(
                    f"💾 Lưu {parsed['vo_code']} vào dự án",
                    type="primary",
                    width="stretch",
                    key=f"vo_save_{pid}_{parsed['vo_code']}",
                    disabled=not can_update,
                ):
                    saved = save_vo(db, pid, parsed)
                    st.success(
                        f"Đã lưu {saved['vo_code']} · revision {saved['revision_no']} · "
                        f"{saved['inserted_items']:,} dòng · giá trị {_money(saved['proposed_amount'])} VND."
                    )
                    st.rerun()
            except Exception as exc:
                st.error(f"Không đọc/lưu được file VO: {exc}")

    orders = list_vos(db, pid)
    if not orders:
        st.info("Chưa có VO Excel nào được lưu trong dự án.")
        return

    proposed_net = sum(float(x.get("proposed_amount") or 0) for x in orders)
    approved_net = sum(float(x.get("approved_amount") or 0) for x in orders)
    total_inc = sum(float(x.get("increase_amount") or 0) for x in orders)
    total_dec = sum(float(x.get("decrease_amount") or 0) for x in orders)
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("VO tăng", f"{_money(total_inc)} VND")
    c2.metric("VO giảm", f"{_money(total_dec)} VND")
    c3.metric("VO ròng đề xuất", f"{_money(proposed_net)} VND")
    c4.metric("VO ròng được duyệt", f"{_money(approved_net)} VND")

    labels = ["📊 Tổng hợp"] + [str(x.get("vo_code") or f"VO {x.get('vo_no')}") for x in orders]
    tabs = st.tabs(labels)
    with tabs[0]:
        rows = []
        for x in orders:
            rows.append({
                "VO": x.get("vo_code") or "",
                "Ngày": x.get("vo_date") or "",
                "Tăng (VND)": float(x.get("increase_amount") or 0),
                "Giảm (VND)": float(x.get("decrease_amount") or 0),
                "Ròng trước VAT": float(x.get("subtotal_before_vat") or 0),
                "VAT": float(x.get("vat_amount") or 0),
                "Đề xuất sau VAT": float(x.get("proposed_amount") or 0),
                "Duyệt": float(x.get("approved_amount") or 0),
                "Revision": int(x.get("latest_revision") or 0),
                "Trạng thái": x.get("status") or "",
            })
        df = pd.DataFrame(rows)
        for col in ("Tăng (VND)", "Giảm (VND)", "Ròng trước VAT", "VAT", "Đề xuất sau VAT", "Duyệt"):
            df[col] = df[col].map(_table_value)
        st.dataframe(df, hide_index=True, width="stretch")

    for tab, order in zip(tabs[1:], orders):
        with tab:
            vo_id = str(order.get("vo_id") or "")
            t1, t2, t3, t4, t5 = st.tabs(["Tổng quan", "Excel", "Chi tiết tăng/giảm", "Đối chiếu BOQ", "Phê duyệt & lịch sử"])
            with t1:
                a, b, c, d, e = st.columns(5)
                a.metric("Tăng", f"{_money(order.get('increase_amount'))} VND")
                b.metric("Giảm", f"{_money(order.get('decrease_amount'))} VND")
                c.metric("Ròng trước VAT", f"{_money(order.get('subtotal_before_vat'))} VND")
                d.metric("VAT", f"{_money(order.get('vat_amount'))} VND")
                e.metric("Sau VAT", f"{_money(order.get('total_after_vat'))} VND")
                st.caption(
                    f"File: {order.get('filename','')} · Ngày: {order.get('vo_date','')} · "
                    f"Revision hiện tại: {order.get('latest_revision',0)} · {order.get('revision_label','')}"
                )
            with t2:
                book = load_vo_workbook(db, vo_id)
                if not book:
                    st.info("Không có snapshot workbook.")
                else:
                    names = list(book.get("workbook_sheet_names") or [])
                    if names:
                        sheet_tabs = st.tabs(names)
                        snapshots = book.get("workbook") or {}
                        for shtab, name in zip(sheet_tabs, names):
                            with shtab:
                                snap = snapshots.get(name) or {}
                                rows = snap.get("rows") or []
                                cols = snap.get("columns") or []
                                if rows and cols:
                                    sdf = pd.DataFrame(rows, columns=cols)
                                    sdf = sdf.map(_table_value)
                                    st.dataframe(sdf, hide_index=True, width="stretch", height=520)
                                    if snap.get("truncated"):
                                        st.caption("Preview đã giới hạn để đảm bảo hiệu năng; dữ liệu chuẩn hóa vẫn được lưu riêng.")
            with t3:
                items = vo_items(db, vo_id)
                if items:
                    idf = pd.DataFrame(items).rename(columns={
                        "sheet_name": "Sheet", "row_no": "Dòng", "description": "Hạng mục", "unit": "ĐVT",
                        "contract_qty": "KL HĐ", "actual_qty": "KL thực tế", "variation_qty": "KL +/-",
                        "material_unit_price": "ĐG vật tư", "labor_unit_price": "ĐG nhân công",
                        "variation_amount": "Giá trị VO", "variation_kind": "Loại", "note": "Ghi chú",
                    })
                    keep = [c for c in ("Sheet","Dòng","Hạng mục","ĐVT","KL HĐ","KL thực tế","KL +/-","ĐG vật tư","ĐG nhân công","Giá trị VO","Loại","Ghi chú") if c in idf.columns]
                    idf = idf[keep]
                    for col in ("KL HĐ","KL thực tế","KL +/-","ĐG vật tư","ĐG nhân công","Giá trị VO"):
                        if col in idf.columns:
                            idf[col] = idf[col].map(_table_value)
                    st.dataframe(idf, hide_index=True, width="stretch", height=560)
                else:
                    st.info("VO chưa có dòng chi tiết.")
            with t4:
                compare = compare_vo_to_boq(db, pid, vo_id)
                if compare:
                    cdf = pd.DataFrame(compare)
                    for col in ("KL BOQ","KL phát sinh","Đơn giá BOQ","Giá trị BOQ","Giá trị VO"):
                        if col in cdf.columns:
                            cdf[col] = cdf[col].map(_table_value)
                    st.dataframe(cdf, hide_index=True, width="stretch", height=520)
                else:
                    st.info("Chưa có dữ liệu để đối chiếu BOQ.")
            with t5:
                st.markdown("#### Phê duyệt VO")
                with st.form(f"vo_finance_{pid}_{vo_id}"):
                    st.caption(
                        "Giá trị đề xuất lấy từ Excel và giữ nguyên dấu. "
                        "VO giảm phải nhập giá trị duyệt âm nếu được phê duyệt giảm chi phí."
                    )
                    proposed = float(order.get("proposed_amount") or 0)
                    st.metric("Giá trị đề xuất", f"{_money(proposed)} VND")
                    approved = st.number_input(
                        "Giá trị được duyệt (VND)",
                        value=float(order.get("approved_amount") or 0),
                        step=1000000.0,
                        key=f"vo_approved_{vo_id}",
                    )
                    statuses = ["Dự thảo","Đã trình","Đang duyệt","Yêu cầu chỉnh sửa","Đã duyệt","Từ chối","Đóng"]
                    current_status = str(order.get("status") or "Dự thảo")
                    status = st.selectbox(
                        "Trạng thái",
                        statuses,
                        index=statuses.index(current_status) if current_status in statuses else 0,
                    )
                    sources = ["","Dự phòng phí","CĐT bổ sung","Điều chuyển ngân sách","Giảm giá trị hợp đồng","Khác"]
                    current_source = str(order.get("funding_source") or "")
                    funding = st.selectbox(
                        "Nguồn/loại điều chỉnh",
                        sources,
                        index=sources.index(current_source) if current_source in sources else 0,
                    )
                    vo_date = st.date_input("Ngày VO", value=_parse_ui_date(order.get("vo_date")))
                    note = st.text_area("Ghi chú", value=str(order.get("note") or ""))
                    submit = st.form_submit_button("💾 Lưu phê duyệt VO", type="primary", disabled=not can_update, width="stretch")
                if submit:
                    update_vo_finance(
                        db, vo_id, approved_amount=float(approved), status=status,
                        funding_source=funding, vo_date=vo_date.strftime("%Y-%m-%d"), note=note,
                    )
                    st.success("Đã cập nhật VO.")
                    st.rerun()

                st.markdown("#### Lịch sử revision")
                revs = vo_revisions(db, vo_id)
                if revs:
                    rdf = pd.DataFrame(revs).rename(columns={
                        "revision_no": "Revision", "revision_label": "Ký hiệu", "filename": "File", "created_at": "Ngày cập nhật",
                    })
                    keep = [c for c in ("Revision","Ký hiệu","File","Ngày cập nhật") if c in rdf.columns]
                    st.dataframe(rdf[keep], hide_index=True, width="stretch")

                if can_update:
                    st.markdown("#### 🗑️ Xóa VO")
                    confirm = st.text_input(
                        f"Nhập {order.get('vo_code','')} để xác nhận xóa toàn bộ VO",
                        key=f"vo_delete_confirm_{vo_id}",
                    )
                    if st.button(
                        f"🗑️ Xóa toàn bộ {order.get('vo_code','')}",
                        key=f"vo_delete_{vo_id}",
                        disabled=confirm.strip().upper() != str(order.get("vo_code") or "").upper(),
                    ):
                        delete_vo(db, vo_id)
                        st.success("Đã xóa VO và toàn bộ dữ liệu Excel/revision liên quan.")
                        st.rerun()
