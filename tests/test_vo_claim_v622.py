import io
import os
import sqlite3
import tempfile
import unittest
from contextlib import contextmanager

from openpyxl import Workbook

from vo_claim_v622 import parse_vo_workbook, save_vo, list_vos, vo_items, update_vo_finance


class _DB:
    def __init__(self, path):
        self.path = path

    @contextmanager
    def connect(self):
        c = sqlite3.connect(self.path)
        c.row_factory = sqlite3.Row
        try:
            yield c
            c.commit()
        finally:
            c.close()


def _workbook_bytes():
    wb = Workbook()
    ws = wb.active
    ws.title = "Tổng hợp"
    ws["A2"] = "TỔNG HỢP GIÁ PHÁT SINH TĂNG GIẢM (VO-03)"
    ws["D3"] = "Lần sửa đổi: R0"
    ws["B7"] = "MÔ TẢ"
    ws["C7"] = "GIÁ PHÁT SINH"
    ws["B13"] = "TỔNG CỘNG (CHƯA BAO GỒM VAT)"
    ws["C13"] = -100
    ws["B14"] = "THUẾ VAT 10%"
    ws["C14"] = -10
    ws["B15"] = "TỔNG CỘNG SAU THUẾ VAT (LÀM TRÒN)"
    ws["C15"] = -110

    d = wb.create_sheet("MEP VO")
    d["A1"] = "BẢNG TỔNG HỢP KHỐI LƯỢNG PHÁT SINH TĂNG GIẢM"
    d["A2"] = "STT"
    d["B2"] = "Nội dung công việc"
    d["C2"] = "Đơn vị"
    d["D2"] = "Phát sinh tăng"
    d["E2"] = "Phát sinh Giảm"
    d["J2"] = "Đơn giá vật tư"
    d["K2"] = "Nhân công"
    d["L2"] = "Thành tiền"
    d["A4"] = 1
    d["B4"] = "Hạng mục tăng"
    d["C4"] = "cái"
    d["D4"] = 2
    d["J4"] = 80
    d["K4"] = 20
    d["L4"] = 200
    d["A5"] = 2
    d["B5"] = "Hạng mục giảm"
    d["C5"] = "cái"
    d["E5"] = 3
    d["J5"] = 80
    d["K5"] = 20
    d["L5"] = -300

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


class VOExcelSignedTests(unittest.TestCase):
    def test_signed_vo_parse_and_persist(self):
        result = parse_vo_workbook(_workbook_bytes(), "2025.09.29 VO-03.xlsx")
        self.assertEqual(result["vo_code"], "VO-03")
        self.assertEqual(result["metadata"]["vo_date"], "2025-09-29")
        self.assertEqual(result["summary"]["subtotal_before_vat"], -100)
        self.assertEqual(result["summary"]["vat_amount"], -10)
        self.assertEqual(result["summary"]["total_after_vat"], -110)
        self.assertEqual(result["summary"]["increase_amount"], 200)
        self.assertEqual(result["summary"]["decrease_amount"], -300)
        self.assertEqual(result["summary"]["detail_discrepancy"], 0)

        reduced = [x for x in result["detail_items"] if x["variation_kind"] == "Giảm"][0]
        self.assertEqual(reduced["decrease_qty"], -3)
        self.assertEqual(reduced["variation_qty"], -3)
        self.assertEqual(reduced["variation_amount"], -300)

        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".db")
        tmp.close()
        try:
            db = _DB(tmp.name)
            with db.connect() as c:
                c.execute("""CREATE TABLE cost_variations(
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    project_id INTEGER NOT NULL,
                    vo_code TEXT NOT NULL,
                    task_ref TEXT DEFAULT '',
                    description TEXT NOT NULL,
                    proposed_amount REAL DEFAULT 0,
                    approved_amount REAL DEFAULT 0,
                    funding_source TEXT DEFAULT '',
                    status TEXT DEFAULT '',
                    vo_date TEXT DEFAULT '',
                    note TEXT DEFAULT '',
                    created_at TEXT DEFAULT '',
                    updated_at TEXT DEFAULT '',
                    UNIQUE(project_id,vo_code)
                )""")
            saved = save_vo(db, 7, result)
            self.assertEqual(saved["proposed_amount"], -110)
            order = list_vos(db, 7)[0]
            self.assertEqual(order["proposed_amount"], -110)
            self.assertEqual(len(vo_items(db, saved["vo_id"])), 2)

            update_vo_finance(
                db, saved["vo_id"], approved_amount=-95, status="Đã duyệt",
                funding_source="Giảm giá trị hợp đồng", vo_date="2025-09-29",
            )
            with db.connect() as c:
                row = c.execute(
                    "SELECT proposed_amount,approved_amount,status FROM cost_variations WHERE project_id=7 AND vo_code='VO-03'"
                ).fetchone()
            self.assertEqual(row["proposed_amount"], -110)
            self.assertEqual(row["approved_amount"], -95)
            self.assertEqual(row["status"], "Đã duyệt")
        finally:
            os.unlink(tmp.name)


if __name__ == "__main__":
    unittest.main()
