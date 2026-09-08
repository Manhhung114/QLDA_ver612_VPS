import io
import os
import unittest

from openpyxl import Workbook

# Force the multicore path even for tiny synthetic workbooks in CI.
os.environ["QLDA_CPU_WORKERS"] = "2"
os.environ["QLDA_PARALLEL_EXCEL_MIN_MB"] = "0"
os.environ["QLDA_PARALLEL_MIN_SHEETS"] = "2"
os.environ["QLDA_MP_START_METHOD"] = "forkserver"

from multicore_excel_v622 import install_multicore_excel, runtime_config
import boq_multisheet_v622 as boq
import ipc_claim_v622 as ipc
import vo_claim_v622 as vo



def _bytes(wb: Workbook) -> bytes:
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()



def _boq_bytes() -> bytes:
    wb = Workbook()
    for idx in range(4):
        ws = wb.active if idx == 0 else wb.create_sheet()
        ws.title = f"BOQ {idx + 1}"
        ws.append(["STT", "Nội dung công việc", "Khối lượng", "ĐVT", "Đơn giá", "Thành tiền"])
        for row in range(1, 11):
            ws.append([row, f"Hạng mục {idx + 1}-{row}", 2, "cái", 100, 200])
    return _bytes(wb)



def _ipc_bytes() -> bytes:
    wb = Workbook()
    decl = wb.active
    decl.title = "KHAI BÁO"
    decl["B4"] = "Dự án test"
    decl["B6"] = "Gói MEP"
    decl["B7"] = "Nhà thầu"
    decl["B12"] = "01"
    decl["B15"] = "HD-01"

    pay = wb.create_sheet("Thanh toán")
    pay["D11"] = 510_000_000_000
    pay["D14"] = 10_000_000_000
    pay["K30"] = 1_000_000_000

    gtht = wb.create_sheet("GTHT")
    gtht["B2"] = "Tên công tác"
    gtht["A6"] = 1
    gtht["B6"] = "Thiết bị test"
    gtht["C6"] = 2
    gtht["D6"] = "cái"
    gtht["I6"] = 100
    gtht["K6"] = 200
    gtht["M6"] = 1
    gtht["N6"] = 1
    gtht["S6"] = 100
    gtht["T6"] = 100
    return _bytes(wb)



def _vo_bytes() -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "Tổng hợp"
    summary["A2"] = "TỔNG HỢP GIÁ PHÁT SINH TĂNG GIẢM (VO-03)"
    summary["D3"] = "Lần sửa đổi: R0"
    summary["B8"] = "Chi tiết MEP"
    summary["C8"] = "='MEP VO'!L6"
    summary["B13"] = "TỔNG CỘNG (CHƯA BAO GỒM VAT)"
    summary["C13"] = -100
    summary["B14"] = "THUẾ VAT 10%"
    summary["C14"] = -10
    summary["B15"] = "TỔNG CỘNG SAU THUẾ VAT (LÀM TRÒN)"
    summary["C15"] = -110

    detail = wb.create_sheet("MEP VO")
    detail["A1"] = "BẢNG TỔNG HỢP KHỐI LƯỢNG PHÁT SINH TĂNG GIẢM"
    detail["A2"] = "STT"
    detail["B2"] = "Nội dung công việc"
    detail["C2"] = "Đơn vị"
    detail["D2"] = "Phát sinh tăng"
    detail["E2"] = "Phát sinh Giảm"
    detail["J2"] = "Đơn giá vật tư"
    detail["K2"] = "Nhân công"
    detail["L2"] = "Thành tiền"
    detail["A4"] = 1
    detail["B4"] = "Hạng mục tăng"
    detail["C4"] = "cái"
    detail["D4"] = 2
    detail["J4"] = 80
    detail["K4"] = 20
    detail["L4"] = 200
    detail["A5"] = 2
    detail["B5"] = "Hạng mục giảm"
    detail["C5"] = "cái"
    detail["E5"] = 3
    detail["J5"] = 80
    detail["K5"] = 20
    detail["L5"] = -300

    extra = wb.create_sheet("Thuyết minh")
    extra["A1"] = "VO test"
    return _bytes(wb)


class MulticoreExcelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        install_multicore_excel()

    def test_runtime_uses_multiple_cpu_when_available(self):
        cfg = runtime_config()
        self.assertGreaterEqual(cfg["cpu_count"], 1)
        if cfg["cpu_count"] > 1:
            self.assertGreaterEqual(cfg["child_workers"], 1)

    def test_boq_multicore(self):
        result = boq.parse_boq_workbook(_boq_bytes(), "BOQ-test.xlsx")
        self.assertEqual(result["detail_line_count"], 40)
        self.assertEqual(result["detail_grand_total"], 8000)
        self.assertTrue(result["_multicore"]["enabled"])
        self.assertGreaterEqual(result["_multicore"]["child_processes_used"], 1)

    def test_ipc_multicore_preview_plus_parent_parse(self):
        result = ipc.parse_ipc_workbook(_ipc_bytes(), "IPC#1.xlsx")
        self.assertEqual(result["claim_code"], "IPC-01")
        self.assertEqual(result["summary"]["contract_value"], 510_000_000_000)
        self.assertEqual(result["detail_line_count"], 1)
        self.assertTrue(result["_multicore"]["enabled"])
        self.assertGreaterEqual(result["_multicore"]["child_processes_used"], 1)

    def test_vo_multicore_preserves_signed_values(self):
        result = vo.parse_vo_workbook(_vo_bytes(), "2025.09.29 VO-03.xlsx")
        self.assertEqual(result["vo_code"], "VO-03")
        self.assertEqual(result["summary"]["increase_amount"], 200)
        self.assertEqual(result["summary"]["decrease_amount"], -300)
        self.assertEqual(result["summary"]["detail_net_amount"], -100)
        self.assertTrue(result["_multicore"]["enabled"])
        self.assertGreaterEqual(result["_multicore"]["child_processes_used"], 1)


if __name__ == "__main__":
    unittest.main()
