from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from boq_multisheet_v622 import (
    AUTO_NOTE_PREFIX,
    SUMMARY_SHEET_NAME,
    build_summary_excel,
    parse_boq_workbook,
    save_boq_summary_to_project,
)
from cloud_db import CloudDatabase


def _sample_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Điện"
    ws.append(["BẢNG BOQ HỆ ĐIỆN"])
    ws.append([])
    ws.append(["STT", "Nội dung công việc", "ĐVT", "Khối lượng", "Đơn giá", "Thành tiền"])
    ws.append([1, "Cáp điện", "m", 1, 100, None])
    ws.append([2, "Tủ điện", "bộ", 2, 200, None])
    ws.append([None, "TỔNG CỘNG", None, None, None, "=SUM(F4:F5)"])

    ws = wb.create_sheet("Nước")
    ws.append(["STT", "Hạng mục", "ĐVT", "Khối lượng", "Đơn giá", "Thành tiền"])
    ws.append([1, "Ống", "m", 3, 100, 300])
    ws.append([2, "Van", "cái", 2, 250, 500])

    ws = wb.create_sheet("Ghi chú")
    ws.append(["Nội dung", "Ghi chú"])
    ws.append(["Không phải BOQ", "Bỏ qua"])

    ws = wb.create_sheet(SUMMARY_SHEET_NAME)
    ws.append(["STT", "Hạng mục", "Thành tiền"])
    ws.append([1, "Điện", 500])
    ws.append([2, "Nước", 800])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class BOQMultiSheetTests(unittest.TestCase):
    def test_detects_detail_sheets_and_avoids_double_counting_summary(self):
        result = parse_boq_workbook(_sample_workbook(), "BOQ_MEP.xlsx")
        self.assertEqual(result["detected_sheets"], ["Điện", "Nước"])
        self.assertEqual([row["line_count"] for row in result["summary"]], [2, 2])
        self.assertEqual(result["grand_total"], 1300.0)
        self.assertTrue(any("bỏ qua sheet tổng hợp" in x.lower() for x in result["warnings"]))

    def test_exports_generated_summary_appendix(self):
        result = parse_boq_workbook(_sample_workbook(), "BOQ_MEP.xlsx")
        payload = build_summary_excel(result)
        wb = load_workbook(io.BytesIO(payload), data_only=True)
        self.assertIn(SUMMARY_SHEET_NAME, wb.sheetnames)
        ws = wb[SUMMARY_SHEET_NAME]
        self.assertEqual(ws["A1"].value, SUMMARY_SHEET_NAME.upper())
        self.assertEqual(ws.cell(row=7, column=4).value, 1300.0)

    def test_saves_per_project_and_preserves_manual_boq(self):
        result = parse_boq_workbook(_sample_workbook(), "BOQ_MEP.xlsx")
        with tempfile.TemporaryDirectory() as tmp:
            db = CloudDatabase(Path(tmp) / "qlda_test.db")
            pid = db.add_project("P01", "Dự án thử")
            db.save_cost_budget(
                pid,
                {
                    "task_ref": "",
                    "boq_item": "BOQ thủ công",
                    "quantity": 1,
                    "unit": "Gói",
                    "unit_price": 99,
                    "budget_total": 99,
                    "contract_type": "",
                    "contractor": "",
                    "note": "manual",
                },
            )

            stats = save_boq_summary_to_project(db, pid, result, replace_existing_excel=True)
            self.assertEqual(stats["inserted"], 2)
            rows = db.cost_budgets(pid)
            self.assertEqual(len(rows), 3)
            self.assertEqual(sum(1 for r in rows if str(r["note"]).startswith(AUTO_NOTE_PREFIX)), 2)
            self.assertTrue(any(r["boq_item"] == "BOQ thủ công" for r in rows))

            # Re-importing the same workbook must not duplicate Excel rows.
            save_boq_summary_to_project(db, pid, result, replace_existing_excel=False)
            rows = db.cost_budgets(pid)
            self.assertEqual(len(rows), 3)
            self.assertTrue(any(r["boq_item"] == "BOQ thủ công" for r in rows))


if __name__ == "__main__":
    unittest.main()
