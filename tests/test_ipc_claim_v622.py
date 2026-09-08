from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from cloud_db import CloudDatabase
from ipc_claim_v622 import (
    compare_claim_to_boq,
    ipc_claim_items,
    ipc_claim_revisions,
    list_ipc_claims,
    load_ipc_workbook,
    parse_ipc_workbook,
    save_ipc_claim,
    update_ipc_claim_finance,
)
from v622_ipc_claim_patch import patch_ipc_claims


def sample_ipc_bytes(claim_no="01", requested=1_226_454_590.68, second_value=500_000.0):
    wb = Workbook()
    khai = wb.active
    khai.title = "KHAI BÁO"
    khai["B4"] = "DỰ ÁN TEST IPC"
    khai["B5"] = "TP.HCM"
    khai["B6"] = "GÓI MEP"
    khai["B7"] = "NHÀ THẦU SIGMA"
    khai["B8"] = "NHÀ THẦU SIGMA"
    khai["B9"] = "123456"
    khai["B10"] = "MIZUHO"
    khai["B11"] = "HÀ NỘI"
    khai["B12"] = claim_no
    khai["B13"] = "04/03/2025"
    khai["E13"] = "25/03/2025"
    khai["B14"] = "VNĐ"
    khai["B15"] = "293/2024/HĐ/SCG-SIGMA"

    payment = wb.create_sheet("Thanh toán")
    payment["D11"] = 510_000_000_000
    payment["D12"] = 94_444_444_444.4
    payment["D14"] = 2_430_038_232
    payment["D15"] = 2_430_038_232
    payment["D18"] = 2_430_038_232
    payment["D19"] = 0
    payment["K24"] = 24_300_382.32
    payment["D25"] = 607_509_558
    payment["D28"] = 571_773_701.18
    payment["D30"] = 1_179_283_259.18
    payment["D31"] = 1_250_754_973
    payment["K30"] = requested
    payment["C32"] = "Một tỷ hai trăm triệu đồng"

    gtht = wb.create_sheet("GTHT")
    gtht["B11"] = "Tên công tác / Diễn giải khối lượng"
    gtht["A15"] = 1
    gtht["B15"] = "Quạt hút tầng hầm"
    gtht["C15"] = 12
    gtht["D15"] = "cái"
    gtht["F15"] = "TDA1000"
    gtht["I15"] = 10_000_000
    gtht["J15"] = 2_000_000
    gtht["K15"] = 144_000_000
    gtht["M15"] = 2
    gtht["N15"] = 2
    gtht["P15"] = 10
    gtht["Q15"] = 10
    gtht["S15"] = 20_000_000
    gtht["T15"] = 20_000_000
    gtht["V15"] = 4_000_000
    gtht["W15"] = 4_000_000
    gtht["Y15"] = 1_000_000
    gtht["Z15"] = 1_000_000
    gtht["AA15"] = 0.2
    gtht["AD15"] = "HVAC"

    gtht["B16"] = "HỆ THỐNG HVAC"
    gtht["S16"] = 999_999_999  # summary/group row must not be imported as detail

    gtht["A17"] = 2
    gtht["B17"] = "Ống gió chống cháy EI45"
    gtht["C17"] = 100
    gtht["D17"] = "m2"
    gtht["S17"] = second_value
    gtht["T17"] = second_value
    gtht["AD17"] = "HVAC"
    gtht["B18"] = "TỔNG CỘNG"

    hidden = wb.create_sheet("DGKL.PCCC")
    hidden.sheet_state = "hidden"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class IPCClaimTests(unittest.TestCase):
    def test_parse_sample_template_and_visible_sheets(self):
        result = parse_ipc_workbook(sample_ipc_bytes(), "IPC#1.xlsx")
        self.assertEqual(result["claim_no"], "01")
        self.assertEqual(result["claim_code"], "IPC-01")
        self.assertEqual(result["metadata"]["contractor"], "NHÀ THẦU SIGMA")
        self.assertAlmostEqual(result["summary"]["requested_amount"], 1_226_454_590.68, places=2)
        self.assertEqual(result["workbook_sheet_names"], ["KHAI BÁO", "Thanh toán", "GTHT"])
        self.assertEqual(result["detail_line_count"], 2)
        self.assertEqual(result["detail_items"][0]["boq_item"], "Quạt hút tầng hầm")
        self.assertAlmostEqual(result["detail_items"][0]["current_value"], 23_000_000, places=2)

    def test_claim_persists_across_db_instances_and_syncs_payment(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "claim.db"
            db = CloudDatabase(path)
            pid = db.add_project("IPC01", "Dự án IPC")
            result = parse_ipc_workbook(sample_ipc_bytes(), "IPC#1.xlsx")
            stats = save_ipc_claim(db, pid, result)
            self.assertEqual(stats["revision_no"], 0)
            self.assertEqual(stats["inserted_items"], 2)

            db2 = CloudDatabase(path)
            claims = list_ipc_claims(db2, pid)
            self.assertEqual(len(claims), 1)
            self.assertEqual(claims[0]["claim_code"], "IPC-01")
            self.assertEqual(len(ipc_claim_items(db2, stats["claim_id"])), 2)
            self.assertEqual(load_ipc_workbook(db2, stats["claim_id"])["claim_no"], "01")
            self.assertEqual(len(ipc_claim_revisions(db2, stats["claim_id"])), 1)

            update_ipc_claim_finance(
                db2,
                stats["claim_id"],
                approved_amount=1_200_000_000,
                disbursed_amount=1_100_000_000,
                payment_status="Đã giải ngân",
                disbursement_date="2025-04-05",
                note="Đã chuyển khoản",
            )
            claim = list_ipc_claims(db2, pid)[0]
            self.assertEqual(claim["payment_status"], "Đã giải ngân")
            self.assertEqual(float(claim["disbursed_amount"]), 1_100_000_000)
            payment = db2.payments(pid)[0]
            self.assertEqual(payment["payment_code"], "IPC-01")
            self.assertEqual(float(payment["paid_amount"]), 1_100_000_000)

            # A changed workbook for the same Claim becomes a new revision, not a new Claim.
            revised = parse_ipc_workbook(sample_ipc_bytes(second_value=700_000), "IPC#1_Rev1.xlsx")
            stats2 = save_ipc_claim(db2, pid, revised)
            self.assertEqual(stats2["claim_id"], stats["claim_id"])
            self.assertEqual(stats2["revision_no"], 1)
            self.assertEqual(len(list_ipc_claims(db2, pid)), 1)
            self.assertEqual(len(ipc_claim_revisions(db2, stats["claim_id"])), 2)
            # Finance approval/disbursement is preserved across workbook revisions.
            claim2 = list_ipc_claims(db2, pid)[0]
            self.assertEqual(float(claim2["approved_amount"]), 1_200_000_000)
            self.assertEqual(float(claim2["disbursed_amount"]), 1_100_000_000)

    def test_claim_to_boq_exact_name_unit_match(self):
        with tempfile.TemporaryDirectory() as tmp:
            db = CloudDatabase(Path(tmp) / "compare.db")
            pid = db.add_project("IPC02", "Dự án đối chiếu")
            with db.connect() as c:
                c.execute(
                    """INSERT INTO cost_budgets(project_id,task_ref,boq_item,quantity,unit,unit_price,budget_total,contract_type,contractor,note,created_at,updated_at)
                       VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (pid, "", "Quạt hút tầng hầm", 12, "cái", 12_000_000, 144_000_000, "", "", "", "", ""),
                )
            result = parse_ipc_workbook(sample_ipc_bytes(), "IPC#1.xlsx")
            stats = save_ipc_claim(db, pid, result)
            rows = compare_claim_to_boq(db, pid, stats["claim_id"])
            matched = [row for row in rows if row["Hạng mục"] == "Quạt hút tầng hầm"][0]
            self.assertEqual(matched["Khớp BOQ"], "Có")
            self.assertEqual(float(matched["KL BOQ"]), 12)

    def test_source_patch_replaces_legacy_payment_body(self):
        source = '''\nclass Dummy:\n    pass\n\ndef render_cost_management(pid: int):\n    cost_tab, payment_tab, vo_tab = st.tabs(["BOQ", "Thanh toán & Giải ngân", "VO"])\n    with cost_tab:\n        st.write("cost")\n    with payment_tab:\n        st.write("legacy payment")\n        st.write("legacy form")\n    with vo_tab:\n        st.write("vo")\n'''
        patched = patch_ipc_claims(source)
        self.assertIn("V6.22 IPC CLAIM PAYMENT V1", patched)
        self.assertIn("_v622_render_ipc_claim_ui(db, pid", patched)
        self.assertNotIn('st.write("legacy payment")', patched)
        compile(patched, "patched.py", "exec")


if __name__ == "__main__":
    unittest.main()
