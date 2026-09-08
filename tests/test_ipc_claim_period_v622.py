from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from cloud_db import CloudDatabase
import ipc_claim_v622 as ipc
from ipc_claim_period_v622 import (
    _clean_date_text,
    sync_claim_periods_from_saved_workbooks,
    update_claim_period,
)


def workbook_bytes(claim_no="03", from_date="0", to_date="0"):
    wb = Workbook()
    ws = wb.active
    ws.title = "KHAI BÁO"
    ws["B4"] = "DỰ ÁN TEST"
    ws["B7"] = "SIGMA"
    ws["B12"] = claim_no
    ws["A13"] = "TỪ NGÀY"
    ws["B13"] = from_date
    ws["D13"] = "ĐẾN NGÀY"
    ws["E13"] = to_date
    ws["B15"] = "HĐ-TEST"

    pay = wb.create_sheet("Thanh toán")
    pay["A10"] = "Giá trị hợp đồng + PLHĐ [1]"
    pay["D10"] = 510_000_000_000
    pay["A15"] = "Lũy kế giá trị nghiệm thu đến nay [4]"
    pay["D15"] = 75_053_143_710
    pay["G31"] = "Giá trị thanh toán kỳ này [18]"
    pay["K31"] = 24_358_633_750.66

    gtht = wb.create_sheet("GTHT")
    gtht["B11"] = "Tên công tác / Diễn giải khối lượng"
    gtht["A15"] = 1
    gtht["B15"] = "Thiết bị"
    gtht["C15"] = 1
    gtht["B16"] = "TỔNG CỘNG"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class IPCClaimPeriodTests(unittest.TestCase):
    def test_zero_dates_are_blank_not_literal_zero(self):
        parsed = ipc.parse_ipc_workbook(workbook_bytes(), "IPC03.xlsx")
        self.assertEqual(parsed["metadata"]["from_date"], "")
        self.assertEqual(parsed["metadata"]["to_date"], "")

    def test_valid_dates_are_read_from_excel_labels(self):
        parsed = ipc.parse_ipc_workbook(
            workbook_bytes(from_date="26/04/2025", to_date="25/05/2025"),
            "IPC03.xlsx",
        )
        self.assertEqual(parsed["metadata"]["from_date"], "2025-04-26")
        self.assertEqual(parsed["metadata"]["to_date"], "2025-05-25")

    def test_manual_update_and_excel_sync(self):
        with tempfile.TemporaryDirectory() as tmp:
            db = CloudDatabase(Path(tmp) / "period.db")
            pid = db.add_project("P01", "Dự án")

            missing = ipc.parse_ipc_workbook(workbook_bytes(), "IPC03.xlsx")
            saved3 = ipc.save_ipc_claim(db, pid, missing)
            claim3 = ipc.list_ipc_claims(db, pid)[0]
            self.assertEqual(_clean_date_text(claim3["from_date"]), "")

            update_claim_period(db, saved3["claim_id"], "26/04/2025", "25/05/2025")
            claim3 = ipc.list_ipc_claims(db, pid)[0]
            self.assertEqual(claim3["from_date"], "2025-04-26")
            self.assertEqual(claim3["to_date"], "2025-05-25")

            valid4 = ipc.parse_ipc_workbook(
                workbook_bytes(claim_no="04", from_date="26/05/2025", to_date="25/06/2025"),
                "IPC04.xlsx",
            )
            ipc.save_ipc_claim(db, pid, valid4)
            # Simulate old DB dates lost while workbook snapshot still contains correct period.
            with db.connect() as c:
                c.execute("UPDATE payment_claims SET from_date='',to_date='' WHERE claim_no='04'")
            result = sync_claim_periods_from_saved_workbooks(db, pid)
            self.assertIn("IPC-04", result["updated"])
            claims = {c["claim_code"]: c for c in ipc.list_ipc_claims(db, pid)}
            self.assertEqual(claims["IPC-04"]["from_date"], "2025-05-26")
            self.assertEqual(claims["IPC-04"]["to_date"], "2025-06-25")


if __name__ == "__main__":
    unittest.main()
