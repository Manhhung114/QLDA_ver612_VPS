from __future__ import annotations

import csv
import io
import json
import math
import re
import unicodedata
from datetime import datetime, timedelta
from typing import Any, Iterable

from qlda.application.google_sheets.service import normalize_production_sheet
from qlda.infrastructure.contractor_data_hub import ContractorDataHubRepository
from qlda.infrastructure.google_sheets.client import GoogleSheetsClient, GoogleSheetsConfigError
from qlda.infrastructure.google_sheets.drive import GoogleWorkspaceClient


GOOGLE_SHEET_MIME = "application/vnd.google-apps.spreadsheet"
GOOGLE_DOC_MIME = "application/vnd.google-apps.document"
GOOGLE_SLIDES_MIME = "application/vnd.google-apps.presentation"
PDF_MIME = "application/pdf"
XLSX_MIMES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
}
TEXT_MIMES = {
    "text/plain",
    "text/csv",
    "text/tab-separated-values",
    "application/json",
}


def _norm(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.lower().replace("đ", "d")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _safe_json(value: Any) -> str:
    try:
        return json.dumps(value, ensure_ascii=False, default=str)
    except Exception:
        return str(value or "")


def _sheet_row_content(headers: list[str], row: list[Any]) -> str:
    pairs: list[str] = []
    for index, value in enumerate(row):
        text = str(value if value is not None else "").strip()
        if not text:
            continue
        label = str(headers[index] if index < len(headers) else f"C{index + 1}").strip() or f"C{index + 1}"
        pairs.append(f"{label}={text}")
    return " | ".join(pairs)


def _choose_header(rows: list[list[Any]]) -> tuple[int, list[str]]:
    """Find a useful header row without assuming one fixed contractor template."""
    if not rows:
        return 0, []
    best_index = 0
    best_score = -1
    best_headers: list[str] = []
    for idx, row in enumerate(rows[:30]):
        headers = [str(x or "").strip() for x in row]
        nonempty = sum(bool(x) for x in headers)
        textual = sum(bool(re.search(r"[A-Za-zÀ-ỹĐđ]", x)) for x in headers)
        zones = sum(x.lower().startswith("zone") for x in headers)
        score = nonempty + textual * 2 + zones * 5
        if score > best_score:
            best_index, best_score, best_headers = idx, score, headers
    return best_index, best_headers


def _chunk_text(text: str, *, size: int = 2600, overlap: int = 180) -> list[str]:
    value = str(text or "").strip()
    if not value:
        return []
    if len(value) <= size:
        return [value]
    out: list[str] = []
    start = 0
    while start < len(value):
        end = min(len(value), start + size)
        chunk = value[start:end].strip()
        if chunk:
            out.append(chunk)
        if end >= len(value):
            break
        start = max(start + 1, end - overlap)
    return out


def _parse_iso(value: str) -> datetime | None:
    text = str(value or "").strip()
    if not text:
        return None
    for candidate in (text, text.replace("Z", "+00:00")):
        try:
            return datetime.fromisoformat(candidate)
        except Exception:
            pass
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(text[:19], fmt)
        except Exception:
            pass
    return None


class ContractorDataHubService:
    """V1-V3 data-space orchestration and synchronization.

    V1: one isolated logical data space per contractor.
    V2: Google Sheets and Drive folder/file connectors.
    V3: durable snapshots, sync history and due-sync support.
    """

    def __init__(self, db, client: GoogleWorkspaceClient | None = None):
        self.db = db
        self.repo = ContractorDataHubRepository(db)
        self.client = client

    def ensure_spaces(self, master_project_id: int, contractors: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        contractor_list = [dict(x) for x in contractors]
        for contractor in contractor_list:
            rows.append(self.repo.ensure_space(int(master_project_id), contractor))
        # Non-destructive bridge from the first Google-Sheets implementation.
        self.repo.migrate_legacy_production_sources(int(master_project_id), contractor_list)
        return rows

    @staticmethod
    def _generic_sheet_records(
        source: dict[str, Any],
        worksheet: str,
        values: list[list[Any]],
        *,
        external_item_id: str,
        external_path: str = "",
        modified_time: str = "",
    ) -> list[dict[str, Any]]:
        if not values:
            return []
        header_idx, headers = _choose_header(values)
        records: list[dict[str, Any]] = []
        for row_no, row in enumerate(values[header_idx + 1 :], start=header_idx + 2):
            content = _sheet_row_content(headers, list(row))
            if not content:
                continue
            records.append({
                "source_name": str(source.get("name") or "Google Sheet"),
                "source_kind": "SHEET",
                "category": str(source.get("category") or "AUTO").upper(),
                "external_item_id": external_item_id,
                "external_path": external_path,
                "worksheet": str(worksheet),
                "record_type": "SHEET_ROW",
                "record_ref": f"row:{row_no}",
                "content": content,
                "source_row": row_no,
                "external_modified_time": modified_time,
            })
        return records

    @staticmethod
    def _production_records(
        source: dict[str, Any],
        worksheet: str,
        values: list[list[Any]],
        *,
        external_item_id: str,
        external_path: str = "",
        modified_time: str = "",
    ) -> list[dict[str, Any]]:
        try:
            rows = normalize_production_sheet(worksheet, values)
        except Exception:
            return []
        out: list[dict[str, Any]] = []
        for row in rows:
            out.append({
                "source_name": str(source.get("name") or "Google Sheet"),
                "source_kind": "SHEET",
                "category": "PRODUCTION",
                "external_item_id": external_item_id,
                "external_path": external_path,
                "worksheet": row.worksheet,
                "record_type": "PRODUCTION",
                "record_ref": f"{row.worksheet}:row:{row.source_row}:{row.zone}",
                "work_item": row.work_item,
                "zone": row.zone,
                "progress_percent": float(row.progress_percent),
                "content": (
                    f"Công tác={row.work_item} | Zone={row.zone} | "
                    f"Tiến độ={float(row.progress_percent):.2f}%"
                ),
                "source_row": int(row.source_row),
                "external_modified_time": modified_time,
            })
        return out

    def _sheet_file_records(
        self,
        source: dict[str, Any],
        spreadsheet_id: str,
        *,
        external_path: str = "",
        modified_time: str = "",
        worksheet_names: list[str] | None = None,
    ) -> tuple[list[dict[str, Any]], int, dict[str, Any]]:
        if self.client is None or not self.client.authorized:
            raise GoogleSheetsConfigError("Nguồn Google riêng tư cần kết nối Google OAuth.")
        meta = self.client.metadata(spreadsheet_id)
        all_tabs = [
            str(x.get("title") or "")
            for x in (meta.get("sheets") or [])
            if not x.get("hidden") and str(x.get("title") or "")
        ]
        wanted = [str(x) for x in (worksheet_names or []) if str(x).strip()]
        tabs = [x for x in all_tabs if not wanted or x in wanted]
        records: list[dict[str, Any]] = []
        production_points = 0
        for worksheet in tabs:
            safe_name = worksheet.replace("'", "''")
            values = self.client.values(spreadsheet_id, f"'{safe_name}'!A:ZZ")
            generic = self._generic_sheet_records(
                source, worksheet, values,
                external_item_id=spreadsheet_id,
                external_path=external_path,
                modified_time=modified_time,
            )
            production = self._production_records(
                source, worksheet, values,
                external_item_id=spreadsheet_id,
                external_path=external_path,
                modified_time=modified_time,
            )
            category = str(source.get("category") or "AUTO").upper()
            if category == "PRODUCTION":
                records.extend(production or generic)
            elif category == "AUTO":
                # Keep generic rows for broad AI retrieval, plus structured
                # production points when the worksheet matches the Zone layout.
                records.extend(generic)
                records.extend(production)
            else:
                records.extend(generic)
            production_points += len(production)
        snapshot = {
            "spreadsheet_id": spreadsheet_id,
            "title": meta.get("title") or source.get("name") or "",
            "worksheets": tabs,
            "record_count": len(records),
            "production_points": production_points,
            "modified_time": modified_time,
        }
        return records, production_points, snapshot

    def _public_sheet_records(self, source: dict[str, Any]) -> tuple[list[dict[str, Any]], int, dict[str, Any]]:
        spreadsheet_id = str(source.get("external_id") or "")
        tabs = list(source.get("worksheet_names") or [])
        # Legacy link-only source encodes gid as __gid__:123.
        gids: list[int] = []
        for value in tabs:
            text = str(value or "")
            if text.startswith("__gid__:") and text.split(":", 1)[1].isdigit():
                gids.append(int(text.split(":", 1)[1]))
        if not gids:
            gids = [0]
        records: list[dict[str, Any]] = []
        production_points = 0
        for gid in gids:
            values = GoogleSheetsClient.public_values(spreadsheet_id, gid)
            worksheet = f"gid={gid}"
            generic = self._generic_sheet_records(
                source, worksheet, values, external_item_id=spreadsheet_id
            )
            production = self._production_records(
                source, worksheet, values, external_item_id=spreadsheet_id
            )
            category = str(source.get("category") or "AUTO").upper()
            records.extend(production if category == "PRODUCTION" and production else generic)
            if category == "AUTO":
                records.extend(generic)
                records.extend(production)
            production_points += len(production)
        return records, production_points, {
            "spreadsheet_id": spreadsheet_id,
            "gids": gids,
            "record_count": len(records),
            "production_points": production_points,
        }

    @staticmethod
    def _text_records(
        source: dict[str, Any],
        text: str,
        *,
        file_meta: dict[str, Any],
        record_type: str = "TEXT_CHUNK",
    ) -> list[dict[str, Any]]:
        out: list[dict[str, Any]] = []
        for index, chunk in enumerate(_chunk_text(text), start=1):
            out.append({
                "source_name": str(file_meta.get("name") or source.get("name") or "File"),
                "source_kind": "DRIVE_FILE",
                "category": str(source.get("category") or "AUTO").upper(),
                "external_item_id": str(file_meta.get("id") or ""),
                "external_path": str(file_meta.get("relative_path") or file_meta.get("name") or ""),
                "record_type": record_type,
                "record_ref": f"chunk:{index}",
                "content": chunk,
                "source_row": index,
                "external_modified_time": str(file_meta.get("modifiedTime") or ""),
            })
        return out

    @staticmethod
    def _pdf_text(data: bytes) -> str:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        parts: list[str] = []
        for page_no, page in enumerate(reader.pages[:250], start=1):
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
            if text.strip():
                parts.append(f"[Trang {page_no}]\n{text.strip()}")
        return "\n\n".join(parts)

    @staticmethod
    def _xlsx_records(
        source: dict[str, Any],
        data: bytes,
        *,
        file_meta: dict[str, Any],
    ) -> tuple[list[dict[str, Any]], int]:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        records: list[dict[str, Any]] = []
        production_points = 0
        try:
            for ws in wb.worksheets[:60]:
                rows: list[list[Any]] = []
                for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    rows.append(list(row))
                    if idx >= 10000:
                        break
                generic = ContractorDataHubService._generic_sheet_records(
                    source, ws.title, rows,
                    external_item_id=str(file_meta.get("id") or ""),
                    external_path=str(file_meta.get("relative_path") or file_meta.get("name") or ""),
                    modified_time=str(file_meta.get("modifiedTime") or ""),
                )
                production = ContractorDataHubService._production_records(
                    source, ws.title, rows,
                    external_item_id=str(file_meta.get("id") or ""),
                    external_path=str(file_meta.get("relative_path") or file_meta.get("name") or ""),
                    modified_time=str(file_meta.get("modifiedTime") or ""),
                )
                category = str(source.get("category") or "AUTO").upper()
                if category == "PRODUCTION":
                    records.extend(production or generic)
                elif category == "AUTO":
                    records.extend(generic)
                    records.extend(production)
                else:
                    records.extend(generic)
                production_points += len(production)
        finally:
            wb.close()
        return records, production_points

    def _drive_file_records(
        self,
        source: dict[str, Any],
        file_meta: dict[str, Any],
    ) -> tuple[list[dict[str, Any]], int]:
        if self.client is None or not self.client.authorized:
            raise GoogleSheetsConfigError("Google Drive folder cần kết nối Google OAuth.")
        fid = str(file_meta.get("id") or "")
        mime = str(file_meta.get("mimeType") or "")
        name = str(file_meta.get("name") or "")
        lower_name = name.lower()
        if mime == GOOGLE_SHEET_MIME:
            records, prod, _ = self._sheet_file_records(
                source, fid,
                external_path=str(file_meta.get("relative_path") or name),
                modified_time=str(file_meta.get("modifiedTime") or ""),
            )
            return records, prod
        if mime == GOOGLE_DOC_MIME:
            data = self.client.export_google_file(fid, "text/plain")
            return self._text_records(
                source, data.decode("utf-8", errors="replace"), file_meta=file_meta, record_type="GOOGLE_DOC"
            ), 0
        if mime == GOOGLE_SLIDES_MIME:
            data = self.client.export_google_file(fid, PDF_MIME)
            return self._text_records(
                source, self._pdf_text(data), file_meta=file_meta, record_type="GOOGLE_SLIDES"
            ), 0

        if mime == PDF_MIME or lower_name.endswith(".pdf"):
            data = self.client.download_file(fid)
            return self._text_records(source, self._pdf_text(data), file_meta=file_meta, record_type="PDF"), 0
        if mime in XLSX_MIMES or lower_name.endswith((".xlsx", ".xlsm", ".xls")):
            if lower_name.endswith(".xls") and not lower_name.endswith(".xlsx"):
                # openpyxl cannot safely parse legacy binary XLS. Keep metadata.
                return self._text_records(
                    source,
                    f"File Excel legacy: {name}. QLDA đã ghi nhận metadata nhưng chưa trích nội dung .xls.",
                    file_meta=file_meta,
                    record_type="FILE_METADATA",
                ), 0
            data = self.client.download_file(fid)
            return self._xlsx_records(source, data, file_meta=file_meta)
        if mime in TEXT_MIMES or lower_name.endswith((".txt", ".csv", ".tsv", ".json", ".md")):
            data = self.client.download_file(fid)
            text = data.decode("utf-8", errors="replace")
            return self._text_records(source, text, file_meta=file_meta, record_type="TEXT_FILE"), 0

        metadata = (
            f"Tên file={name} | MIME={mime} | đường dẫn={file_meta.get('relative_path','')} | "
            f"modified={file_meta.get('modifiedTime','')} | web={file_meta.get('webViewLink','')}"
        )
        return self._text_records(source, metadata, file_meta=file_meta, record_type="FILE_METADATA"), 0

    def sync_source(self, source: dict[str, Any], *, trigger_type: str = "MANUAL") -> dict[str, Any]:
        source = dict(source or {})
        master = int(source.get("master_project_id") or 0)
        run_id = self.repo.begin_sync_run(
            master,
            contractor_id=int(source.get("contractor_id") or 0) or None,
            data_space_id=str(source.get("data_space_id") or ""),
            source_id=str(source.get("source_id") or ""),
            trigger_type=trigger_type,
        )
        discovered = 0
        production_points = 0
        try:
            kind = str(source.get("source_kind") or "SHEET").upper()
            access_mode = str(source.get("access_mode") or "GOOGLE_OAUTH").upper()
            records: list[dict[str, Any]] = []
            snapshot: dict[str, Any] = {"kind": kind, "source": source.get("name")}

            if kind == "SHEET":
                if access_mode == "PUBLIC_LINK":
                    records, production_points, snapshot = self._public_sheet_records(source)
                else:
                    records, production_points, snapshot = self._sheet_file_records(
                        source,
                        str(source.get("external_id") or ""),
                        worksheet_names=list(source.get("worksheet_names") or []),
                    )
            elif kind == "FOLDER":
                if self.client is None or not self.client.authorized:
                    raise GoogleSheetsConfigError("Google Drive folder cần đăng nhập Google.")
                files = self.client.list_folder_tree(
                    str(source.get("external_id") or ""), max_depth=6, max_files=1000
                )
                discovered = len(files)
                for file_meta in files:
                    if len(records) >= 20000:
                        break
                    try:
                        child, prod = self._drive_file_records(source, file_meta)
                        records.extend(child[: max(0, 20000 - len(records))])
                        production_points += prod
                    except Exception as exc:
                        records.append({
                            "source_name": str(file_meta.get("name") or source.get("name") or "File"),
                            "source_kind": "DRIVE_FILE",
                            "category": str(source.get("category") or "AUTO").upper(),
                            "external_item_id": str(file_meta.get("id") or ""),
                            "external_path": str(file_meta.get("relative_path") or ""),
                            "record_type": "SYNC_ERROR",
                            "record_ref": "error",
                            "content": f"Không đọc được file: {exc}",
                            "external_modified_time": str(file_meta.get("modifiedTime") or ""),
                        })
                snapshot = {
                    "kind": "FOLDER",
                    "folder_id": source.get("external_id"),
                    "discovered_files": discovered,
                    "files": [
                        {
                            "id": x.get("id"), "name": x.get("name"), "mimeType": x.get("mimeType"),
                            "modifiedTime": x.get("modifiedTime"), "relative_path": x.get("relative_path"),
                        }
                        for x in files[:1000]
                    ],
                    "record_count": len(records),
                    "production_points": production_points,
                }
            elif kind == "DRIVE_FILE":
                if self.client is None or not self.client.authorized:
                    raise GoogleSheetsConfigError("Google Drive file cần đăng nhập Google.")
                meta = self.client.file_metadata(str(source.get("external_id") or ""))
                records, production_points = self._drive_file_records(source, meta)
                discovered = 1
                snapshot = {"kind": "DRIVE_FILE", "file": meta, "record_count": len(records)}
            else:
                raise ValueError(f"source_kind chưa hỗ trợ: {kind}")

            result = self.repo.replace_records(source, records, snapshot_payload=snapshot)
            self.repo.finish_sync_run(
                run_id,
                status="SUCCESS",
                discovered_files=discovered,
                records_written=int(result.get("records") or 0),
                production_points=production_points,
                message="Đồng bộ thành công.",
            )
            return {
                "ok": True,
                "run_id": run_id,
                "records": int(result.get("records") or 0),
                "discovered_files": discovered,
                "production_points": production_points,
                "snapshot_created": bool(result.get("snapshot_created")),
            }
        except Exception as exc:
            self.repo.mark_source_error(str(source.get("source_id") or ""), str(exc))
            self.repo.finish_sync_run(
                run_id,
                status="ERROR",
                discovered_files=discovered,
                production_points=production_points,
                message=str(exc),
            )
            raise

    def sync_space(self, data_space_id: str, *, trigger_type: str = "MANUAL") -> dict[str, Any]:
        sources = self.repo.list_sources(str(data_space_id), enabled_only=True)
        summary = {"sources": 0, "success": 0, "errors": 0, "records": 0, "files": 0, "production_points": 0, "messages": []}
        for source in sources:
            summary["sources"] += 1
            try:
                result = self.sync_source(source, trigger_type=trigger_type)
                summary["success"] += 1
                summary["records"] += int(result.get("records") or 0)
                summary["files"] += int(result.get("discovered_files") or 0)
                summary["production_points"] += int(result.get("production_points") or 0)
            except Exception as exc:
                summary["errors"] += 1
                summary["messages"].append(f"{source.get('name')}: {exc}")
        return summary

    def sync_project(
        self,
        master_project_id: int,
        *,
        workspace_ids: Iterable[int] | None = None,
        trigger_type: str = "MANUAL",
    ) -> dict[str, Any]:
        spaces = self.repo.list_spaces(
            int(master_project_id), enabled_only=True
        )
        allowed = {int(x) for x in (workspace_ids or []) if int(x) > 0}
        if allowed:
            spaces = [x for x in spaces if int(x.get("workspace_project_id") or 0) in allowed]
        total = {"spaces": len(spaces), "success": 0, "errors": 0, "records": 0, "files": 0, "production_points": 0, "messages": []}
        for space in spaces:
            result = self.sync_space(str(space["data_space_id"]), trigger_type=trigger_type)
            total["records"] += int(result.get("records") or 0)
            total["files"] += int(result.get("files") or 0)
            total["production_points"] += int(result.get("production_points") or 0)
            if int(result.get("errors") or 0):
                total["errors"] += 1
                total["messages"].extend(result.get("messages") or [])
            else:
                total["success"] += 1
        return total

    def sync_due_spaces(self, master_project_id: int) -> dict[str, Any]:
        spaces = self.repo.list_spaces(int(master_project_id), enabled_only=True)
        now = datetime.now()
        due: list[dict[str, Any]] = []
        for space in spaces:
            last = _parse_iso(str(space.get("last_sync") or ""))
            interval = max(15, int(space.get("sync_interval_minutes") or 60))
            if last is None or now - last >= timedelta(minutes=interval):
                due.append(space)
        result = {"due": len(due), "success": 0, "errors": 0, "records": 0}
        for space in due:
            summary = self.sync_space(str(space["data_space_id"]), trigger_type="SCHEDULED")
            result["records"] += int(summary.get("records") or 0)
            if int(summary.get("errors") or 0):
                result["errors"] += 1
            else:
                result["success"] += 1
        return result

    def project_alerts(
        self,
        master_project_id: int,
        *,
        workspace_ids: Iterable[int] | None = None,
    ) -> list[dict[str, Any]]:
        metrics = self.repo.project_metrics(int(master_project_id), workspace_ids=workspace_ids)
        alerts: list[dict[str, Any]] = []
        now = datetime.now()
        for row in metrics:
            code = str(row.get("contractor_code") or "")
            name = str(row.get("contractor_name") or "")
            label = f"{code} - {name}".strip(" -")
            last = _parse_iso(str(row.get("last_sync") or ""))
            if row.get("last_error"):
                alerts.append({"level": "error", "contractor": label, "message": str(row.get("last_error"))})
            if int(row.get("source_count") or 0) == 0:
                alerts.append({"level": "warning", "contractor": label, "message": "Chưa khai báo nguồn dữ liệu."})
            if last is None:
                alerts.append({"level": "warning", "contractor": label, "message": "Chưa đồng bộ dữ liệu lần nào."})
            elif now - last > timedelta(hours=24):
                alerts.append({"level": "warning", "contractor": label, "message": f"Dữ liệu đã cũ hơn 24 giờ (sync {row.get('last_sync')})."})
            avg = row.get("avg_progress")
            try:
                if avg is not None and not math.isnan(float(avg)) and float(avg) < 50:
                    alerts.append({"level": "info", "contractor": label, "message": f"Tiến độ sản lượng trung bình đang ở {float(avg):.1f}%."})
            except Exception:
                pass
        return alerts


class ContractorDataHubAI:
    """V4/V5 retrieval + AI analysis over contractor data spaces.

    Retrieval happens in QLDA first.  Only the compact, authorized context is
    sent to the configured AI provider; raw warehouses are not blindly copied to
    the model prompt.
    """

    def __init__(self, db):
        self.db = db
        self.repo = ContractorDataHubRepository(db)

    @staticmethod
    def _tokens(question: str) -> list[str]:
        stop = {"la", "va", "cua", "cho", "trong", "nhung", "cac", "voi", "theo", "bao", "nhieu", "nao", "gi"}
        return [x for x in _norm(question).split() if len(x) >= 2 and x not in stop][:18]

    def retrieve(
        self,
        master_project_id: int,
        question: str,
        *,
        workspace_ids: Iterable[int] | None = None,
        limit: int = 120,
    ) -> list[dict[str, Any]]:
        rows = self.repo.records(
            int(master_project_id), workspace_ids=workspace_ids, limit=12000
        )
        tokens = self._tokens(question)
        scored: list[tuple[float, dict[str, Any]]] = []
        for row in rows:
            haystack = _norm(" ".join([
                str(row.get("source_name") or ""), str(row.get("category") or ""),
                str(row.get("worksheet") or ""), str(row.get("work_item") or ""),
                str(row.get("zone") or ""), str(row.get("content") or ""),
            ]))
            score = 0.0
            for token in tokens:
                if token in haystack:
                    score += 2.0 + min(3.0, haystack.count(token) * 0.25)
            if str(row.get("record_type") or "") == "PRODUCTION":
                score += 0.6
            if str(row.get("record_type") or "") == "SYNC_ERROR":
                score += 0.4
            if not tokens:
                score = 1.0
            if score > 0:
                scored.append((score, row))
        scored.sort(key=lambda pair: (pair[0], str(pair[1].get("synced_at") or "")), reverse=True)
        selected = [dict(row) for _, row in scored[: max(10, min(int(limit), 250))]]
        if selected:
            return selected
        return [dict(x) for x in rows[: max(10, min(int(limit), 250))]]

    def build_context(
        self,
        master_project_id: int,
        question: str,
        *,
        workspace_ids: Iterable[int] | None = None,
    ) -> str:
        metrics = self.repo.project_metrics(int(master_project_id), workspace_ids=workspace_ids)
        records = self.retrieve(
            int(master_project_id), question, workspace_ids=workspace_ids, limit=140
        )
        alerts = ContractorDataHubService(self.db).project_alerts(
            int(master_project_id), workspace_ids=workspace_ids
        )
        lines = [
            "# CONTRACTOR DATA HUB — DỮ LIỆU ĐÃ ĐỒNG BỘ",
            "QUY TẮC: Mỗi nhà thầu là một kho riêng. Không trộn số liệu giữa hai nhà thầu nếu câu hỏi không yêu cầu tổng hợp.",
            "Mọi kết luận phải dựa trên các dòng DATA bên dưới; nếu dữ liệu thiếu/cũ phải nói rõ.",
            "",
            "## TỔNG QUAN KHO NHÀ THẦU",
        ]
        contractor_by_workspace: dict[int, dict[str, Any]] = {}
        for row in metrics:
            wid = int(row.get("workspace_project_id") or 0)
            contractor_by_workspace[wid] = row
            avg = row.get("avg_progress")
            avg_text = "—" if avg is None else f"{float(avg):.1f}%"
            lines.append(
                f"[WAREHOUSE:{row.get('contractor_code','')}] {row.get('contractor_name','')} | "
                f"nguồn={int(row.get('source_count') or 0)} | records={int(row.get('record_count') or 0)} | "
                f"điểm sản lượng={int(row.get('production_points') or 0)} | tiến độ TB={avg_text} | "
                f"sync={row.get('last_sync') or 'chưa sync'}"
            )
        if alerts:
            lines += ["", "## CẢNH BÁO DỮ LIỆU"]
            for alert in alerts[:30]:
                lines.append(f"[{str(alert.get('level') or '').upper()}] {alert.get('contractor','')}: {alert.get('message','')}")

        lines += ["", "## DỮ LIỆU LIÊN QUAN ĐẾN CÂU HỎI"]
        for row in records:
            meta = contractor_by_workspace.get(int(row.get("workspace_project_id") or 0), {})
            code = str(meta.get("contractor_code") or row.get("contractor_id") or "?")
            source = str(row.get("source_name") or "Nguồn")
            worksheet = str(row.get("worksheet") or "")
            ref = str(row.get("record_ref") or row.get("record_key") or "")[:120]
            tag = f"[DATA:{code}/{source}/{worksheet}/{ref}]"
            extra = ""
            if str(row.get("record_type") or "") == "PRODUCTION":
                extra = (
                    f" | công tác={row.get('work_item','')} | zone={row.get('zone','')} | "
                    f"tiến độ={float(row.get('progress_percent') or 0):.2f}%"
                )
            lines.append(
                f"{tag} type={row.get('record_type','')} | category={row.get('category','')} | "
                f"path={row.get('external_path','')}{extra} | {str(row.get('content') or '')[:1800]}"
            )
        return "\n".join(lines)

    def answer(
        self,
        master_project_id: int,
        question: str,
        *,
        workspace_ids: Iterable[int] | None = None,
    ) -> str:
        q = str(question or "").strip()
        if not q:
            raise ValueError("Câu hỏi AI đang trống.")
        context = self.build_context(
            int(master_project_id), q, workspace_ids=workspace_ids
        )
        prompt = f"""Bạn là AI kiểm soát dữ liệu dự án xây dựng QLDA.

{context}

CÂU HỎI:
{q}

YÊU CẦU TRẢ LỜI:
- Trả lời bằng tiếng Việt, ưu tiên bảng khi so sánh nhiều nhà thầu.
- Giữ nguyên các mã [DATA:...] làm bằng chứng cho kết luận quan trọng.
- Không tự suy đoán dữ liệu không có trong kho.
- Nếu phát hiện dữ liệu cũ, lỗi sync, thiếu nguồn hoặc mâu thuẫn, nêu rõ trước khi kết luận.
- Với tổng hợp toàn dự án, tách số liệu từng nhà thầu trước rồi mới tổng hợp.
"""
        from qlda.runtime_core import settings_store as ss
        from qlda.runtime_core.ai_service import (
            AIServiceError,
            gemini_error_to_service_error,
            openai_error_to_service_error,
        )

        settings = ss.get_ai_runtime_settings()
        provider = str(settings.get("provider") or "openai").lower()
        if provider == "gemini":
            key = str(settings.get("api_key") or "").strip()
            if not key:
                raise AIServiceError("Chưa cấu hình Gemini API key trong Cài đặt hệ thống.")
            try:
                from google import genai
                from google.genai import types

                client = genai.Client(api_key=key)
                model = str(settings.get("model") or "auto").strip() or "auto"
                if model.lower() in {"auto", "default"}:
                    candidates: list[str] = []
                    try:
                        for item in client.models.list():
                            name = str(getattr(item, "name", "") or "").removeprefix("models/")
                            actions = list(getattr(item, "supported_actions", None) or [])
                            if name and (not actions or "generateContent" in actions) and "gemini" in name.lower():
                                candidates.append(name)
                    except Exception:
                        candidates = []
                    preferred = ["gemini-3.7-flash", "gemini-3.5-flash", "gemini-2.5-flash"]
                    model = next((x for x in preferred if x in candidates), candidates[0] if candidates else "gemini-2.5-flash")
                response = client.models.generate_content(
                    model=model,
                    contents=prompt,
                    config=types.GenerateContentConfig(
                        system_instruction="Chỉ phân tích dữ liệu QLDA được cung cấp. Không bịa số liệu."
                    ),
                )
                text = str(getattr(response, "text", "") or "").strip()
                try:
                    client.close()
                except Exception:
                    pass
                return text or "AI không trả về nội dung."
            except AIServiceError:
                raise
            except Exception as exc:
                raise gemini_error_to_service_error(exc) from exc

        key = str(settings.get("api_key") or "").strip()
        if not key:
            raise AIServiceError("Chưa cấu hình OpenAI API key trong Cài đặt hệ thống.")
        try:
            from openai import OpenAI

            client = OpenAI(api_key=key)
            response = client.responses.create(
                model=str(settings.get("model") or "gpt-5-mini"),
                store=False,
                input=[
                    {"role": "developer", "content": "Chỉ phân tích dữ liệu QLDA được cung cấp. Không bịa số liệu."},
                    {"role": "user", "content": prompt},
                ],
            )
            return str(getattr(response, "output_text", "") or "").strip() or "AI không trả về nội dung."
        except AIServiceError:
            raise
        except Exception as exc:
            raise openai_error_to_service_error(exc) from exc
