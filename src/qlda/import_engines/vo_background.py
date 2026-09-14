from __future__ import annotations

import hashlib
import os
from pathlib import Path
from typing import Any, Callable


PATCH_VERSION = "V6.24.4 VO BACKGROUND PIPELINE"

ProgressFn = Callable[[int, str, str], None]
CancelFn = Callable[[], bool]


def _sha256_file(path: Path, chunk_size: int = 4 * 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while True:
            chunk = handle.read(chunk_size)
            if not chunk:
                return digest.hexdigest()
            digest.update(chunk)


def _max_file_bytes() -> int:
    try:
        value = int(os.environ.get("QLDA_VO_WORKER_MAX_FILE_MB", "512") or 512)
    except Exception:
        value = 512
    return max(25, min(value, 2048)) * 1024 * 1024


def _value(values: tuple[Any, ...], column: int | None, default: Any = "") -> Any:
    if not column or column <= 0 or column > len(values):
        return default
    value = values[column - 1]
    return default if value is None else value


def _stream_detail_sheet(
    core,
    value_ws,
    formula_ws,
    *,
    cancelled: CancelFn | None = None,
) -> list[dict[str, Any]]:
    """Preserve VO V6.22 rules while scanning each detail sheet sequentially."""
    headers = core._header_map(value_ws)
    desc_col = core._pick_col(headers, ("noi dung cong viec", "mo ta", "noi dung"))
    if not desc_col:
        return []
    seq_col = core._pick_col(headers, ("stt", " tt "))
    unit_col = core._pick_col(headers, ("don vi",))
    variation_col = core._pick_col(headers, ("phat sinh tang giam",))
    increase_col = core._pick_col(headers, ("phat sinh tang",), ("phat sinh tang giam",))
    decrease_col = core._pick_col(headers, ("phat sinh giam",))
    contract_col = core._pick_col(headers, ("theo hop dong",))
    actual_col = core._pick_col(headers, ("thuc te thi cong",))
    mat_col = core._pick_col(headers, ("don gia vat tu", "don gia vat lieu"))
    labor_col = core._pick_col(headers, ("nhan cong",))
    amount_col = core._pick_col(headers, ("thanh tien",))
    spec_col = core._pick_col(headers, ("quy cach",), ("dieu chinh",))
    code_col = core._pick_col(headers, ("ma hieu",), ("dieu chinh",))
    brand_col = core._pick_col(headers, ("thuong hieu",))
    origin_col = core._pick_col(headers, ("xuat xu",))
    note_col = core._pick_col(headers, ("ghi chu",))

    max_row = int(value_ws.max_row or 0)
    max_col = min(max(int(value_ws.max_column or 1), int(getattr(formula_ws, "max_column", 1) or 1)), 40)
    value_rows = value_ws.iter_rows(min_row=4, max_row=max_row, min_col=1, max_col=max_col, values_only=True)
    formula_rows = (
        formula_ws.iter_rows(min_row=4, max_row=max_row, min_col=1, max_col=max_col, values_only=True)
        if formula_ws is not None
        else None
    )

    items: list[dict[str, Any]] = []
    for offset, value_row in enumerate(value_rows):
        row_no = offset + 4
        if cancelled and row_no % 250 == 0 and cancelled():
            raise InterruptedError("Job VO đã được yêu cầu hủy.")
        values = tuple(value_row)
        formulas = tuple(next(formula_rows)) if formula_rows is not None else ()
        description = str(_value(values, desc_col) or "").strip()
        if not description:
            continue
        normalized = core._norm(description)
        if any(
            token in normalized
            for token in (
                "cong gia tri truoc thue",
                "tong cong",
                "thue vat",
                "ban qlda",
                "tu van giam sat",
                "tong thau",
                "nha thau thi cong truc tiep",
            )
        ):
            continue
        if amount_col and formulas and core._is_rollup_formula(_value(formulas, amount_col)):
            continue

        number = core._to_number
        contract_qty = number(_value(values, contract_col)) if contract_col else None
        actual_qty = number(_value(values, actual_col)) if actual_col else None
        variation_qty = number(_value(values, variation_col)) if variation_col else None
        increase_qty = number(_value(values, increase_col)) if increase_col else None
        decrease_qty = number(_value(values, decrease_col)) if decrease_col else None
        if decrease_qty is not None and decrease_qty > 0:
            decrease_qty = -abs(decrease_qty)
        if variation_qty is None and contract_qty is not None and actual_qty is not None:
            variation_qty = actual_qty - contract_qty
        if variation_qty is None and (increase_qty is not None or decrease_qty is not None):
            variation_qty = float(increase_qty or 0) + float(decrease_qty or 0)

        material_price = number(_value(values, mat_col)) if mat_col else None
        labor_price = number(_value(values, labor_col)) if labor_col else None
        direct_amount = number(_value(values, amount_col)) if amount_col else None
        amount = float(direct_amount or 0)
        quantity = float(variation_qty or 0)
        if abs(quantity) < 1e-12 and abs(amount) < 0.5:
            continue
        kind = (
            "Tăng"
            if amount > 0.5 or (abs(amount) <= 0.5 and quantity > 0)
            else "Giảm"
            if amount < -0.5 or quantity < 0
            else "Chưa định giá"
        )
        items.append(
            {
                "sheet_name": str(value_ws.title),
                "row_no": row_no,
                "seq": str(_value(values, seq_col) or "").strip() if seq_col else "",
                "description": description,
                "unit": str(_value(values, unit_col) or "").strip() if unit_col else "",
                "contract_qty": float(contract_qty or 0),
                "actual_qty": float(actual_qty or 0),
                "increase_qty": float(increase_qty or 0),
                "decrease_qty": float(decrease_qty or 0),
                "variation_qty": quantity,
                "spec": str(_value(values, spec_col) or "").strip() if spec_col else "",
                "item_code": str(_value(values, code_col) or "").strip() if code_col else "",
                "brand": str(_value(values, brand_col) or "").strip() if brand_col else "",
                "origin": str(_value(values, origin_col) or "").strip() if origin_col else "",
                "material_unit_price": float(material_price or 0),
                "labor_unit_price": float(labor_price or 0),
                "unit_price_total": float((material_price or 0) + (labor_price or 0)),
                "variation_amount": amount,
                "variation_kind": kind,
                "note": str(_value(values, note_col) or "").strip() if note_col else "",
            }
        )
    return items


def parse_vo_path(
    path: str | Path,
    filename: str = "VO.xlsx",
    *,
    progress: ProgressFn | None = None,
    cancelled: CancelFn | None = None,
) -> dict[str, Any]:
    """Parse a VO from its durable VPS path without loading the file as bytes."""
    from openpyxl import load_workbook
    import qlda.runtime_core.vo_claim as core
    source = Path(path)
    if not source.exists() or not source.is_file():
        raise FileNotFoundError(f"Không tìm thấy file VO trên VPS: {source}")
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise core.VOWorkbookError("VO Background chỉ hỗ trợ file .xlsx hoặc .xlsm.")
    size = int(source.stat().st_size)
    if size <= 0:
        raise core.VOWorkbookError("File VO đang trống.")
    limit = _max_file_bytes()
    if size > limit:
        raise core.VOWorkbookError(
            f"File VO {size / 1024 / 1024:.1f} MB vượt giới hạn worker {limit / 1024 / 1024:.0f} MB."
        )
    if cancelled and cancelled():
        raise InterruptedError("Job VO đã được yêu cầu hủy.")

    display_name = Path(str(filename or source.name)).name
    if progress:
        progress(8, "Đang kiểm tra SHA file VO", "")
    source_sha256 = _sha256_file(source)
    if progress:
        progress(12, "Đang mở VO từ SSD ở chế độ read-only", "")
    try:
        workbook = load_workbook(str(source), data_only=True, read_only=True, keep_links=False)
        formula_book = load_workbook(str(source), data_only=False, read_only=True, keep_links=False)
    except Exception as exc:
        raise core.VOWorkbookError(f"Không đọc được workbook VO: {exc}") from exc

    try:
        summary_ws = core._find_summary_sheet(workbook)
        if summary_ws is None:
            raise core.VOWorkbookError("Không tìm thấy sheet Tổng hợp VO.")
        formula_summary = formula_book[summary_ws.title] if summary_ws.title in formula_book.sheetnames else None
        metadata = core._summary_metadata(summary_ws, display_name)
        summary, summary_lines, subtotal_row = core._summary_values(summary_ws)
        source_sheets = core._formula_source_sheets(formula_summary, subtotal_row, set(workbook.sheetnames))
        if not source_sheets:
            source_sheets = core._fallback_source_sheets(workbook, summary_ws.title)

        detail_items: list[dict[str, Any]] = []
        total_sheets = max(1, len(source_sheets))
        for index, name in enumerate(source_sheets, start=1):
            if cancelled and cancelled():
                raise InterruptedError("Job VO đã được yêu cầu hủy.")
            if name not in workbook.sheetnames:
                continue
            if progress:
                progress(35 + int((index - 1) * 30 / total_sheets), "Đang quét chi tiết VO", name)
            detail_items.extend(
                _stream_detail_sheet(
                    core,
                    workbook[name],
                    formula_book[name] if name in formula_book.sheetnames else None,
                    cancelled=cancelled,
                )
            )

        visible = [ws for ws in workbook.worksheets if getattr(ws, "sheet_state", "visible") == "visible"]
        snapshots: dict[str, dict[str, Any]] = {}
        for index, sheet in enumerate(visible, start=1):
            if cancelled and cancelled():
                raise InterruptedError("Job VO đã được yêu cầu hủy.")
            if progress:
                progress(66 + int(index * 5 / max(1, len(visible))), "Đang tạo preview VO", sheet.title)
            snapshots[str(sheet.title)] = core._sheet_snapshot(sheet)
        visible_names = [str(ws.title) for ws in visible]
    finally:
        workbook.close()
        formula_book.close()

    priced = [item for item in detail_items if abs(float(item.get("variation_amount") or 0)) >= 0.5]
    increase_amount = sum(max(0.0, float(item.get("variation_amount") or 0)) for item in priced)
    decrease_amount = sum(min(0.0, float(item.get("variation_amount") or 0)) for item in priced)
    detail_net = increase_amount + decrease_amount
    subtotal = float(summary.get("subtotal_before_vat") or 0)
    discrepancy = detail_net - subtotal
    warnings = list(metadata.pop("warnings", []))
    if not source_sheets:
        warnings.append("Không nhận được các sheet chi tiết từ công thức Tổng hợp; parser dùng dò nội dung.")
    if abs(discrepancy) > max(1.0, abs(subtotal) * 0.0001):
        warnings.append(
            f"Tổng dòng chi tiết có giá {core._money(detail_net)} VND lệch "
            f"{core._money(discrepancy)} VND so với Tổng hợp {core._money(subtotal)} VND. "
            "Không tự sửa số liệu; cần kiểm tra workbook."
        )
    confidence = min(
        100,
        55
        + (15 if metadata.get("vo_code") else 0)
        + (10 if subtotal_row else 0)
        + (10 if source_sheets else 0)
        + (10 if abs(discrepancy) <= max(1.0, abs(subtotal) * 0.0001) else 0),
    )
    result = {
        "schema": "qlda_vo_excel_v2",
        "filename": display_name,
        "batch_id": source_sha256,
        "metadata": metadata,
        "vo_no": int(metadata["vo_no"]),
        "vo_code": str(metadata["vo_code"]),
        "summary": {
            **summary,
            "increase_amount": float(increase_amount),
            "decrease_amount": float(decrease_amount),
            "detail_net_amount": float(detail_net),
            "detail_discrepancy": float(discrepancy),
        },
        "summary_lines": summary_lines,
        "source_sheets": source_sheets,
        "detail_items": detail_items,
        "detail_line_count": len(detail_items),
        "workbook": snapshots,
        "workbook_sheet_names": visible_names,
        "confidence_pct": int(confidence),
        "warnings": warnings,
        "pipeline": PATCH_VERSION,
        "source_file_size": size,
        "source_sha256": source_sha256,
    }
    if progress:
        progress(72, f"Đã quét {len(detail_items):,} dòng VO", str(result["vo_code"]))
    return result
