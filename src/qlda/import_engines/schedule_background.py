from __future__ import annotations

import hashlib
import math
import os
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable


PATCH_VERSION = "V6.24.5 SCHEDULE EXCEL BACKGROUND PIPELINE"

ProgressFn = Callable[[int, str, str], None]
CancelFn = Callable[[], bool]


def _sha256_file(path: Path, chunk_size: int = 4 * 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while True:
            chunk = handle.read(chunk_size)
            if not chunk:
                return digest.hexdigest()
            digest.update(chunk)


def _norm(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return " ".join(text.lower().replace("đ", "d").strip().split())


ALIASES = {
    "wbs": "wbs",
    "cong viec": "name",
    "name": "name",
    "bat dau": "start",
    "start": "start",
    "ket thuc": "end",
    "finish": "end",
    "kh %": "planned",
    "tt %": "actual",
    "phu trach": "responsible",
    "predecessor": "predecessor",
    "ghi chu": "note",
}


def _date_text(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(float(value)):
        try:
            from openpyxl.utils.datetime import from_excel

            converted = from_excel(float(value))
            return converted.date().isoformat() if isinstance(converted, datetime) else converted.isoformat()
        except Exception:
            return ""
    text = str(value or "").strip()
    if not text:
        return ""
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except Exception:
            pass
    try:
        return datetime.fromisoformat(text).date().isoformat()
    except Exception:
        return ""


def _percent(value: Any, default: int) -> int:
    if value in (None, ""):
        return max(0, min(100, int(default)))
    try:
        number = float(value)
        if 0 <= number <= 1 and not float(number).is_integer():
            number *= 100
        return max(0, min(100, int(round(number))))
    except Exception:
        return max(0, min(100, int(default)))


def parse_schedule_excel_path(
    path: str | Path,
    filename: str = "TienDo.xlsx",
    *,
    status_date: date | str | None = None,
    progress: ProgressFn | None = None,
    cancelled: CancelFn | None = None,
) -> dict[str, Any]:
    """Stream the first worksheet using the same columns as the legacy importer."""
    from openpyxl import load_workbook
    from qlda.runtime_core.project_store import planned_progress

    source = Path(path)
    if not source.exists() or not source.is_file():
        raise FileNotFoundError(f"Không tìm thấy file tiến độ trên VPS: {source}")
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Schedule Background chỉ hỗ trợ .xlsx/.xlsm; file .xls vẫn dùng luồng nhập cũ.")
    size = int(source.stat().st_size)
    if size <= 0:
        raise ValueError("File Excel tiến độ đang trống.")
    try:
        max_mb = max(25, min(int(os.environ.get("QLDA_SCHEDULE_WORKER_MAX_FILE_MB", "512") or 512), 2048))
    except Exception:
        max_mb = 512
    if size > max_mb * 1024 * 1024:
        raise ValueError(f"File tiến độ vượt giới hạn worker {max_mb} MB.")

    report_date = _date_text(status_date) or date.today().isoformat()
    report_day = datetime.strptime(report_date, "%Y-%m-%d").date()
    if progress:
        progress(8, "Đang kiểm tra SHA file tiến độ", "")
    source_sha256 = _sha256_file(source)
    if cancelled and cancelled():
        raise InterruptedError("Job tiến độ đã được yêu cầu hủy.")
    if progress:
        progress(15, "Đang mở Excel tiến độ từ SSD", "")
    try:
        workbook = load_workbook(str(source), data_only=True, read_only=True, keep_links=False)
    except Exception as exc:
        raise ValueError(f"Không đọc được Excel tiến độ: {exc}") from exc

    try:
        visible = [sheet for sheet in workbook.worksheets if getattr(sheet, "sheet_state", "visible") == "visible"]
        if not visible:
            raise ValueError("Workbook tiến độ không có sheet hiển thị.")
        sheet = visible[0]
        rows = sheet.iter_rows(values_only=True)
        try:
            header = tuple(next(rows))
        except StopIteration as exc:
            raise ValueError("Sheet tiến độ đang trống.") from exc
        columns: dict[str, int] = {}
        for index, value in enumerate(header):
            field = ALIASES.get(_norm(value))
            if field and field not in columns:
                columns[field] = index
        missing = {"name", "start", "end"} - set(columns)
        if missing:
            raise ValueError("Excel cần tối thiểu các cột: Công việc/Name, Bắt đầu/Start, Kết thúc/Finish.")

        tasks: list[dict[str, Any]] = []
        source_rows = 0
        skipped: list[dict[str, Any]] = []
        for row_no, row in enumerate(rows, start=2):
            if cancelled and row_no % 250 == 0 and cancelled():
                raise InterruptedError("Job tiến độ đã được yêu cầu hủy.")
            values = tuple(row)
            if not any(value not in (None, "") for value in values):
                continue
            source_rows += 1

            def cell(field: str, default: Any = "") -> Any:
                index = columns.get(field)
                if index is None or index >= len(values):
                    return default
                value = values[index]
                return default if value is None else value

            name = str(cell("name") or "").strip()
            start = _date_text(cell("start"))
            end = _date_text(cell("end"))
            if not name or not start or not end or end < start:
                skipped.append({"row_no": row_no, "reason": "Thiếu tên/ngày hoặc ngày kết thúc trước ngày bắt đầu"})
                continue
            start_day = datetime.strptime(start, "%Y-%m-%d").date()
            end_day = datetime.strptime(end, "%Y-%m-%d").date()
            planned_default = planned_progress(start, end, report_day)
            tasks.append(
                {
                    "source_row_no": row_no,
                    "wbs": str(cell("wbs") or "").strip(),
                    "name": name,
                    "responsible": str(cell("responsible") or "").strip(),
                    "start_date": start,
                    "end_date": end,
                    "duration": max(1, (end_day - start_day).days + 1),
                    "planned_progress": _percent(cell("planned", None), planned_default),
                    "actual_progress": _percent(cell("actual", None), 0),
                    "predecessor": str(cell("predecessor") or "").strip(),
                    "note": str(cell("note") or "").strip(),
                }
            )
            if progress and row_no % 500 == 0:
                progress(min(68, 20 + int(row_no * 48 / max(1, int(sheet.max_row or row_no)))), "Đang quét công việc", sheet.title)
        sheet_name = str(sheet.title)
    finally:
        workbook.close()

    if not tasks:
        raise ValueError("Không có công việc hợp lệ để ghi PostgreSQL.")
    warnings = []
    if skipped:
        warnings.append(f"Đã bỏ qua {len(skipped):,} dòng thiếu tên/ngày hoặc có khoảng ngày không hợp lệ.")
    result = {
        "pipeline": PATCH_VERSION,
        "filename": Path(str(filename or source.name)).name,
        "batch_id": source_sha256,
        "source_sha256": source_sha256,
        "source_file_size": size,
        "sheet_name": sheet_name,
        "status_date": report_date,
        "tasks": tasks,
        "task_count": len(tasks),
        "detail_line_count": len(tasks),
        "source_row_count": source_rows,
        "skipped_rows": len(skipped),
        "skipped_details": skipped[:50],
        "warnings": warnings,
    }
    if progress:
        progress(72, f"Đã quét {len(tasks):,} công việc hợp lệ", sheet_name)
    return result
