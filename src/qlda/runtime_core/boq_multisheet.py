from __future__ import annotations

import hashlib
import io
import math
import re
import unicodedata
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


AUTO_NOTE_PREFIX = "[QLDA_BOQ_EXCEL]"
SUMMARY_SHEET_NAME = "Phụ lục tổng hợp giá trị"
MAX_WORKBOOK_BYTES = 100 * 1024 * 1024
MAX_PREVIEW_ROWS = 5000
MAX_PREVIEW_COLS = 30


class BOQWorkbookError(ValueError):
    pass


def _norm(value: Any) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.lower().replace("đ", "d")
    text = re.sub(r"[^a-z0-9%]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


_HEADER_ALIASES = {
    "seq": ("stt", "tt", "so thu tu"),
    "item": (
        "noi dung cong viec", "noi dung cong tac", "danh muc cong tac", "ten cong tac",
        "ten hang muc", "hang muc", "noi dung", "dien giai", "mo ta", "description",
        "boq item", "work item", "ten vat tu", "ten thiet bi",
    ),
    "quantity": ("khoi luong", "so luong", "quantity", "qty", "kl"),
    "unit": ("don vi tinh", "dvt", "don vi", "unit"),
    "unit_price": (
        "don gia du toan", "don gia hop dong", "don gia", "unit price", "rate", "gia don vi",
    ),
    "amount": (
        "thanh tien du toan", "thanh tien", "gia tri du toan", "gia tri", "amount", "total amount",
        "tong gia tri", "gia thanh", "tong",
    ),
    "task_ref": ("ma task", "task", "wbs", "ma cong viec", "ma hieu", "ma hang muc"),
}

_SUMMARY_NAME_HINTS = (
    "tong hop", "summary", "phu luc tong hop", "bia", "cover", "muc luc", "tong gia tri",
)
_BOQ_NAME_HINTS = ("boq", "du toan", "khoi luong", "quantity", "bill of quantity", "bill quantities")


def _is_summary_sheet_name(name: Any) -> bool:
    text = _norm(name)
    return text in {"sum", "summary", "tong hop"} or any(hint in text for hint in _SUMMARY_NAME_HINTS)


def _header_kind(value: Any) -> str | None:
    text = _norm(value)
    if not text:
        return None
    matches: list[tuple[int, str]] = []
    for kind, aliases in _HEADER_ALIASES.items():
        for alias in aliases:
            if text == alias or (len(alias) >= 4 and alias in text):
                matches.append((len(alias), kind))
    return max(matches)[1] if matches else None


def _to_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        try:
            number = float(value)
            return number if math.isfinite(number) else None
        except Exception:
            return None
    text = str(value).strip()
    if not text or text.startswith("="):
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    text = re.sub(r"[^0-9,\.\-]", "", text)
    if not text or text in {"-", ".", ","}:
        return None
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        parts = text.split(",")
        if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3 and len(parts[0]) >= 1):
            text = "".join(parts)
        else:
            text = text.replace(",", ".")
    elif "." in text:
        parts = text.split(".")
        if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3 and len(parts[0]) >= 1):
            text = "".join(parts)
    try:
        number = float(text)
        if negative:
            number = -abs(number)
        return number if math.isfinite(number) else None
    except Exception:
        return None


def _header_mapping(values: list[Any], previous: list[Any] | None = None) -> dict[str, int]:
    candidates: dict[str, tuple[int, int]] = {}
    for idx, value in enumerate(values):
        texts = [value]
        if previous is not None and idx < len(previous):
            texts.append(f"{previous[idx] or ''} {value or ''}")
        for text in texts:
            kind = _header_kind(text)
            if not kind:
                continue
            normalized = _norm(text)
            score = 10 if normalized in _HEADER_ALIASES.get(kind, ()) else 5
            old = candidates.get(kind)
            if old is None or score > old[0] or (score == old[0] and kind == "amount" and idx > old[1]):
                candidates[kind] = (score, idx)
    return {kind: idx for kind, (_, idx) in candidates.items()}


def _mapping_score(mapping: dict[str, int], sheet_name: str) -> int:
    score = 0
    score += 5 if "item" in mapping else 0
    score += 3 if "amount" in mapping else 0
    score += 2 if "quantity" in mapping else 0
    score += 2 if "unit_price" in mapping else 0
    score += 1 if "unit" in mapping else 0
    score += 1 if "seq" in mapping else 0
    if any(hint in _norm(sheet_name) for hint in _BOQ_NAME_HINTS):
        score += 2
    return score


def _valid_mapping(mapping: dict[str, int], sheet_name: str) -> bool:
    has_value = "amount" in mapping or ("quantity" in mapping and "unit_price" in mapping)
    if "item" in mapping and has_value:
        return True
    named = any(hint in _norm(sheet_name) for hint in _BOQ_NAME_HINTS)
    return named and "item" in mapping and len(set(mapping) & {"quantity", "unit_price", "amount"}) >= 2


def _find_header(ws, max_scan_rows: int = 45) -> tuple[int, dict[str, int]] | None:
    rows = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, max_scan_rows), values_only=True):
        rows.append(list(row))
    best: tuple[int, int, dict[str, int]] | None = None
    for i, values in enumerate(rows):
        previous = rows[i - 1] if i else None
        mapping = _header_mapping(values, previous)
        if not _valid_mapping(mapping, ws.title):
            continue
        score = _mapping_score(mapping, ws.title)
        candidate = (score, i + 1, mapping)
        if best is None or candidate[0] > best[0]:
            best = candidate
    return (best[1], best[2]) if best else None


def _cell(values: tuple[Any, ...], index: int | None) -> Any:
    if index is None or index < 0 or index >= len(values):
        return None
    return values[index]


def _is_total_label(value: Any) -> bool:
    text = _norm(value)
    if not text:
        return False
    if text in {"tong", "tong cong", "cong", "subtotal", "grand total", "vat", "thue gtgt", "thue vat"}:
        return True
    return (
        text.startswith("tong cong ")
        or text.startswith("tong gia tri ")
        or text.startswith("cong gia tri ")
        or text.startswith("gia tri truoc thue")
        or text.startswith("gia tri sau thue")
        or text.startswith("thue gtgt ")
        or text.startswith("thue vat ")
    )


def _display_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    return value


def _sheet_snapshot(ws) -> dict[str, Any]:
    max_row = min(int(ws.max_row or 1), MAX_PREVIEW_ROWS)
    max_col = min(int(ws.max_column or 1), MAX_PREVIEW_COLS)
    rows: list[list[Any]] = []
    for row_no, values in enumerate(
        ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col, values_only=True),
        start=1,
    ):
        rows.append([row_no] + [_display_value(v) for v in values])
    return {
        "sheet": ws.title,
        "columns": ["Dòng"] + [get_column_letter(i) for i in range(1, max_col + 1)],
        "rows": rows,
        "row_count": int(ws.max_row or 0),
        "col_count": int(ws.max_column or 0),
        "truncated": bool((ws.max_row or 0) > max_row or (ws.max_column or 0) > max_col),
    }


def _parse_sheet(ws) -> dict[str, Any] | None:
    if _is_summary_sheet_name(ws.title):
        return None
    found = _find_header(ws)
    if not found:
        return None
    header_row, mapping = found
    items: list[dict[str, Any]] = []
    skipped_totals = 0
    skipped_groups = 0

    for row_no, values in enumerate(
        ws.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        values = tuple(values)
        item_raw = _cell(values, mapping.get("item"))
        item = str(item_raw or "").strip()
        if not item:
            continue

        seq = _to_number(_cell(values, mapping.get("seq")))
        qty = _to_number(_cell(values, mapping.get("quantity")))
        price = _to_number(_cell(values, mapping.get("unit_price")))
        amount = _to_number(_cell(values, mapping.get("amount")))

        if amount is None and qty is not None and price is not None:
            amount = qty * price
        if amount is None:
            continue

        if _is_total_label(item):
            skipped_totals += 1
            continue

        # Chỉ lấy dòng BOQ thực; bỏ tiêu đề/nhóm/subtotal để không cộng trùng.
        is_detail = bool((seq is not None and seq > 0) or (qty is not None and qty > 0))
        if not is_detail:
            skipped_groups += 1
            continue

        # File MEP thường tách Đơn giá thành Vật tư + Nhân công. Database chỉ có
        # một unit_price nên lấy Thành tiền / Khối lượng để giữ đúng đơn giá all-in.
        all_in_price = price
        if qty not in (None, 0):
            all_in_price = amount / qty

        items.append(
            {
                "row_no": int(row_no),
                "boq_item": item,
                "quantity": float(qty or 0),
                "unit": str(_cell(values, mapping.get("unit")) or "").strip(),
                "unit_price": float(all_in_price or 0),
                "budget_total": float(amount),
                "task_ref": str(_cell(values, mapping.get("task_ref")) or "").strip(),
            }
        )

    if not items:
        return None
    return {
        "sheet": ws.title,
        "header_row": header_row,
        "line_count": len(items),
        "budget_total": float(sum(item["budget_total"] for item in items)),
        "items": items,
        "skipped_totals": skipped_totals,
        "skipped_groups": skipped_groups,
    }


def _extract_summary_sheet(ws) -> dict[str, Any] | None:
    found = _find_header(ws)
    if not found:
        return None
    header_row, mapping = found
    if "item" not in mapping or "amount" not in mapping:
        return None

    item_col = mapping["item"]
    amount_col = mapping["amount"]
    seq_col = mapping.get("seq")
    note_col = None
    header_values = list(next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True)))
    for idx, value in enumerate(header_values):
        if "ghi chu" in _norm(value):
            note_col = idx
            break

    rows: list[dict[str, Any]] = []
    blank_run = 0
    for row_no, values in enumerate(
        ws.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        values = tuple(values)
        label = str(_cell(values, item_col) or "").strip()
        amount = _to_number(_cell(values, amount_col))
        if not label and amount is None:
            blank_run += 1
            if blank_run >= 5 and rows:
                break
            continue
        blank_run = 0
        if not label or amount is None:
            continue
        rows.append(
            {
                "row_no": int(row_no),
                "stt": _cell(values, seq_col),
                "item": label,
                "amount": float(amount),
                "note": str(_cell(values, note_col) or "").strip() if note_col is not None else "",
            }
        )

    if not rows:
        return None

    before_tax = vat = after_tax = None
    for item in rows:
        label = _norm(item["item"])
        if "truoc thue" in label:
            before_tax = item["amount"]
        elif "thue vat" in label or "thue gtgt" in label or label.startswith("vat"):
            vat = item["amount"]
        elif "sau thue" in label:
            after_tax = item["amount"]

    return {
        "sheet": ws.title,
        "header_row": int(header_row),
        "rows": rows,
        "before_tax_total": before_tax,
        "vat_total": vat,
        "after_tax_total": after_tax,
    }


def parse_boq_workbook(data: bytes, filename: str = "BOQ.xlsx") -> dict[str, Any]:
    raw = bytes(data or b"")
    if not raw:
        raise BOQWorkbookError("File BOQ đang trống.")
    if len(raw) > MAX_WORKBOOK_BYTES:
        raise BOQWorkbookError("File BOQ lớn hơn 100 MB; hãy tách file trước khi nhập.")
    suffix = Path(str(filename or "")).suffix.lower()
    if suffix and suffix not in {".xlsx", ".xlsm"}:
        raise BOQWorkbookError("Chỉ hỗ trợ BOQ Excel .xlsx hoặc .xlsm.")

    try:
        wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    except Exception as exc:
        raise BOQWorkbookError(f"Không đọc được workbook Excel: {exc}") from exc

    try:
        visible_sheets = [ws for ws in wb.worksheets if getattr(ws, "sheet_state", "visible") == "visible"]
        snapshots = {ws.title: _sheet_snapshot(ws) for ws in visible_sheets}

        summary_sheet = None
        for ws in visible_sheets:
            if _is_summary_sheet_name(ws.title):
                summary_sheet = _extract_summary_sheet(ws)
                if summary_sheet:
                    break

        parsed = []
        for ws in visible_sheets:
            item = _parse_sheet(ws)
            if item:
                parsed.append(item)
    finally:
        wb.close()

    if not parsed:
        raise BOQWorkbookError(
            "Không nhận diện được sheet BOQ chi tiết. Cần có cột Nội dung công việc và Thành tiền, "
            "hoặc bộ Khối lượng + Đơn giá."
        )

    summary = [
        {
            "sheet": item["sheet"],
            "line_count": int(item["line_count"]),
            "budget_total": float(item["budget_total"]),
            "header_row": int(item["header_row"]),
        }
        for item in parsed
    ]
    detail_items: list[dict[str, Any]] = []
    for sheet_data in parsed:
        sheet_name = sheet_data["sheet"]
        for item in sheet_data["items"]:
            detail_items.append({"sheet": sheet_name, **item})

    detail_total = float(sum(item["budget_total"] for item in detail_items))
    warnings: list[str] = []
    if summary_sheet:
        warnings.append(
            f"Sheet tổng hợp '{summary_sheet['sheet']}' được hiển thị riêng và không cộng lặp vào BOQ chi tiết."
        )
        before_tax = summary_sheet.get("before_tax_total")
        if before_tax is not None:
            tolerance = max(1.0, abs(float(before_tax)) * 1e-8)
            if abs(detail_total - float(before_tax)) > tolerance:
                warnings.append(
                    "Tổng các dòng BOQ chi tiết lệch so với 'Cộng giá trị trước thuế' trong sheet tổng hợp: "
                    f"{detail_total:,.0f} so với {float(before_tax):,.0f} VND."
                )

    before_tax_total = (
        float(summary_sheet["before_tax_total"])
        if summary_sheet and summary_sheet.get("before_tax_total") is not None
        else detail_total
    )
    vat_total = (
        float(summary_sheet["vat_total"])
        if summary_sheet and summary_sheet.get("vat_total") is not None
        else 0.0
    )
    after_tax_total = (
        float(summary_sheet["after_tax_total"])
        if summary_sheet and summary_sheet.get("after_tax_total") is not None
        else before_tax_total + vat_total
    )

    return {
        "filename": Path(str(filename or "BOQ.xlsx")).name,
        "batch_id": hashlib.sha256(raw).hexdigest()[:16],
        "workbook_sheet_names": [ws.title for ws in visible_sheets],
        "workbook_sheets": snapshots,
        "summary_sheet": summary_sheet,
        "detected_sheets": [item["sheet"] for item in parsed],
        "summary": summary,
        "detail_items": detail_items,
        "detail_line_count": len(detail_items),
        "detail_grand_total": detail_total,
        "before_tax_total": before_tax_total,
        "vat_total": vat_total,
        "after_tax_total": after_tax_total,
        "grand_total": after_tax_total,
        "warnings": warnings,
    }


def build_summary_excel(result: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = SUMMARY_SHEET_NAME
    ws.merge_cells("A1:D1")
    ws["A1"] = SUMMARY_SHEET_NAME.upper()
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A2"] = "Nguồn file"
    ws["B2"] = str(result.get("filename") or "")

    appendix = (result.get("summary_sheet") or {}).get("rows") or []
    headers = ["STT", "NỘI DUNG CÔNG VIỆC", "TỔNG (VND)", "GHI CHÚ"]
    for col, value in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=value)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    if appendix:
        for idx, item in enumerate(appendix, 1):
            row = 4 + idx
            ws.cell(row=row, column=1, value=item.get("stt"))
            ws.cell(row=row, column=2, value=str(item.get("item") or ""))
            ws.cell(row=row, column=3, value=float(item.get("amount") or 0))
            ws.cell(row=row, column=4, value=str(item.get("note") or ""))
            ws.cell(row=row, column=3).number_format = "#,##0"
    else:
        for idx, item in enumerate(result.get("summary") or [], 1):
            row = 4 + idx
            ws.cell(row=row, column=1, value=idx)
            ws.cell(row=row, column=2, value=str(item.get("sheet") or ""))
            ws.cell(row=row, column=3, value=float(item.get("budget_total") or 0))
            ws.cell(row=row, column=3).number_format = "#,##0"
        total_row = 5 + len(result.get("summary") or [])
        ws.cell(row=total_row, column=2, value="TỔNG CỘNG").font = Font(bold=True)
        ws.cell(row=total_row, column=3, value=float(result.get("detail_grand_total") or 0)).font = Font(bold=True)
        ws.cell(row=total_row, column=3).number_format = "#,##0"

    widths = [10, 52, 24, 32]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A5"

    detail_ws = wb.create_sheet("Tổng theo sheet BOQ")
    detail_headers = ["STT", "Sheet BOQ", "Số dòng chi tiết", "Giá trị trước thuế (VND)"]
    for col, value in enumerate(detail_headers, 1):
        detail_ws.cell(row=1, column=col, value=value).font = Font(bold=True)
    for idx, item in enumerate(result.get("summary") or [], 1):
        detail_ws.cell(row=idx + 1, column=1, value=idx)
        detail_ws.cell(row=idx + 1, column=2, value=str(item.get("sheet") or ""))
        detail_ws.cell(row=idx + 1, column=3, value=int(item.get("line_count") or 0))
        detail_ws.cell(row=idx + 1, column=4, value=float(item.get("budget_total") or 0))
        detail_ws.cell(row=idx + 1, column=4).number_format = "#,##0"
    detail_ws.freeze_panes = "A2"
    detail_ws.column_dimensions["A"].width = 8
    detail_ws.column_dimensions["B"].width = 34
    detail_ws.column_dimensions["C"].width = 18
    detail_ws.column_dimensions["D"].width = 26

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def save_boq_summary_to_project(
    db,
    project_id: int,
    result: dict[str, Any],
    *,
    replace_existing_excel: bool = True,
) -> dict[str, Any]:
    detail_items = list(result.get("detail_items") or [])
    if not detail_items:
        raise BOQWorkbookError("Không có dòng BOQ chi tiết để lưu.")

    pid = int(project_id)
    filename = Path(str(result.get("filename") or "BOQ.xlsx")).name.replace("|", "_")
    batch_id = re.sub(r"[^a-fA-F0-9]", "", str(result.get("batch_id") or ""))[:32]
    if not batch_id:
        batch_id = hashlib.sha256(repr(detail_items).encode("utf-8")).hexdigest()[:16]
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    deleted = 0

    sql = (
        "INSERT INTO cost_budgets("
        "project_id,task_ref,boq_item,quantity,unit,unit_price,budget_total,"
        "contract_type,contractor,note,created_at,updated_at"
        ") VALUES(?,?,?,?,?,?,?,?,?,?,?,?)"
    )
    params = []
    for item in detail_items:
        sheet = str(item.get("sheet") or "BOQ").replace("|", "_")
        row_no = int(item.get("row_no") or 0)
        note = f"{AUTO_NOTE_PREFIX} file={filename}|sheet={sheet}|row={row_no}|batch={batch_id}"
        params.append(
            (
                pid,
                str(item.get("task_ref") or ""),
                str(item.get("boq_item") or ""),
                float(item.get("quantity") or 0),
                str(item.get("unit") or ""),
                float(item.get("unit_price") or 0),
                float(item.get("budget_total") or 0),
                "",
                "",
                note,
                now,
                now,
            )
        )

    with db.connect() as c:
        if replace_existing_excel:
            cur = c.execute(
                "DELETE FROM cost_budgets WHERE project_id=? AND note LIKE ?",
                (pid, AUTO_NOTE_PREFIX + "%"),
            )
        else:
            cur = c.execute(
                "DELETE FROM cost_budgets WHERE project_id=? AND note LIKE ?",
                (pid, f"{AUTO_NOTE_PREFIX}%batch={batch_id}%"),
            )
        try:
            deleted = max(0, int(cur.rowcount or 0))
        except Exception:
            deleted = 0
        c.executemany(sql, params)

    return {
        "inserted": len(params),
        "deleted": deleted,
        "batch_id": batch_id,
        "detail_grand_total": float(result.get("detail_grand_total") or 0),
        "before_tax_total": float(result.get("before_tax_total") or 0),
        "vat_total": float(result.get("vat_total") or 0),
        "after_tax_total": float(result.get("after_tax_total") or 0),
        "grand_total": float(result.get("after_tax_total") or result.get("detail_grand_total") or 0),
    }
