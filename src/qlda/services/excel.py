from __future__ import annotations

from pathlib import Path
from typing import Any

from qlda.infrastructure.database import make_database
from qlda.services.base import CancelFn, DbFactory, ProgressFn, require_project_id
from qlda.services.boq import BOQService
from qlda.services.ipc import IPCService
from qlda.services.schedule import ScheduleService
from qlda.services.vo import VOService


class ExcelImportService:
    """Dispatch heavy Excel use cases through domain services.

    This is the stable application boundary for both the systemd worker and the
    V6.27 HTTP API. It intentionally contains no Streamlit or systemd code.
    """

    def __init__(
        self,
        db_factory: DbFactory = make_database,
        *,
        boq: BOQService | None = None,
        ipc: IPCService | None = None,
        vo: VOService | None = None,
        schedule: ScheduleService | None = None,
    ):
        self.boq = boq or BOQService(db_factory)
        self.ipc = ipc or IPCService(db_factory)
        self.vo = vo or VOService(db_factory)
        self.schedule = schedule or ScheduleService(db_factory)

    @staticmethod
    def normalize_job_type(value: Any) -> str:
        job_type = str(value or "").strip().upper()
        return "BOQ" if job_type == "BOQ_IMPORT" else job_type

    @staticmethod
    def scan_workbook(
        path: str | Path,
        *,
        progress: ProgressFn | None = None,
        cancelled: CancelFn | None = None,
    ) -> dict[str, Any]:
        from openpyxl import load_workbook

        source = Path(path)
        if source.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ValueError("Background Excel V6.26 hiện hỗ trợ .xlsx/.xlsm.")
        if not source.exists() or not source.is_file():
            raise FileNotFoundError(f"Không tìm thấy file Excel trên VPS: {source}")
        if progress:
            progress(5, "Đang mở workbook từ ổ đĩa", "")

        workbook = load_workbook(
            str(source),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        try:
            sheets = list(workbook.worksheets)
            total = max(1, len(sheets))
            result_sheets: list[dict[str, Any]] = []
            for index, sheet in enumerate(sheets, start=1):
                if cancelled and cancelled():
                    raise InterruptedError("Job đã được yêu cầu hủy.")
                if progress:
                    progress(
                        10 + int((index - 1) * 80 / total),
                        "Đang quét cấu trúc workbook",
                        str(sheet.title),
                    )
                nonempty_rows = 0
                nonempty_cells = 0
                sample_rows: list[list[str]] = []
                for row_no, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                    if cancelled and row_no % 500 == 0 and cancelled():
                        raise InterruptedError("Job đã được yêu cầu hủy.")
                    used = [value for value in values if value not in (None, "")]
                    if used:
                        nonempty_rows += 1
                        nonempty_cells += len(used)
                        if len(sample_rows) < 5:
                            sample_rows.append([str(value)[:120] for value in values[:12]])
                result_sheets.append(
                    {
                        "name": str(sheet.title),
                        "max_row": int(getattr(sheet, "max_row", 0) or 0),
                        "max_column": int(getattr(sheet, "max_column", 0) or 0),
                        "nonempty_rows": nonempty_rows,
                        "nonempty_cells": nonempty_cells,
                        "sample_rows": sample_rows,
                    }
                )
                if progress:
                    progress(
                        10 + int(index * 80 / total),
                        "Đang quét cấu trúc workbook",
                        str(sheet.title),
                    )
            if progress:
                progress(95, "Đang hoàn tất kiểm tra workbook", "")
            return {
                "pipeline": "V6.26 service-layer workbook scan",
                "file_name": source.name,
                "file_size": int(source.stat().st_size),
                "sheet_count": len(result_sheets),
                "sheets": result_sheets,
            }
        finally:
            workbook.close()

    def process_job(
        self,
        job: dict[str, Any],
        path: str | Path,
        file_row: dict[str, Any],
        *,
        progress: ProgressFn | None = None,
        cancelled: CancelFn | None = None,
    ) -> dict[str, Any]:
        job_type = self.normalize_job_type(job.get("job_type"))
        if job_type == "WORKBOOK_SCAN":
            return self.scan_workbook(path, progress=progress, cancelled=cancelled)

        project_id = require_project_id(
            job.get("workspace_project_id") or job.get("project_id"),
            f"workspace_project_id {job_type}",
        )
        filename = str(file_row.get("name") or job.get("file_name") or "").strip()
        file_size = int(file_row.get("size") or 0)
        options = dict(job.get("options") or {})

        if job_type == "BOQ":
            return self.boq.import_file(
                project_id,
                path,
                filename or "BOQ.xlsx",
                file_size=file_size,
                options=options,
                progress=progress,
                cancelled=cancelled,
            )
        if job_type == "IPC":
            return self.ipc.import_file(
                project_id,
                path,
                filename or "IPC.xlsx",
                file_size=file_size,
                progress=progress,
                cancelled=cancelled,
            )
        if job_type == "VO":
            return self.vo.import_file(
                project_id,
                path,
                filename or "VO.xlsx",
                file_size=file_size,
                progress=progress,
                cancelled=cancelled,
            )
        if job_type == "SCHEDULE_EXCEL":
            return self.schedule.import_file(
                project_id,
                path,
                filename or "TienDo.xlsx",
                file_size=file_size,
                options=options,
                progress=progress,
                cancelled=cancelled,
            )
        raise RuntimeError(f"Pipeline {job_type} chưa được bật trong V6.26.")
