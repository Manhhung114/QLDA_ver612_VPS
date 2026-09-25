from __future__ import annotations

"""Document ingestion with explicit provenance for the native AI context layer."""

import csv
import hashlib
import io
import json
import mimetypes
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
from xml.etree import ElementTree as ET

from qlda.application.ai.ports import AIChunk


@dataclass(frozen=True, slots=True)
class DocumentProvenance:
    workspace_project_id: int
    source_name: str
    source_ref: str
    mime_type: str
    checksum: str
    byte_size: int


def document_checksum(data: bytes) -> str:
    return hashlib.sha256(bytes(data or b"")).hexdigest()


def detect_mime(name: str, supplied: str = "") -> str:
    if str(supplied or "").strip():
        return str(supplied).strip().lower()
    return str(mimetypes.guess_type(str(name or ""))[0] or "application/octet-stream").lower()


def extract_document_text(name: str, data: bytes, mime_type: str = "") -> str:
    """Extract text from supported office/text/PDF formats without runtime_core."""
    raw = bytes(data or b"")
    if not raw:
        return ""
    suffix = Path(str(name or "")).suffix.lower()
    mime = detect_mime(name, mime_type)

    if suffix in {".txt", ".md", ".log", ".json", ".xml", ".html", ".htm"} or mime.startswith("text/"):
        return _decode_text(raw)
    if suffix == ".csv" or mime in {"text/csv", "application/csv"}:
        text = _decode_text(raw)
        return "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(text)))
    if suffix == ".pdf" or mime == "application/pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        out: list[str] = []
        for page_no, page in enumerate(reader.pages, start=1):
            text = str(page.extract_text() or "").strip()
            if text:
                out.append(f"[TRANG {page_no}]\n{text}")
        return "\n\n".join(out)
    if suffix in {".xlsx", ".xlsm"} or "spreadsheetml" in mime:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        blocks: list[str] = []
        for ws in wb.worksheets:
            rows: list[str] = []
            for values in ws.iter_rows(values_only=True):
                cells = [str(value).strip() for value in values if value not in (None, "")]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                blocks.append(f"[SHEET {ws.title}]\n" + "\n".join(rows))
        return "\n\n".join(blocks)
    if suffix == ".docx" or "wordprocessingml" in mime:
        return _docx_text(raw)
    return ""


def build_document_chunks(
    workspace_project_id: int,
    name: str,
    data: bytes,
    *,
    mime_type: str = "",
    source_ref: str = "",
    chunk_chars: int = 4200,
    overlap_chars: int = 350,
) -> tuple[DocumentProvenance, tuple[AIChunk, ...]]:
    tenant = int(workspace_project_id)
    if tenant <= 0:
        raise ValueError("workspace_project_id phải > 0")
    raw = bytes(data or b"")
    checksum = document_checksum(raw)
    mime = detect_mime(name, mime_type)
    ref = str(source_ref or f"document:{checksum[:16]}:{Path(str(name or 'file')).name}")
    provenance = DocumentProvenance(tenant, str(name or "file"), ref, mime, checksum, len(raw))
    text = extract_document_text(name, raw, mime)
    if not text.strip():
        return provenance, ()

    size = max(800, min(int(chunk_chars), 12000))
    overlap = max(0, min(int(overlap_chars), size // 3))
    chunks: list[AIChunk] = []
    start = 0
    index = 1
    while start < len(text):
        end = min(len(text), start + size)
        body = text[start:end].strip()
        if body:
            chunk_ref = f"{ref}#chunk={index}"
            chunks.append(AIChunk(
                workspace_project_id=tenant,
                source_kind="DOCUMENT",
                source_name=str(name or "file"),
                source_ref=chunk_ref,
                content=body,
                checksum=hashlib.sha256(f"{tenant}|{chunk_ref}|{body}".encode("utf-8")).hexdigest(),
                metadata={
                    "domain": "document",
                    "mime_type": mime,
                    "document_checksum": checksum,
                    "byte_size": len(raw),
                    "chunk_index": index,
                },
            ))
            index += 1
        if end >= len(text):
            break
        start = max(start + 1, end - overlap)
    return provenance, tuple(chunks)


def ingest_document(store, workspace_project_id: int, name: str, data: bytes, *, mime_type: str = "", source_ref: str = "") -> DocumentProvenance:
    provenance, chunks = build_document_chunks(
        workspace_project_id, name, data, mime_type=mime_type, source_ref=source_ref
    )
    if chunks:
        store.upsert_chunks(chunks)
    return provenance


def _decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1258", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except Exception:
            continue
    return raw.decode("utf-8", errors="replace")


def _docx_text(raw: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        xml = archive.read("word/document.xml")
    root = ET.fromstring(xml)
    words: list[str] = []
    paragraphs: list[str] = []
    # Namespace-independent handling keeps this tiny and avoids adding python-docx.
    for node in root.iter():
        tag = node.tag.rsplit("}", 1)[-1]
        if tag == "t" and node.text:
            words.append(node.text)
        elif tag in {"p", "br"} and words:
            paragraphs.append("".join(words).strip())
            words = []
    if words:
        paragraphs.append("".join(words).strip())
    return "\n".join(p for p in paragraphs if p)


__all__ = [
    "DocumentProvenance",
    "build_document_chunks",
    "detect_mime",
    "document_checksum",
    "extract_document_text",
    "ingest_document",
]
