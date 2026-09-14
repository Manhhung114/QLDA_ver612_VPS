from __future__ import annotations

import argparse
import gc
import os
import signal
import socket
import time
from pathlib import Path
from typing import Any, Callable

from excel_jobs_v624 import (
    PATCH_VERSION,
    cancel_requested,
    claim_next_job,
    complete_job,
    ensure_schema,
    fail_job,
    memory_used_percent,
    recover_stale_jobs,
    resource_limits,
    update_progress,
    worker_heartbeat,
)

_STOP = False


def _stop(*_args) -> None:
    global _STOP
    _STOP = True


def _worker_id() -> str:
    return f"{socket.gethostname()}:{os.getpid()}"


def _make_db():
    """Create the same PostgreSQL-backed CloudDatabase used by Streamlit."""
    import postgres_backend_v622 as pg

    pg.install_postgres_backend()
    from cloud_db import CloudDatabase

    return CloudDatabase(Path(os.environ.get("QLDA_WORKER_DB_LABEL", "/opt/qlda/shared/qlda-worker.db")))


def scan_workbook_path(
    path: str | Path,
    *,
    progress: Callable[[int, str, str], None] | None = None,
    cancelled: Callable[[], bool] | None = None,
) -> dict[str, Any]:
    """Read an xlsx/xlsm from disk in read-only mode without loading file bytes."""
    from openpyxl import load_workbook

    source = Path(path)
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Background Excel V6.24 hiện hỗ trợ .xlsx/.xlsm.")
    if not source.exists() or not source.is_file():
        raise FileNotFoundError(f"Không tìm thấy file Excel trên VPS: {source}")

    if progress:
        progress(5, "Đang mở workbook từ ổ đĩa", "")
    wb = load_workbook(str(source), read_only=True, data_only=True, keep_links=False)
    try:
        sheets = list(wb.worksheets)
        total = max(1, len(sheets))
        result_sheets: list[dict[str, Any]] = []
        for index, ws in enumerate(sheets, start=1):
            if cancelled and cancelled():
                raise InterruptedError("Job đã được yêu cầu hủy.")
            base = 10 + int((index - 1) * 80 / total)
            if progress:
                progress(base, "Đang quét cấu trúc workbook", ws.title)

            nonempty_rows = 0
            nonempty_cells = 0
            sample_rows: list[list[str]] = []
            for row_no, values in enumerate(ws.iter_rows(values_only=True), start=1):
                if cancelled and row_no % 500 == 0 and cancelled():
                    raise InterruptedError("Job đã được yêu cầu hủy.")
                used = [value for value in values if value not in (None, "")]
                if used:
                    nonempty_rows += 1
                    nonempty_cells += len(used)
                    if len(sample_rows) < 5:
                        sample_rows.append([str(value)[:120] for value in values[:12]])
            result_sheets.append(
                {
                    "name": str(ws.title),
                    "max_row": int(getattr(ws, "max_row", 0) or 0),
                    "max_column": int(getattr(ws, "max_column", 0) or 0),
                    "nonempty_rows": nonempty_rows,
                    "nonempty_cells": nonempty_cells,
                    "sample_rows": sample_rows,
                }
            )
            if progress:
                progress(10 + int(index * 80 / total), "Đang quét cấu trúc workbook", ws.title)

        if progress:
            progress(95, "Đang hoàn tất kiểm tra workbook", "")
        return {
            "pipeline": "V6.24.2 background disk-stream scan",
            "file_name": source.name,
            "file_size": int(source.stat().st_size),
            "sheet_count": len(result_sheets),
            "sheets": result_sheets,
        }
    finally:
        wb.close()


def _process_boq(
    job: dict[str, Any],
    path: Path,
    file_row: dict[str, Any],
    *,
    report: Callable[[int, str, str], None],
    is_cancelled: Callable[[], bool],
) -> dict[str, Any]:
    from boq_background_v624 import parse_boq_path
    from boq_persist_v624 import save_boq_result_batched
    from boq_persistence_v622 import save_saved_boq_workbook

    job_id = int(job["id"])
    workspace_pid = int(job.get("workspace_project_id") or job.get("project_id") or 0)
    if workspace_pid <= 0:
        raise ValueError("Excel job thiếu workspace_project_id hợp lệ.")

    filename = str(file_row.get("name") or job.get("file_name") or "BOQ.xlsx")
    report(5, "Worker đã nhận BOQ", "")
    result = parse_boq_path(
        path,
        filename,
        progress=report,
        cancelled=is_cancelled,
    )
    if is_cancelled():
        raise InterruptedError("Job BOQ đã được yêu cầu hủy.")

    db = _make_db()
    options = dict(job.get("options") or {})
    report(75, "Đang chuẩn bị ghi BOQ vào PostgreSQL", "")
    stats = save_boq_result_batched(
        db,
        workspace_pid,
        result,
        replace_existing_excel=bool(options.get("replace_existing_excel", True)),
        progress=report,
        cancelled=is_cancelled,
    )
    if is_cancelled():
        raise InterruptedError("Job BOQ đã được yêu cầu hủy.")

    report(94, "Đang lưu snapshot workbook dùng chung", "")
    save_saved_boq_workbook(db, workspace_pid, result)
    report(98, "Đang hoàn tất BOQ", "")

    summary = {
        "pipeline": "V6.24.2 BOQ background",
        "job_type": "BOQ",
        "workspace_project_id": workspace_pid,
        "filename": filename,
        "file_size": int(file_row.get("size") or path.stat().st_size),
        "inserted": int(stats.get("inserted") or 0),
        "expected_rows": int(stats.get("expected_rows") or 0),
        "scanned_rows": int(stats.get("scanned_rows") or 0),
        "prepared_rows": int(stats.get("prepared_rows") or 0),
        "written_rows": int(stats.get("written_rows") or 0),
        "failed_rows": int(stats.get("failed_rows") or 0),
        "verification_status": str(stats.get("verification_status") or "CHƯA ĐỦ"),
        "verified_postgresql": bool(stats.get("verified_postgresql")),
        "deleted": int(stats.get("deleted") or 0),
        "detail_line_count": int(result.get("detail_line_count") or 0),
        "before_tax_total": float(stats.get("before_tax_total") or 0),
        "vat_total": float(stats.get("vat_total") or 0),
        "after_tax_total": float(stats.get("after_tax_total") or 0),
        "material_cost_total": float(stats.get("material_cost_total") or 0),
        "labor_cost_total": float(stats.get("labor_cost_total") or 0),
        "material_component_line_count": int(stats.get("material_component_line_count") or 0),
        "labor_component_line_count": int(stats.get("labor_component_line_count") or 0),
        "batch_id": str(stats.get("batch_id") or result.get("batch_id") or ""),
    }
    del result
    gc.collect()
    print(
        f"Excel job #{job_id} BOQ PostgreSQL verification "
        f"workspace={workspace_pid} expected={summary['expected_rows']} "
        f"scanned={summary['scanned_rows']} written={summary['written_rows']} "
        f"failed={summary['failed_rows']} status={summary['verification_status']}",
        flush=True,
    )
    return summary


def _process_ipc(
    job: dict[str, Any],
    path: Path,
    file_row: dict[str, Any],
    *,
    report: Callable[[int, str, str], None],
    is_cancelled: Callable[[], bool],
) -> dict[str, Any]:
    from ipc_background_v624 import parse_ipc_path
    from ipc_persist_v624 import save_ipc_result_batched

    job_id = int(job["id"])
    workspace_pid = int(job.get("workspace_project_id") or job.get("project_id") or 0)
    if workspace_pid <= 0:
        raise ValueError("Excel job IPC thiếu workspace_project_id hợp lệ.")

    filename = str(file_row.get("name") or job.get("file_name") or "IPC.xlsx")
    report(5, "Worker đã nhận IPC", "")
    result = parse_ipc_path(
        path,
        filename,
        progress=report,
        cancelled=is_cancelled,
    )
    if is_cancelled():
        raise InterruptedError("Job IPC đã được yêu cầu hủy.")

    db = _make_db()
    report(75, "Đang chuẩn bị ghi IPC vào PostgreSQL", str(result.get("claim_code") or ""))
    stats = save_ipc_result_batched(
        db,
        workspace_pid,
        result,
        progress=report,
        cancelled=is_cancelled,
    )
    if is_cancelled():
        raise InterruptedError("Job IPC đã được yêu cầu hủy.")
    report(98, "Đang hoàn tất IPC", str(stats.get("claim_code") or ""))

    summary = {
        "pipeline": "V6.24.3 IPC background",
        "job_type": "IPC",
        "workspace_project_id": workspace_pid,
        "filename": filename,
        "file_size": int(file_row.get("size") or path.stat().st_size),
        "claim_id": str(stats.get("claim_id") or ""),
        "claim_no": str(stats.get("claim_no") or ""),
        "claim_code": str(stats.get("claim_code") or ""),
        "revision_no": int(stats.get("revision_no") or 0),
        "expected_rows": int(stats.get("expected_rows") or 0),
        "scanned_rows": int(stats.get("scanned_rows") or 0),
        "prepared_rows": int(stats.get("prepared_rows") or 0),
        "inserted_rows": int(stats.get("inserted_rows") or 0),
        "written_rows": int(stats.get("written_rows") or 0),
        "failed_rows": int(stats.get("failed_rows") or 0),
        "verification_status": str(stats.get("verification_status") or "CHƯA ĐỦ"),
        "verified_postgresql": bool(stats.get("verified_postgresql")),
        "requested_amount": float(stats.get("requested_amount") or 0),
        "certified_cumulative": float(stats.get("certified_cumulative") or 0),
        "detail_line_count": int(result.get("detail_line_count") or 0),
        "batch_id": str(stats.get("batch_id") or result.get("batch_id") or ""),
        "source_sha256": str(result.get("source_sha256") or ""),
    }
    del result
    gc.collect()
    print(
        f"Excel job #{job_id} IPC PostgreSQL verification "
        f"workspace={workspace_pid} claim={summary['claim_code']} "
        f"expected={summary['expected_rows']} scanned={summary['scanned_rows']} "
        f"written={summary['written_rows']} failed={summary['failed_rows']} "
        f"status={summary['verification_status']}",
        flush=True,
    )
    return summary


def _process_job(job: dict[str, Any]) -> dict[str, Any]:
    job_id = int(job["id"])
    job_type = str(job.get("job_type") or "").strip().upper()
    if job_type == "BOQ_IMPORT":
        job_type = "BOQ"

    from local_vps_backend_v622 import local_file_path

    file_row, path = local_file_path(str(job.get("file_id") or ""))

    def report(percent: int, stage: str, sheet: str = "") -> None:
        update_progress(job_id, percent, stage, sheet)

    def is_cancelled() -> bool:
        return _STOP or cancel_requested(job_id)

    if job_type == "WORKBOOK_SCAN":
        return scan_workbook_path(path, progress=report, cancelled=is_cancelled)
    if job_type == "BOQ":
        return _process_boq(
            job,
            path,
            file_row,
            report=report,
            is_cancelled=is_cancelled,
        )
    if job_type == "IPC":
        return _process_ipc(
            job,
            path,
            file_row,
            report=report,
            is_cancelled=is_cancelled,
        )
    raise RuntimeError(
        f"Pipeline {job_type} chưa được bật. V6.24.3 hiện chạy production BOQ + IPC; "
        "VO/Tiến độ Excel sẽ chuyển ở các bước tiếp theo."
    )


def run_once(worker_id: str) -> bool:
    limits = resource_limits()
    ram = memory_used_percent()
    if ram >= float(limits["hold_percent"]):
        state = "PAUSED_RAM" if ram >= float(limits["pause_percent"]) else "HOLD_RAM"
        worker_heartbeat(worker_id, status=f"{state}:{ram:.1f}%")
        return False

    job = claim_next_job(worker_id)
    if not job:
        worker_heartbeat(worker_id, status="IDLE")
        return False

    job_id = int(job["id"])
    worker_heartbeat(worker_id, status="RUNNING", current_job_id=job_id)
    print(
        f"Excel job #{job_id} started type={job.get('job_type')} file={job.get('file_name')}",
        flush=True,
    )
    try:
        summary = _process_job(job)
        if cancel_requested(job_id) or _STOP:
            fail_job(job_id, "Job đã được yêu cầu hủy.", retry=False)
        else:
            complete_job(job_id, summary)
            print(f"Excel job #{job_id} DONE", flush=True)
    except InterruptedError as exc:
        fail_job(job_id, str(exc), retry=False)
        print(f"Excel job #{job_id} CANCELLED: {exc}", flush=True)
    except Exception as exc:
        fail_job(job_id, f"{type(exc).__name__}: {exc}", retry=True)
        print(f"Excel job #{job_id} ERROR: {type(exc).__name__}: {exc}", flush=True)
    finally:
        worker_heartbeat(worker_id, status="IDLE")
        gc.collect()
    return True


def _recycle_after_job() -> bool:
    value = str(os.environ.get("QLDA_EXCEL_RECYCLE_AFTER_JOB", "1") or "1").strip().lower()
    return value not in {"0", "false", "no", "off"}


def main() -> None:
    parser = argparse.ArgumentParser(description="QLDA V6.24.3 background Excel worker")
    parser.add_argument("--once", action="store_true", help="Process at most one job then exit")
    parser.add_argument("--poll-seconds", type=float, default=float(os.environ.get("QLDA_EXCEL_POLL_SECONDS", "2")))
    args = parser.parse_args()

    signal.signal(signal.SIGTERM, _stop)
    signal.signal(signal.SIGINT, _stop)
    ensure_schema()
    try:
        recovered = recover_stale_jobs()
        if recovered:
            print(f"Recovered {recovered} stale Excel job(s).", flush=True)
    except Exception as exc:
        print(f"Excel stale-job recovery warning: {exc}", flush=True)

    wid = _worker_id()
    worker_heartbeat(wid, status="STARTING")
    poll = max(0.5, min(float(args.poll_seconds), 30.0))

    while not _STOP:
        processed = run_once(wid)
        if args.once:
            break
        # Recycle after each heavy job so openpyxl/Python arenas are returned to
        # Linux. systemd Restart=always starts the next clean worker process.
        if processed and _recycle_after_job():
            break
        if not processed:
            time.sleep(poll)

    worker_heartbeat(wid, status="STOPPED")
    print(f"{PATCH_VERSION} worker stopped: {wid}", flush=True)


if __name__ == "__main__":
    main()
