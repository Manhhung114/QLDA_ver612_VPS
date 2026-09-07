from __future__ import annotations

import hashlib
import io
import math
import re
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


AUTO_NOTE_PREFIX = "[QLDA_BOQ_EXCEL]"
SUMMARY_SHEET_NAME = "Phụ lục tổng hợp giá trị"
MAX_WORKBOOK_BYTES = 100 * 1024 * 1024


class BOQWorkbookError(ValueError):
    pass


def _norm(value: Any) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.lower().replace("đ", "d")
    text = re.sub(r"[^a-z0-9%]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


_HEADER_ALIASES = {
    "item": (
        "noi dung cong viec", "noi dung cong tac", "danh muc cong tac", "ten cong tac",
        "ten hang muc", "hang muc", "noi dung", "dien giai", "mo ta", "description",
        "boq item", "work item", "ten vat tu", "ten thiet bi",
    ),
    "quantity": ("khoi luong", "so luong", "quantity", "qty", "kl"),
    "unit": ("don vi tinh", "dvt", "don vi", "unit"),
    "unit_price": (
        "don gia du toan", "don gia hop dong", "don gia", "unit price", "rate", "gia don vi",
    ),
    "amount": (
        "thanh tien du toan", "thanh tien", "gia tri du toan", "gia tri", "amount", "total amount",
        "tong gia tri", "gia thanh",
    ),
    "task_ref": ("ma task", "task", "wbs", "ma cong viec", "ma hieu", "ma hang muc"),
}

_SUMMARY_NAME_HINTS = (
    "tong hop", "summary", "phu luc tong hop", "bia", "cover", "muc luc", "tong gia tri",
)
_BOQ_NAME_HINTS = ("boq", "du toan", "khoi luong", "quantity", "bill of quantity", "bill quantities")
_TOTAL_LABELS = (
    "tong", "tong cong", "cong", "subtotal", "grand total", "tong gia tri", "gia tri truoc thue",
    "gia tri sau thue", "thue gtgt", "vat",
)


def _header_kind(value: Any) -> str | None:
    text = _norm(value)
    if not text:
        return None
    matches: list[tuple[int, str]] = []
    for kind, aliases in _HEADER_ALIASES.items():
        for alias in aliases:
            if text == alias or (len(alias) >= 4 and alias in text):
                matches.append((len(alias), kind))
    return max(matches)[1] if matches else None


def _to_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        try:
            number = float(value)
            return number if math.isfinite(number) else None
        except Exception:
            return None
    text = str(value).strip()
    if not text or text.startswith("="):
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    text = re.sub(r"[^0-9,\.\-]", "", text)
    if not text or text in {"-", ".", ","}:
        return None
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        parts = text.split(",")
        if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3 and len(parts[0]) >= 1):
            text = "".join(parts)
        else:
            text = text.replace(",", ".")
    elif "." in text:
        parts = text.split(".")
        if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3 and len(parts[0]) >= 1):
            text = "".join(parts)
    try:
        number = float(text)
        if negative:
            number = -abs(number)
        return number if math.isfinite(number) else None
    except Exception:
        return None


def _header_mapping(values: list[Any], previous: list[Any] | None = None) -> dict[str, int]:
    candidates: dict[str, tuple[int, int]] = {}
    for idx, value in enumerate(values):
        texts = [value]
        if previous is not None and idx < len(previous):
            texts.append(f"{previous[idx] or ''} {value or ''}")
        for text in texts:
            kind = _header_kind(text)
            if not kind:
                continue
            normalized = _norm(text)
            score = 10 if normalized in _HEADER_ALIASES.get(kind, ()) else 5
            old = candidates.get(kind)
            if old is None or score > old[0] or (score == old[0] and kind == "amount" and idx > old[1]):
                candidates[kind] = (score, idx)
    return {kind: idx for kind, (_, idx) in candidates.items()}


def _mapping_score(mapping: dict[str, int], sheet_name: str) -> int:
    score = 0
    score += 5 if "item" in mapping else 0
    score += 3 if "amount" in mapping else 0
    score += 2 if "quantity" in mapping else 0
    score += 2 if "unit_price" in mapping else 0
    score += 1 if "unit" in mapping else 0
    if any(hint in _norm(sheet_name) for hint in _BOQ_NAME_HINTS):
        score += 2
    return score


def _valid_mapping(mapping: dict[str, int], sheet_name: str) -> bool:
    has_value = "amount" in mapping or ("quantity" in mapping and "unit_price" in mapping)
    if "item" in mapping and has_value:
        return True
    named = any(hint in _norm(sheet_name) for hint in _BOQ_NAME_HINTS)
    return named and "item" in mapping and len(set(mapping) & {"quantity", "unit_price", "amount"}) >= 2


def _find_header(ws, max_scan_rows: int = 45) -> tuple[int, dict[str, int]] | None:
    rows = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, max_scan_rows), values_only=True):
        rows.append(list(row))
    best: tuple[int, int, dict[str, int]] | None = None
    for i, values in enumerate(rows):
        previous = rows[i - 1] if i else None
        mapping = _header_mapping(values, previous)
        if not _valid_mapping(mapping, ws.title):
            continue
        score = _mapping_score(mapping, ws.title)
        candidate = (score, i + 1, mapping)
        if best is None or candidate[0] > best[0]:
            best = candidate
    return (best[1], best[2]) if best else None


def _cell(values: tuple[Any, ...], index: int | None) -> Any:
    if index is None or index < 0 or index >= len(values):
        return None
    return values[index]


def _is_total_label(value: Any) -> bool:
    text = _norm(value)
    if not text:
        return False
    return any(text == token or text.startswith(token + " ") for token in _TOTAL_LABELS)


def _parse_sheet(ws) -> dict[str, Any] | None:
    found = _find_header(ws)
    if not found:
        return None
    header_row, mapping = found
    items: list[dict[str, Any]] = []
    skipped_totals = 0
    for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
        values = tuple(values)
        item_raw = _cell(values, mapping.get("item"))
        item = str(item_raw or "").strip()
        if not item:
            continue
        qty = _to_number(_cell(values, mapping.get("quantity")))
        price = _to_number(_cell(values, mapping.get("unit_price")))
        amount = _to_number(_cell(values, mapping.get("amount")))
        if amount is None and qty is not None and price is not None:
            amount = qty * price
        if amount is None:
            continue
        if _is_total_label(item) and (qty is None or price is None):
            skipped_totals += 1
            continue
        if price is None and qty not in (None, 0):
            price = amount / qty
        items.append(
            {
                "boq_item": item,
                "quantity": float(qty or 0),
                "unit": str(_cell(values, mapping.get("unit")) or "").strip(),
                "unit_price": float(price or 0),
                "budget_total": float(amount),
                "task_ref": str(_cell(values, mapping.get("task_ref")) or "").strip(),
            }
        )
    if not items:
        return None
    return {
        "sheet": ws.title,
        "header_row": header_row,
        "line_count": len(items),
        "budget_total": float(sum(item["budget_total"] for item in items)),
        "items": items,
        "skipped_totals": skipped_totals,
        "summary_name_hint": any(hint in _norm(ws.title) for hint in _SUMMARY_NAME_HINTS),
    }


def parse_boq_workbook(data: bytes, filename: str = "BOQ.xlsx") -> dict[str, Any]:
    raw = bytes(data or b"")
    if not raw:
        raise BOQWorkbookError("File BOQ đang trống.")
    if len(raw) > MAX_WORKBOOK_BYTES:
        raise BOQWorkbookError("File BOQ lớn hơn 100 MB; hãy tách file trước khi nhập.")
    suffix = Path(str(filename or "")).suffix.lower()
    if suffix and suffix not in {".xlsx", ".xlsm"}:
        raise BOQWorkbookError("Chỉ hỗ trợ BOQ Excel .xlsx hoặc .xlsm.")
    try:
        wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    except Exception as exc:
        raise BOQWorkbookError(f"Không đọc được workbook Excel: {exc}") from exc
    try:
        parsed = []
        for ws in wb.worksheets:
            item = _parse_sheet(ws)
            if item:
                parsed.append(item)
    finally:
        wb.close()
    if not parsed:
        raise BOQWorkbookError(
            "Không nhận diện được sheet BOQ. Cần có cột nội dung/hạng mục và cột Thành tiền, "
            "hoặc bộ Khối lượng + Đơn giá."
        )
    normal = [item for item in parsed if not item["summary_name_hint"]]
    selected = normal or parsed
    warnings: list[str] = []
    ignored = [item["sheet"] for item in parsed if item not in selected]
    if ignored:
        warnings.append("Đã bỏ qua sheet tổng hợp để tránh cộng trùng: " + ", ".join(ignored))
    summary = [
        {
            "sheet": item["sheet"],
            "line_count": int(item["line_count"]),
            "budget_total": float(item["budget_total"]),
            "header_row": int(item["header_row"]),
        }
        for item in selected
    ]
    return {
        "filename": Path(str(filename or "BOQ.xlsx")).name,
        "batch_id": hashlib.sha256(raw).hexdigest()[:16],
        "detected_sheets": [item["sheet"] for item in selected],
        "summary": summary,
        "grand_total": float(sum(item["budget_total"] for item in summary)),
        "warnings": warnings,
    }


def build_summary_excel(result: dict[str, Any]) -> bytes:
    summary = list(result.get("summary") or [])
    if not summary:
        raise BOQWorkbookError("Không có dữ liệu tổng hợp để xuất Excel.")
    wb = Workbook()
    ws = wb.active
    ws.title = SUMMARY_SHEET_NAME
    ws.merge_cells("A1:E1")
    ws["A1"] = SUMMARY_SHEET_NAME.upper()
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A2"] = "Nguồn file"
    ws["B2"] = str(result.get("filename") or "")
    headers = ["STT", "Hạng mục / Sheet BOQ", "Số dòng BOQ", "Giá trị dự toán (VND)", "Ghi chú"]
    for col, value in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=value)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for idx, item in enumerate(summary, 1):
        row = 4 + idx
        ws.cell(row=row, column=1, value=idx)
        ws.cell(row=row, column=2, value=str(item.get("sheet") or ""))
        ws.cell(row=row, column=3, value=int(item.get("line_count") or 0))
        ws.cell(row=row, column=4, value=float(item.get("budget_total") or 0))
        ws.cell(row=row, column=5, value=f"Header dòng {int(item.get('header_row') or 0)}")
        ws.cell(row=row, column=4).number_format = '#,##0'
    total_row = 5 + len(summary)
    ws.cell(row=total_row, column=2, value="TỔNG CỘNG").font = Font(bold=True)
    ws.cell(row=total_row, column=4, value=float(result.get("grand_total") or 0)).font = Font(bold=True)
    ws.cell(row=total_row, column=4).number_format = '#,##0'
    widths = [8, 34, 16, 24, 22]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A5"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def save_boq_summary_to_project(
    db,
    project_id: int,
    result: dict[str, Any],
    *,
    replace_existing_excel: bool = True,
) -> dict[str, Any]:
    summary = list(result.get("summary") or [])
    if not summary:
        raise BOQWorkbookError("Không có dữ liệu BOQ để lưu.")
    pid = int(project_id)
    filename = Path(str(result.get("filename") or "BOQ.xlsx")).name.replace("|", "_")
    batch_id = re.sub(r"[^a-fA-F0-9]", "", str(result.get("batch_id") or ""))[:32]
    if not batch_id:
        batch_id = hashlib.sha256(repr(summary).encode("utf-8")).hexdigest()[:16]
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    deleted = 0
    inserted = 0
    with db.connect() as c:
        if replace_existing_excel:
            cur = c.execute(
                "DELETE FROM cost_budgets WHERE project_id=? AND note LIKE ?",
                (pid, AUTO_NOTE_PREFIX + "%"),
            )
        else:
            cur = c.execute(
                "DELETE FROM cost_budgets WHERE project_id=? AND note LIKE ?",
                (pid, f"{AUTO_NOTE_PREFIX}%batch={batch_id}%"),
            )
        try:
            deleted = max(0, int(cur.rowcount or 0))
        except Exception:
            deleted = 0
        sql = (
            "INSERT INTO cost_budgets("
            "project_id,task_ref,boq_item,quantity,unit,unit_price,budget_total,"
            "contract_type,contractor,note,created_at,updated_at"
            ") VALUES(?,?,?,?,?,?,?,?,?,?,?,?)"
        )
        for item in summary:
            sheet = str(item.get("sheet") or "BOQ").replace("|", "_")
            total = float(item.get("budget_total") or 0)
            lines = int(item.get("line_count") or 0)
            note = f"{AUTO_NOTE_PREFIX} file={filename}|sheet={sheet}|lines={lines}|batch={batch_id}"
            c.execute(
                sql,
                (
                    pid, "", sheet, 1.0, "Gói", total, total,
                    "", "", note, now, now,
                ),
            )
            inserted += 1
    return {
        "inserted": inserted,
        "deleted": deleted,
        "batch_id": batch_id,
        "grand_total": float(result.get("grand_total") or 0),
    }
